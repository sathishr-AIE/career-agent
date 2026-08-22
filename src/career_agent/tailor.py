import copy
import json
import re
from pathlib import Path
from typing import Awaitable, Callable

import docx
from docx.text.paragraph import Paragraph

from career_agent.config import CareerBrief
from career_agent.gate import MIN_FACTS_HARD, InsufficientFacts
from career_agent.models import Job, TailorResult
from career_agent import store

TAILOR_PROMPT_VERSION = "tailor-v1"

TEMPLATE = """You are tailoring resume content for one job, from a candidate's
verified facts only.

CAREER BRIEF
Target titles: {titles}

VERIFIED FACTS (numbered by id; cite only these ids, never state anything
not listed here)
{facts}

JOB
Company: {company}
Title: {title}
Description:
{description}

Select the facts most relevant to this job, then write:
- a 2-3 sentence professional summary tailored to this job, built only from
  the facts above
- a list of resume bullets, each rephrased from one or more of the facts
  above for clarity and relevance to this job -- never inventing a claim,
  metric, or skill that is not listed above

Every bullet must cite the id of every fact it draws from.

Reply with a single JSON object and nothing else:
{{"summary": "...", "bullets": [{{"text": "...", "fact_ids": [int, ...]}}]}}
"""


def build_prompt(job: Job, brief: CareerBrief, facts: list[tuple[int, str]]) -> str:
    facts_block = "\n".join(f"[{fid}] {text}" for fid, text in facts)
    return TEMPLATE.format(
        titles=", ".join(brief.target_titles),
        facts=facts_block,
        company=job.company,
        title=job.title,
        description=(job.description or "")[:6000],
    )


def parse_tailor_result(text: str) -> TailorResult:
    """Tolerate prose around the JSON; reject anything that is not a
    TailorResult (including zero bullets or a bullet with no fact_ids,
    both enforced by the model's own min_length constraints)."""
    match = re.search(r"\{.*\}", text, re.DOTALL)
    if not match:
        raise ValueError("no JSON object found in model output")
    try:
        return TailorResult(**json.loads(match.group(0)))
    except Exception as exc:
        raise ValueError(f"could not parse tailoring result: {exc}") from exc


def _validate_fact_ids(result: TailorResult, valid_ids: set[int]) -> None:
    for bullet in result.bullets:
        unknown = [fid for fid in bullet.fact_ids if fid not in valid_ids]
        if unknown:
            raise ValueError(f"bullet cites unknown fact id(s): {unknown}")


async def tailor(job: Job, brief: CareerBrief, facts: list[tuple[int, str]],
                 ask: Callable[[str], Awaitable[str]]) -> TailorResult:
    if len(facts) < MIN_FACTS_HARD:
        raise InsufficientFacts(
            f"{len(facts)} facts recorded, need at least {MIN_FACTS_HARD} to "
            "tailor a resume honestly.")

    valid_ids = {fid for fid, _ in facts}
    prompt = build_prompt(job, brief, facts)
    try:
        result = parse_tailor_result(await ask(prompt))
        _validate_fact_ids(result, valid_ids)
        return result
    except ValueError:
        retry = prompt + ("\n\nYour previous reply was invalid: either not "
                          "valid JSON, missing a required field, or citing a "
                          "fact id not listed above. Reply with ONLY a valid "
                          "JSON object, citing only the fact ids listed.")
        result = parse_tailor_result(await ask(retry))
        _validate_fact_ids(result, valid_ids)
        return result


def _find_marker(doc, marker: str) -> Paragraph:
    for p in doc.paragraphs:
        if p.text.strip() == marker:
            return p
    raise ValueError(f"template is missing the {marker!r} marker paragraph")


def _set_text(paragraph: Paragraph, text: str) -> None:
    """Preserve the paragraph's first run's formatting; drop any others."""
    for run in paragraph.runs[1:]:
        run.text = ""
    if paragraph.runs:
        paragraph.runs[0].text = text
    else:
        paragraph.add_run(text)


def render_docx(template_path: Path, result: TailorResult, out_path: Path) -> None:
    """Deterministic, not the LLM: find the two marker paragraphs by exact
    text, replace <<SUMMARY>>'s text in place, and clone <<PROJECT_BULLET>>
    once per bullet before removing the marker itself. Everything else in
    the template -- header, contact info, education, layout -- is untouched."""
    doc = docx.Document(str(template_path))

    summary_para = _find_marker(doc, "<<SUMMARY>>")
    _set_text(summary_para, result.summary)

    bullet_marker = _find_marker(doc, "<<PROJECT_BULLET>>")
    for bullet in result.bullets:
        clone_el = copy.deepcopy(bullet_marker._p)
        bullet_marker._p.addprevious(clone_el)
        _set_text(Paragraph(clone_el, bullet_marker._parent), bullet.text)
    bullet_marker._p.getparent().remove(bullet_marker._p)

    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out_path))


TEMPLATE_PATH = Path("resume/master.docx")
OUTPUT_DIR = Path("resume/generated")


async def ensure_tailored(conn, job_id: int, job: Job, brief: CareerBrief,
                          ask: Callable[[str], Awaitable[str]]) -> str:
    """Tailor and render this job's resume if one doesn't exist yet, else
    reuse the most recent version -- there is no "Retailor" action. Reads
    TEMPLATE_PATH/OUTPUT_DIR as module attributes (not function defaults)
    so tests can monkeypatch them without needing to reload this module."""
    existing = store.latest_resume_version(conn, job_id)
    if existing:
        return existing

    result = await tailor(job, brief, store.fact_rows(conn), ask)

    version = store.next_resume_version(conn, job_id)
    out_path = OUTPUT_DIR / f"{version}.docx"
    render_docx(TEMPLATE_PATH, result, out_path)

    content = json.dumps({"summary": result.summary,
                          "bullets": [b.model_dump() for b in result.bullets]})
    return store.insert_resume(conn, job_id, version, str(out_path), content)
