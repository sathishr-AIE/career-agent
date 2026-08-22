import docx


def build_tailor_template(path):
    """A minimal DOCX carrying tailor.py's two marker paragraphs, built with
    python-docx rather than committed as a binary fixture -- keeps the
    fixture readable and diffable in the test files that use it."""
    doc = docx.Document()
    doc.add_paragraph("Sathish R -- AI Engineer")
    doc.add_paragraph("<<SUMMARY>>")
    doc.add_paragraph("<<PROJECT_BULLET>>", style="List Bullet")
    doc.add_paragraph("Education")
    doc.save(str(path))
