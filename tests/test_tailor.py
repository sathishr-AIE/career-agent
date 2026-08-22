import json
import sqlite3
from pathlib import Path

import docx
import pytest

from conftest import build_tailor_template
from career_agent import db, store, tailor
from career_agent.config import CareerBrief
from career_agent.gate import InsufficientFacts
from career_agent.models import Bullet, Job, TailorResult
from career_agent.tailor import build_prompt, parse_tailor_result, render_docx
from career_agent.tailor import tailor as tailor_func

BRIEF = CareerBrief(target_titles=["AI Engineer"], search_locations=["Chennai"])
JOB = Job(source="ats", external_id="1", company="Acme", title="AI Engineer",
          description="Build LLM-backed features in Python.")
FACTS = [(i, f"claim {i} (evidence: evidence {i})") for i in range(1, 11)]

GOOD = json.dumps({"summary": "Tailored summary.", "bullets": [
    {"text": "Built an LLM feature", "fact_ids": [1, 2]}]})


def test_build_prompt_includes_fact_ids_and_job_description():
    prompt = build_prompt(JOB, BRIEF, FACTS)
    assert "[1] claim 1" in prompt
    assert "Build LLM-backed features" in prompt


def test_parses_json_out_of_prose():
    r = parse_tailor_result(f"Here it is.\n```json\n{GOOD}\n```")
    assert r.summary == "Tailored summary."
    assert r.bullets[0].fact_ids == [1, 2]


def test_rejects_unparseable():
    with pytest.raises(ValueError):
        parse_tailor_result("I could not decide.")


def test_rejects_zero_bullets():
    with pytest.raises(ValueError):
        parse_tailor_result(json.dumps({"summary": "S", "bullets": []}))


def test_rejects_a_bullet_with_no_fact_ids():
    with pytest.raises(ValueError):
        parse_tailor_result(json.dumps(
            {"summary": "S", "bullets": [{"text": "x", "fact_ids": []}]}))


async def test_raises_below_ten_facts():
    async def ask(_):
        raise AssertionError("must not call the model")

    with pytest.raises(InsufficientFacts):
        await tailor_func(JOB, BRIEF, FACTS[:2], ask)


async def test_retries_once_on_bad_output():
    calls = []

    async def ask(prompt):
        calls.append(prompt)
        return "nonsense" if len(calls) == 1 else GOOD

    r = await tailor_func(JOB, BRIEF, FACTS, ask)
    assert r.summary == "Tailored summary."
    assert len(calls) == 2


async def test_retries_once_on_an_unknown_fact_id():
    calls = []
    bad = json.dumps({"summary": "S", "bullets": [
        {"text": "x", "fact_ids": [999]}]})

    async def ask(prompt):
        calls.append(prompt)
        return bad if len(calls) == 1 else GOOD

    r = await tailor_func(JOB, BRIEF, FACTS, ask)
    assert r.summary == "Tailored summary."
    assert len(calls) == 2


async def test_raises_after_second_failure():
    async def ask(_):
        return "still nonsense"

    with pytest.raises(ValueError):
        await tailor_func(JOB, BRIEF, FACTS, ask)


def test_render_docx_replaces_markers_and_clones_bullets(tmp_path):
    template = tmp_path / "template.docx"
    build_tailor_template(template)

    result = TailorResult(summary="Tailored summary.", bullets=[
        Bullet(text="Bullet one", fact_ids=[1]),
        Bullet(text="Bullet two", fact_ids=[2]),
    ])
    out = tmp_path / "out.docx"
    render_docx(template, result, out)

    doc = docx.Document(str(out))
    texts = [p.text for p in doc.paragraphs]
    assert "<<SUMMARY>>" not in texts
    assert "<<PROJECT_BULLET>>" not in texts
    assert "Tailored summary." in texts
    assert "Bullet one" in texts
    assert "Bullet two" in texts
    # header/footer paragraphs outside the markers are untouched
    assert "Sathish R -- AI Engineer" in texts
    assert "Education" in texts


def test_render_docx_raises_on_a_template_missing_a_marker(tmp_path):
    template = tmp_path / "bad.docx"
    doc = docx.Document()
    doc.add_paragraph("<<SUMMARY>>")  # no <<PROJECT_BULLET>>
    doc.save(str(template))

    result = TailorResult(summary="S", bullets=[Bullet(text="B", fact_ids=[1])])
    with pytest.raises(ValueError):
        render_docx(template, result, tmp_path / "out.docx")


def test_render_docx_creates_missing_output_directories(tmp_path):
    template = tmp_path / "template.docx"
    build_tailor_template(template)
    result = TailorResult(summary="S", bullets=[Bullet(text="B", fact_ids=[1])])
    out = tmp_path / "nested" / "dir" / "out.docx"
    render_docx(template, result, out)
    assert out.exists()


@pytest.fixture
def conn(tmp_path):
    c = db.connect(tmp_path / "t.db")
    db.init_schema(c)
    c.execute("INSERT INTO job (fingerprint, source, external_id, company,"
              " company_normalized, title, title_normalized)"
              " VALUES ('fp1', 'ats', '1', 'Acme', 'acme', 'AI Engineer',"
              " 'aiengineer')")
    for i in range(10):
        c.execute("INSERT INTO fact (claim, evidence) VALUES (?, ?)",
                  (f"claim {i}", f"evidence {i}"))
    c.commit()
    return c


async def _ask(_prompt):
    return json.dumps({"summary": "S", "bullets": [
        {"text": "B1", "fact_ids": [1]}]})


async def test_ensure_tailored_creates_a_resume_row_and_file(tmp_path, conn, monkeypatch):
    template = tmp_path / "master.docx"
    build_tailor_template(template)
    monkeypatch.setattr(tailor, "TEMPLATE_PATH", template)
    monkeypatch.setattr(tailor, "OUTPUT_DIR", tmp_path / "generated")

    version = await tailor.ensure_tailored(conn, 1, JOB, BRIEF, _ask)

    assert version == "tailored-1-r1"
    row = conn.execute("SELECT * FROM resume WHERE version = ?",
                       (version,)).fetchone()
    assert row["job_id"] == 1
    assert Path(row["path"]).exists()
    content = json.loads(row["content"])
    assert content["bullets"][0]["fact_ids"] == [1]


async def test_ensure_tailored_reuses_an_existing_version(conn):
    conn.execute("INSERT INTO resume (version, path, job_id)"
                 " VALUES ('tailored-1-r1', 'x.docx', 1)")
    conn.commit()

    async def must_not_call(_prompt):
        raise AssertionError("must not tailor again")

    version = await tailor.ensure_tailored(conn, 1, JOB, BRIEF, must_not_call)
    assert version == "tailored-1-r1"
