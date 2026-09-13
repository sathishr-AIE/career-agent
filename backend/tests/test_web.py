import datetime as dt
import json
import threading
import time
from pathlib import Path

import pytest
from fastapi.testclient import TestClient

from conftest import build_tailor_template
from career_agent import db, store
from career_agent.config import load_brief, load_candidate_profile
from career_agent.web import app as web
from career_agent.web import pipeline
from career_agent.web import worker


@pytest.fixture
def client(tmp_path, monkeypatch):
    path = tmp_path / "t.db"
    conn = db.connect(path)
    db.init_schema(conn)
    conn.execute("INSERT INTO job (fingerprint, source, external_id, company,"
                 " company_normalized, title, title_normalized, url)"
                 " VALUES ('fp1','ats','1','Acme','acme','AI Engineer',"
                 " 'aiengineer','https://x/1')")
    conn.execute("INSERT INTO assessment (job_id, stage, role_fit, credibility,"
                 " opportunity, application_quality, eligibility_soft,"
                 " weighted_score, verdict, rationale, model, prompt_version)"
                 " VALUES (1,'scored',90,90,90,90,90,90,'submit',"
                 " 'strong match','m','gate-v1')")
    conn.execute("INSERT INTO job (fingerprint, source, external_id, company,"
                 " company_normalized, title, title_normalized, url)"
                 " VALUES ('fp2','ats','2','Globex','globex','ML Engineer',"
                 " 'engineerml','https://x/2')")
    conn.execute("INSERT INTO assessment (job_id, stage, role_fit, credibility,"
                 " opportunity, application_quality, eligibility_soft,"
                 " weighted_score, verdict, rationale, model, prompt_version)"
                 " VALUES (2,'scored',80,50,80,80,80,72,'skip',"
                 " 'credibility below floor','m','gate-v1')")
    for i in range(10):
        conn.execute("INSERT INTO fact (claim, evidence) VALUES (?, ?)",
                     (f"claim {i}", f"evidence {i}"))
    conn.commit()
    monkeypatch.setattr(web, "DB_PATH", path)
    monkeypatch.setattr(web, "CANDIDATE_PROFILE_PATH",
                        tmp_path / "candidate_profile.toml")
    (tmp_path / "candidate_profile.toml").write_text(
        'candidate_name = "Jane Doe"\ncandidate_email = "jane@example.com"\n'
        'candidate_phone = "+91-90000-00000"\n')

    # Every /apply and /override route now tailors before drafting. Default
    # every test to a safe, deterministic tailoring path -- no real LLM
    # call, no CLAUDE_CODE_OAUTH_TOKEN dependency -- so tests that only care
    # about the apply/send state machine get real tailoring for free.
    template = tmp_path / "master.docx"
    build_tailor_template(template)
    monkeypatch.setattr(web.tailor, "TEMPLATE_PATH", template)
    monkeypatch.setattr(web.tailor, "OUTPUT_DIR", tmp_path / "generated")
    monkeypatch.setattr(web.worker.run_module, "verify_auth", lambda: None)

    async def _default_ask(prompt, model=None):
        return json.dumps({"summary": "Tailored summary.",
                           "bullets": [{"text": "Relevant bullet",
                                       "fact_ids": [1]}]})
    monkeypatch.setattr(web.worker.run_module, "_ask", _default_ask)

    return TestClient(web.app)


def test_apply_tailors_before_drafting_and_threads_the_version(client, monkeypatch):
    captured = {}

    async def fake_submit(conn, job_id, mode, brief=None, profile=None,
                          resume_version=None, **kw):
        captured["resume_version"] = resume_version
        return {"ok": True, "job_id": job_id, "status": "draft"}

    monkeypatch.setattr(web.ats_apply, "submit", fake_submit)

    r = client.post("/apply/1")
    assert r.status_code == 200
    assert captured["resume_version"] == "tailored-1-r1"

    conn = db.connect(web.DB_PATH)
    row = conn.execute("SELECT * FROM resume WHERE version = ?",
                       (captured["resume_version"],)).fetchone()
    assert row is not None
    assert Path(row["path"]).exists()


def test_a_captcha_hold_does_not_lose_the_already_tailored_resume(client, monkeypatch):
    """Ordering matters: tailoring is committed before ats_apply.submit is
    even called, so a captcha (or any submission failure) afterward must not
    make the resume row or file disappear."""
    async def held(*args, **kwargs):
        return {"ok": False, "held": True, "reason": "captcha encountered"}

    monkeypatch.setattr(web.ats_apply, "submit", held)

    client.post("/apply/1")

    conn = db.connect(web.DB_PATH)
    row = conn.execute("SELECT * FROM resume WHERE job_id = 1").fetchone()
    assert row is not None
    assert Path(row["path"]).exists()


def test_two_sequential_apply_clicks_reuse_the_same_resume(client, monkeypatch):
    """Verifies ensure_tailored's idempotent-reuse path: a second Apply click
    on the same job finds the first click's resume row via
    latest_resume_version and returns early, rather than retailoring or
    inserting a second row. This is two sequential clicks through the same
    connection, not a concurrency race -- ensure_tailored's early-return
    means the second call never reaches insert_resume's IntegrityError
    fallback. That actual insert-time collision race is covered separately
    by test_insert_resume_on_a_version_collision_returns_the_winner in
    tests/test_store.py, which calls store.insert_resume directly twice with
    the same version to force it."""
    async def fake_submit(conn, job_id, mode, brief=None, profile=None,
                          resume_version=None, **kw):
        return {"ok": True, "job_id": job_id, "status": "draft"}

    monkeypatch.setattr(web.ats_apply, "submit", fake_submit)

    r1 = client.post("/apply/1")
    r2 = client.post("/apply/1")
    assert r1.status_code == 200 and r2.status_code == 200

    conn = db.connect(web.DB_PATH)
    count = conn.execute(
        "SELECT COUNT(*) n FROM resume WHERE job_id = 1").fetchone()["n"]
    assert count == 1, "the second click must reuse the first click's resume"


@pytest.fixture
def running_pipeline(monkeypatch):
    """A run_once that keeps the pipeline in 'running' until the test ends.
    It has to be *releasable*: run_background hands run_once to
    asyncio.to_thread, whose worker thread is non-daemon, so a fake that
    never returns would hold the interpreter open at exit."""
    release = threading.Event()

    async def blocks(args, progress=None):
        release.wait(10)

    monkeypatch.setattr(pipeline.run_module, "run_once", blocks)
    yield
    release.set()


def test_index_shows_submit_and_hold_with_rationale(client):
    r = client.get("/applications")
    assert "AI Engineer" in r.text
    assert "strong match" in r.text


def test_queue_tab_offers_an_apply_button_per_job(client):
    """The Queue tab previously only had priority/Skip controls -- no way to
    manually draft a specific job out of turn without starting the worker."""
    r = client.get("/applications")
    queue_html = r.text.split('id="tab-queue"')[1].split('id="tab-all"')[0]
    assert 'hx-post="/apply/1"' in queue_html


def test_index_hides_skips_by_default(client):
    """The Skipped tab's markup is always present in the response (tabs are
    client-side CSS toggles, per the applications.html design), but the
    Queue tab itself must not list a skip-verdict job."""
    r = client.get("/applications")
    queue_html = r.text.split('id="tab-queue"')[1].split('id="tab-all"')[0]
    assert "ML Engineer" not in queue_html


def test_skipped_view_shows_them(client):
    r = client.get("/applications?show=skipped")
    assert "ML Engineer" in r.text
    assert "credibility below floor" in r.text


def test_dismiss_records_the_human_decision(client):
    r = client.post("/dismiss/1")
    assert r.status_code == 200
    conn = db.connect(web.DB_PATH)
    types = {e["type"] for e in conn.execute("SELECT type FROM event")}
    assert "human_dismissed" in types


def test_override_on_a_skip_records_the_override(client, monkeypatch):
    # Without a stub this ran the real apply engine (Chrome + a paid
    # `claude` session) on every suite run; the event is all it checks.
    async def fake_submit(conn, job_id, mode, brief=None, profile=None,
                          resume_version=None, **kw):
        return {"ok": True, "job_id": job_id, "status": "draft"}

    monkeypatch.setattr(web.ats_apply, "submit", fake_submit)
    r = client.post("/override/2")
    assert r.status_code == 200
    conn = db.connect(web.DB_PATH)
    types = {e["type"] for e in conn.execute("SELECT type FROM event")}
    assert "human_override" in types


def test_plain_apply_refuses_a_skip(client):
    r = client.post("/apply/2")
    assert "skip" in r.text.lower() or "override" in r.text.lower()


def test_submit_failure_message_is_escaped(client, monkeypatch):
    async def boom(*args, **kwargs):
        raise RuntimeError("<script>bad</script>")

    monkeypatch.setattr(web.ats_apply, "submit", boom)
    r = client.post("/override/2")
    assert r.status_code == 200
    assert "<script>" not in r.text


def test_apply_denial_reason_is_escaped(client, monkeypatch):
    async def denied(*args, **kwargs):
        return {"ok": False, "reason": "<script>bad</script>"}

    monkeypatch.setattr(web.ats_apply, "submit", denied)
    r = client.post("/apply/1")
    assert r.status_code == 200
    assert "<script>" not in r.text


def test_apply_parks_the_job_so_the_draft_shows_in_the_status_card(client, monkeypatch):
    """The Queue tab's Apply button posts to /apply/{job_id}, but until this
    sets current_job_id, the status card (the only place Send/Skip render)
    has no idea a draft exists -- it only ever shows the job matching
    current_job_id, same as apply_tick's own park-before-draft pattern."""
    async def fake_submit(conn, job_id, mode, brief=None, profile=None,
                          resume_version=None, **kw):
        return {"ok": True, "job_id": job_id, "status": "draft"}

    monkeypatch.setattr(web.ats_apply, "submit", fake_submit)

    r = client.post("/apply/1")
    assert r.status_code == 200

    conn = db.connect(web.DB_PATH)
    assert worker.get_run_state(conn, "apply")["current_job_id"] == 1

    status = client.get("/run/status")
    assert "AI Engineer" in status.text
    assert 'hx-post="/send/1"' in status.text


def test_apply_refuses_when_another_job_is_already_parked(client, monkeypatch):
    """Without this guard, clicking Apply on a second job while the first
    one is awaiting review would silently overwrite current_job_id and
    orphan the first draft -- invisible in the status card, but still a
    live 'draft' row nobody can find a Send button for."""
    conn = db.connect(web.DB_PATH)
    worker.set_run_state(conn, "apply", current_job_id=2)

    calls = []

    async def spy(*args, **kwargs):
        calls.append(args)
        return {"ok": True, "job_id": 1, "status": "draft"}

    monkeypatch.setattr(web.ats_apply, "submit", spy)

    r = client.post("/apply/1")
    assert r.status_code == 200
    assert "denied" in r.text
    assert calls == [], "must refuse before ever calling submit()"
    assert worker.get_run_state(conn, "apply")["current_job_id"] == 2


@pytest.mark.parametrize("path", ["/send/1", "/api/send/1"])
def test_send_is_retired_in_favour_of_the_chat_review_card(client, monkeypatch, path):
    """S2: a live run's CONFIRM decision is the send. The route stays so an old
    page's button gets a clear answer -- and it never starts a run."""
    calls = []

    async def spy(*args, **kwargs):
        calls.append(args)
        return {"ok": True}

    monkeypatch.setattr(web.ats_apply, "submit", spy)
    conn = db.connect(web.DB_PATH)
    conn.execute("INSERT INTO application (job_id, resume_version, status)"
                 " VALUES (1, 'base-v1', 'draft')")
    conn.commit()
    r = client.post(path)
    assert calls == []
    assert "Answer the review card" in r.text    # Jinja escapes the apostrophe


def test_index_offers_mark_applied_when_a_draft_exists(client):
    """The All Applications tab has no Send/Apply UI (Task 4 removed it): a
    job with only a draft application row has no submitted row, so it is
    untracked and gets the same Mark applied control as any other job."""
    conn = db.connect(web.DB_PATH)
    conn.execute("INSERT INTO application (job_id, resume_version, status)"
                 " VALUES (1, 'v1', 'draft')")
    conn.commit()
    r = client.get("/applications")
    assert '/send/1' not in r.text
    assert 'hx-post="/applied/1"' in r.text
    # and it leaves the Queue tab: a drafted job is in progress, not queued
    queue_html = r.text.split('id="tab-queue"')[1].split('id="tab-all"')[0]
    assert "AI Engineer" not in queue_html


def test_index_requeues_a_job_that_drafted_then_failed(client):
    """submit()'s real-send path never deletes the earlier draft row, so a
    job that drafted and then failed transiently carries both a 'draft' row
    and a later 'failed' row. has_draft must track the LATEST row (here,
    'failed'), not blanket row-membership -- else the job stays permanently
    hidden from the Queue tab, even though worker.next_candidate now treats
    it as retryable again."""
    conn = db.connect(web.DB_PATH)
    conn.execute("INSERT INTO application (job_id, resume_version, status)"
                 " VALUES (1, 'v1', 'draft')")
    conn.execute("INSERT INTO application (job_id, resume_version, status)"
                 " VALUES (1, 'v1', 'failed')")
    conn.commit()
    r = client.get("/applications")
    # back in the Queue tab, like any other retryable job
    queue_html = r.text.split('id="tab-queue"')[1].split('id="tab-all"')[0]
    assert "AI Engineer" in queue_html
    # All Applications: Mark applied/Dismiss again, not a stale Send button
    all_html = r.text.split('id="tab-all"')[1].split('id="tab-skipped"')[0]
    assert 'hx-post="/applied/1"' in all_html
    assert '/send/1' not in all_html


def test_index_shows_held_unknown_and_hides_apply_button(client):
    conn = db.connect(web.DB_PATH)
    conn.execute("INSERT INTO application (job_id, resume_version, status)"
                 " VALUES (1, 'v1', 'held_unknown')")
    conn.commit()
    r = client.get("/applications")
    assert "Held" in r.text
    assert "/apply/1" not in r.text


def test_index_shows_failed_permanent_and_hides_apply_button(client):
    conn = db.connect(web.DB_PATH)
    conn.execute("INSERT INTO application (job_id, resume_version, status)"
                 " VALUES (1, 'v1', 'failed_permanent')")
    conn.commit()
    r = client.get("/applications")
    assert "Failed permanently" in r.text
    assert "/apply/1" not in r.text


def test_scheduled_task_installed_is_false_for_a_missing_task():
    assert web.scheduled_task_installed("NoSuchCareerAgentTask") is False


def test_index_shows_no_schedule_banner_by_default(client):
    r = client.get("/applications")
    assert "No scheduled run is installed" in r.text


def test_index_hides_banner_when_a_schedule_is_installed(client, monkeypatch):
    monkeypatch.setattr(web, "scheduled_task_installed", lambda: True)
    r = client.get("/applications")
    assert "No scheduled run is installed" not in r.text


def test_run_start_sets_status_running_and_ticks_once(client, monkeypatch):
    async def fake_submit(conn, job_id, mode, brief=None, profile=None,
                          resume_version=None, **kw):
        conn.execute("INSERT INTO application (job_id, resume_version,"
                     " status) VALUES (?, 'v1', 'draft')", (job_id,))
        conn.commit()
        return {"ok": True, "job_id": job_id, "status": "draft"}

    monkeypatch.setattr(web.ats_apply, "submit", fake_submit)
    r = client.post("/run/start", data={"mode": "manual"})
    assert r.status_code == 200
    conn = db.connect(web.DB_PATH)
    state = worker.get_run_state(conn, "apply")
    assert state["status"] == "running"
    assert state["mode"] == "manual"
    assert state["current_job_id"] == 1  # job 1 is the higher-scored fixture row


def test_run_pause_sets_status_paused(client, monkeypatch):
    async def fake_submit(conn, job_id, mode, brief=None, profile=None,
                          resume_version=None, **kw):
        conn.execute("INSERT INTO application (job_id, resume_version,"
                     " status) VALUES (?, 'v1', 'draft')", (job_id,))
        conn.commit()
        return {"ok": True, "job_id": job_id, "status": "draft"}

    monkeypatch.setattr(web.ats_apply, "submit", fake_submit)
    client.post("/run/start", data={"mode": "manual"})
    r = client.post("/run/pause")
    assert r.status_code == 200
    conn = db.connect(web.DB_PATH)
    assert worker.get_run_state(conn, "apply")["status"] == "paused"


def test_run_resume_sets_status_running(client, monkeypatch):
    async def fake_submit(conn, job_id, mode, brief=None, profile=None,
                          resume_version=None, **kw):
        conn.execute("INSERT INTO application (job_id, resume_version,"
                     " status) VALUES (?, 'v1', 'draft')", (job_id,))
        conn.commit()
        return {"ok": True, "job_id": job_id, "status": "draft"}

    monkeypatch.setattr(web.ats_apply, "submit", fake_submit)
    client.post("/run/start", data={"mode": "manual"})
    client.post("/run/pause")
    r = client.post("/run/resume")
    assert r.status_code == 200
    conn = db.connect(web.DB_PATH)
    assert worker.get_run_state(conn, "apply")["status"] == "running"


def test_run_stop_clears_current_job(client, monkeypatch):
    async def fake_submit(conn, job_id, mode, brief=None, profile=None,
                          resume_version=None, **kw):
        conn.execute("INSERT INTO application (job_id, resume_version,"
                     " status) VALUES (?, 'v1', 'draft')", (job_id,))
        conn.commit()
        return {"ok": True, "job_id": job_id, "status": "draft"}

    monkeypatch.setattr(web.ats_apply, "submit", fake_submit)
    client.post("/run/start", data={"mode": "manual"})
    r = client.post("/run/stop")
    assert r.status_code == 200
    conn = db.connect(web.DB_PATH)
    state = worker.get_run_state(conn, "apply")
    assert state["status"] == "stopped"
    assert state["current_job_id"] is None


def test_queue_skip_clears_current_job_and_logs(client, monkeypatch):
    async def fake_submit(conn, job_id, mode, brief=None, profile=None,
                          resume_version=None, **kw):
        conn.execute("INSERT INTO application (job_id, resume_version,"
                     " status) VALUES (?, 'v1', 'draft')", (job_id,))
        conn.commit()
        return {"ok": True, "job_id": job_id, "status": "draft"}

    monkeypatch.setattr(web.ats_apply, "submit", fake_submit)
    client.post("/run/start", data={"mode": "manual"})
    r = client.post("/queue/1/skip")
    assert r.status_code == 200
    conn = db.connect(web.DB_PATH)
    assert worker.get_run_state(conn, "apply")["current_job_id"] is None
    types = {e["type"] for e in conn.execute("SELECT type FROM event")}
    assert "job_skipped" in types


def test_queue_skip_advances_to_a_higher_ranked_candidate(client, monkeypatch):
    """Job 1 is stuck as 'current' (e.g. left over from before job 2's
    priority was bumped above it). Skipping job 1 must not just refuse to
    re-tick (that's the single-candidate case above) -- it must pick up the
    now-higher-ranked job 2, proving the fix's job-identity check advances
    the queue rather than freezing it whenever there IS somewhere to go."""
    conn = db.connect(web.DB_PATH)
    conn.execute("UPDATE assessment SET verdict = 'submit' WHERE job_id = 2")
    conn.execute("UPDATE job SET priority = 1 WHERE id = 1")
    conn.execute("UPDATE job SET priority = 0 WHERE id = 2")
    conn.commit()
    worker.set_run_state(conn, "apply", status="running", mode="manual",
                         current_job_id=1)

    async def fake_submit(conn, job_id, mode, brief=None, profile=None,
                          resume_version=None, **kw):
        conn.execute("INSERT INTO application (job_id, resume_version,"
                     " status) VALUES (?, 'v1', 'draft')", (job_id,))
        conn.commit()
        return {"ok": True, "job_id": job_id, "status": "draft"}

    monkeypatch.setattr(web.ats_apply, "submit", fake_submit)
    r = client.post("/queue/1/skip")
    assert r.status_code == 200
    conn = db.connect(web.DB_PATH)
    state = worker.get_run_state(conn, "apply")
    assert state["current_job_id"] == 2


def test_queue_retry_only_accepts_failed_status(client):
    conn = db.connect(web.DB_PATH)
    conn.execute("INSERT INTO application (job_id, resume_version, status)"
                 " VALUES (1, 'v1', 'submitted')")
    conn.commit()
    r = client.post("/queue/1/retry")
    assert r.status_code == 200
    assert "failed" in r.text.lower()


def test_queue_retry_on_a_failed_job_bumps_priority_to_front(client):
    conn = db.connect(web.DB_PATH)
    conn.execute("INSERT INTO application (job_id, resume_version, status)"
                 " VALUES (1, 'v1', 'failed')")
    conn.execute("UPDATE job SET priority = 5 WHERE id = 2")
    conn.commit()
    r = client.post("/queue/1/retry")
    assert r.status_code == 200
    conn = db.connect(web.DB_PATH)
    row = conn.execute("SELECT priority FROM job WHERE id = 1").fetchone()
    assert row["priority"] < 5


def test_queue_priority_swaps_with_neighbor(client):
    conn = db.connect(web.DB_PATH)
    conn.execute("UPDATE assessment SET verdict = 'submit' WHERE job_id = 2")
    conn.execute("UPDATE job SET priority = 0 WHERE id = 1")
    conn.execute("UPDATE job SET priority = 1 WHERE id = 2")
    conn.commit()
    r = client.post("/queue/2/priority", data={"direction": "up"})
    assert r.status_code == 200
    conn = db.connect(web.DB_PATH)
    p1 = conn.execute("SELECT priority FROM job WHERE id = 1").fetchone()["priority"]
    p2 = conn.execute("SELECT priority FROM job WHERE id = 2").fetchone()["priority"]
    assert p2 < p1


def test_queue_priority_ignores_jobs_that_are_not_in_the_queue(client):
    """Job 2's verdict is 'skip', so it isn't in the Queue tab at all.
    Nudging job 1 down must be a no-op rather than a silent swap against a
    row the user cannot see -- which is what made the buttons look dead."""
    conn = db.connect(web.DB_PATH)
    conn.execute("UPDATE job SET priority = 0 WHERE id = 1")
    conn.execute("UPDATE job SET priority = 1 WHERE id = 2")
    conn.commit()
    r = client.post("/queue/1/priority", data={"direction": "down"})
    assert r.status_code == 200
    conn = db.connect(web.DB_PATH)
    p1 = conn.execute("SELECT priority FROM job WHERE id = 1").fetchone()["priority"]
    p2 = conn.execute("SELECT priority FROM job WHERE id = 2").fetchone()["priority"]
    assert p1 < p2  # unchanged: the skipped job was never a neighbor


def test_queue_priority_refuses_a_job_that_is_not_a_candidate(client):
    conn = db.connect(web.DB_PATH)
    conn.execute("INSERT INTO application (job_id, resume_version, status)"
                 " VALUES (1, 'v1', 'draft')")
    conn.commit()
    r = client.post("/queue/1/priority", data={"direction": "up"})
    assert "not in the queue" in r.text.lower()


def test_run_start_clears_a_stale_last_error(client, monkeypatch):
    async def fake_submit(conn, job_id, mode, brief=None, profile=None,
                          resume_version=None, **kw):
        conn.execute("INSERT INTO application (job_id, resume_version,"
                     " status) VALUES (?, 'v1', 'draft')", (job_id,))
        conn.commit()
        return {"ok": True, "job_id": job_id, "status": "draft"}

    monkeypatch.setattr(web.ats_apply, "submit", fake_submit)
    conn = db.connect(web.DB_PATH)
    worker.set_run_state(conn, "apply", status="error",
                         last_error="browser crashed")
    client.post("/run/start", data={"mode": "manual"})
    conn = db.connect(web.DB_PATH)
    assert worker.get_run_state(conn, "apply")["last_error"] is None


def _mode_toggle(html: str) -> str:
    return html.split('class="mode-toggle"')[1].split("</span>")[0]


def test_applications_page_offers_a_mode_toggle(client):
    toggle = _mode_toggle(client.get("/applications").text)
    assert 'name="mode"' in toggle
    assert 'value="auto"' in toggle
    assert 'value="manual"' in toggle
    # and Start sends whatever the toggle holds
    assert "hx-include=\"[name='mode']\"" in client.get("/run/status").text


def test_mode_toggle_is_disabled_while_a_run_is_running(client, monkeypatch):
    async def fake_submit(conn, job_id, mode, brief=None, profile=None,
                          resume_version=None, **kw):
        conn.execute("INSERT INTO application (job_id, resume_version,"
                     " status) VALUES (?, 'v1', 'draft')", (job_id,))
        conn.commit()
        return {"ok": True, "job_id": job_id, "status": "draft"}

    monkeypatch.setattr(web.ats_apply, "submit", fake_submit)
    assert "disabled" not in _mode_toggle(client.get("/applications").text)
    client.post("/run/start", data={"mode": "manual"})
    assert "disabled" in _mode_toggle(client.get("/applications").text)


@pytest.mark.parametrize("mode", ["auto", "manual"])
def test_run_start_records_the_mode_it_was_given(client, monkeypatch, mode):
    async def fake_submit(conn, job_id, mode, brief=None, profile=None,
                          resume_version=None, **kw):
        status = "draft" if mode == "manual" else "submitted"
        conn.execute("INSERT INTO application (job_id, resume_version,"
                     " status) VALUES (?, 'v1', ?)", (job_id, status))
        conn.commit()
        return {"ok": True, "job_id": job_id, "status": status}

    monkeypatch.setattr(web.ats_apply, "submit", fake_submit)
    client.post("/run/start", data={"mode": mode})
    conn = db.connect(web.DB_PATH)
    assert worker.get_run_state(conn, "apply")["mode"] == mode


def test_run_status_shows_idle_start_button(client):
    r = client.get("/run/status")
    assert r.status_code == 200
    assert "Start" in r.text
    assert 'hx-post="/run/start"' in r.text


def test_run_status_shows_pause_button_and_current_job_when_running(client, monkeypatch):
    async def fake_submit(conn, job_id, mode, brief=None, profile=None,
                          resume_version=None, **kw):
        conn.execute("INSERT INTO application (job_id, resume_version,"
                     " status) VALUES (?, 'v1', 'draft')", (job_id,))
        conn.commit()
        return {"ok": True, "job_id": job_id, "status": "draft"}

    monkeypatch.setattr(web.ats_apply, "submit", fake_submit)
    client.post("/run/start", data={"mode": "manual"})
    r = client.get("/run/status")
    assert "Pause" in r.text
    assert "AI Engineer" in r.text  # current job's title, from the fixture


def test_applications_page_has_nav_and_tabs(client):
    r = client.get("/applications")
    assert 'href="/"' in r.text  # Dashboard nav link
    assert 'href="/applications"' in r.text
    assert "Queue" in r.text and "Skipped" in r.text


def test_applications_page_embeds_run_status_polling(client):
    r = client.get("/applications")
    assert 'hx-get="/run/status"' in r.text
    assert "every 2s" in r.text


def test_applications_page_shows_career_brief_panel(client):
    r = client.get("/applications")
    assert "Career Brief" in r.text


def test_applications_page_shows_no_rollout_note_while_flag_is_off(client):
    r = client.get("/applications")
    assert "outcomes recorded" not in r.text


def test_applications_page_shows_the_rollout_note_once_the_flag_is_on(
        client, monkeypatch):
    monkeypatch.setattr(web.ats_apply, "SUBMISSION_IMPLEMENTED", True)
    r = client.get("/applications")
    assert "outcomes recorded" in r.text


def test_failed_skipped_counts_gate_skips_until_they_are_overridden(client, monkeypatch):
    """Job 2 is a skip nobody acted on -- 'Failed/Skipped' should say so
    instead of reporting 0 until an application actually fails."""
    conn = db.connect(web.DB_PATH)
    stats = web._run_status_context(conn)["stats"]
    assert stats["failed_skipped"] == 1
    assert stats["total_applied"] == stats["successful"] == 0

    async def fake_submit(conn, job_id, mode, brief=None, profile=None,
                          resume_version=None, **kw):
        conn.execute("INSERT INTO application (job_id, resume_version,"
                     " status) VALUES (?, 'v1', 'draft')", (job_id,))
        conn.commit()
        return {"ok": True, "job_id": job_id, "status": "draft"}

    monkeypatch.setattr(web.ats_apply, "submit", fake_submit)
    client.post("/override/2")
    conn = db.connect(web.DB_PATH)
    assert web._run_status_context(conn)["stats"]["failed_skipped"] == 0


def test_run_status_shows_stats(client, monkeypatch):
    async def fake_submit(conn, job_id, mode, brief=None, profile=None,
                          resume_version=None, **kw):
        conn.execute("INSERT INTO application (job_id, resume_version,"
                     " status) VALUES (?, 'v1', 'submitted')", (job_id,))
        conn.commit()
        return {"ok": True, "job_id": job_id, "status": "submitted"}

    monkeypatch.setattr(web.ats_apply, "submit", fake_submit)
    client.post("/run/start", data={"mode": "auto"})
    r = client.get("/run/status")
    assert r.status_code == 200
    assert "Total Applied" in r.text


def test_pipeline_run_now_flips_status_and_logs_synchronously(client, running_pipeline):
    r = client.post("/pipeline/run-now")
    assert r.status_code == 200
    conn = db.connect(web.DB_PATH)
    assert worker.get_run_state(conn, "pipeline")["status"] == "running"
    types = {e["type"] for e in conn.execute("SELECT type FROM event")}
    assert "pipeline_started" in types


def test_pipeline_run_now_refuses_a_second_concurrent_run(client, running_pipeline):
    client.post("/pipeline/run-now")
    r = client.post("/pipeline/run-now")
    assert "already" in r.text.lower()


def test_pipeline_status_shows_run_now_button_when_idle(client):
    r = client.get("/pipeline/status")
    assert r.status_code == 200
    assert 'hx-post="/pipeline/run-now"' in r.text
    assert "Run Now" in r.text


def test_pipeline_status_shows_running_state(client, running_pipeline):
    client.post("/pipeline/run-now")
    r = client.get("/pipeline/status")
    assert "Running" in r.text


def test_pipeline_status_renders_the_stage_track(client):
    r = client.get("/pipeline/status")
    for label in ("Discover", "Clean", "Filter", "Score", "Ready"):
        assert label in r.text


def test_pipeline_status_renders_the_counters(client):
    conn = db.connect(web.DB_PATH)
    worker.set_run_state(conn, "pipeline", stage="score", found=91,
                         duplicates=7, passed=27, scored=14, shortlisted=6)
    r = client.get("/pipeline/status")
    assert "Jobs Found" in r.text and "91" in r.text
    assert "Passed Filter" in r.text and "27" in r.text
    assert "Shortlisted" in r.text and "6" in r.text


def test_pipeline_status_renders_the_activity_feed(client):
    conn = db.connect(web.DB_PATH)
    # The feed is scoped to the current run (started_at or later) so a
    # freshly started run never shows a previous run's lines -- see
    # test_the_activity_feed_is_scoped_to_the_current_run below. That scoping
    # needs a real started_at for this event to be in range.
    worker.set_run_state(conn, "pipeline", status="running")
    conn.execute("UPDATE run_state SET started_at = datetime('now', '-1 minute')"
                 " WHERE kind = 'pipeline'")
    conn.execute("INSERT INTO event (job_id, type, payload)"
                 " VALUES (NULL, 'pipeline_progress', 'Discovery returned 40.')")
    conn.commit()
    r = client.get("/pipeline/status")
    assert "Discovery returned 40." in r.text


def test_the_no_auto_apply_note_is_present(client):
    r = client.get("/pipeline/status")
    assert "No automatic applications" in r.text


def test_progress_bar_interpolates_within_the_score_stage(client):
    """Review fix: pct used to jump straight to the end of the score band
    ((idx+1)*100//5 == 80) and sit there for the run's longest phase. It must
    now move as `scored` climbs toward max_score_per_run (25, the fixture's
    default setting), not just jump between five fixed positions."""
    conn = db.connect(web.DB_PATH)
    worker.set_run_state(conn, "pipeline", stage="score", scored=0)
    start = client.get("/pipeline/status").text
    start_pct = int(start.split('style="width:')[1].split('%')[0])

    worker.set_run_state(conn, "pipeline", stage="score", scored=20)
    later = client.get("/pipeline/status").text
    later_pct = int(later.split('style="width:')[1].split('%')[0])

    assert start_pct == 60   # band start: 3 whole stages done (3*20)
    assert later_pct == 76   # 60 + (20 * 20 // 25)


def test_the_activity_feed_is_scoped_to_the_current_run(client):
    """Review fix: the feed query had no run scoping, so opening the monitor
    on a freshly started run showed the *previous* run's twelve lines next to
    all-zero counters. Also locks in newest-first DOM order, which .activity's
    flex-direction:column-reverse (base.html) depends on to stay anchored to
    the newest line without JS re-scrolling on every 3s poll."""
    conn = db.connect(web.DB_PATH)
    conn.execute("INSERT INTO event (job_id, type, payload, occurred_at)"
                 " VALUES (NULL, 'pipeline_progress', 'stale from last run',"
                 " datetime('now', '-1 hour'))")
    conn.commit()

    worker.set_run_state(conn, "pipeline", status="running", stage="discover")
    conn.execute("UPDATE run_state SET started_at = datetime('now', '-1 minute')"
                 " WHERE kind = 'pipeline'")
    conn.execute("INSERT INTO event (job_id, type, payload)"
                 " VALUES (NULL, 'pipeline_progress', 'first message')")
    conn.execute("INSERT INTO event (job_id, type, payload)"
                 " VALUES (NULL, 'pipeline_progress', 'second message')")
    conn.commit()

    r = client.get("/pipeline/status")
    assert "stale from last run" not in r.text
    assert r.text.index("second message") < r.text.index("first message")


def test_run_now_discards_the_response_instead_of_blanking_the_monitor(client):
    """Review fix: the button used to hx-target="#pipeline-status"
    hx-swap="outerHTML" the bare string "ok" over the div this task widened
    from a two-element strip into the entire monitor -- nuking the
    track/counters/feed/note for up to 3s on every click. hx-swap="none"
    discards the response; the poller's own independent 3s tick picks up the
    new state instead."""
    button = client.get("/pipeline/status").text.split("Run Now")[0]
    assert 'hx-swap="none"' in button
    assert "outerHTML" not in button


def test_the_run_now_button_is_reachable_in_a_rendered_page(client):
    """The Critical review finding: `.live-overlay{display:none}` used to
    hide everything inside #runOverlay unconditionally, including the
    idle-state Run Now button -- the only element that can ever add `.open`.
    Every path to opening the modal was unreachable.

    No other test in this suite can catch that class of bug: TestClient only
    ever inspects HTML text, never the CSS that decides what is actually
    visible. This is the one check that renders the real page and looks at
    computed visibility, via a headless, offline (all network blocked)
    Playwright browser -- already a hard dependency of this project
    (career_agent.apply.ats) -- against the exact HTML GET / returns, no
    live server needed. Skips itself if a browser isn't installed rather
    than failing the suite in an environment that lacks one."""
    try:
        from playwright.sync_api import sync_playwright
    except Exception as exc:
        pytest.skip(f"playwright not importable: {exc}")

    html = client.get("/").text

    with sync_playwright() as p:
        try:
            browser = p.chromium.launch(headless=True)
        except Exception as exc:
            pytest.skip(f"no browser available for playwright: {exc}")

        try:
            page = browser.new_page()
            page.route("**/*", lambda route: route.abort())  # fully offline
            page.set_content(html)

            run_now = page.locator("#pipeline-status .btn.primary")
            assert run_now.is_visible(), (
                "Run Now must be visible in the closed (idle) state -- "
                "it is the only path to opening the modal")
            assert not page.locator(".live-head").is_visible(), (
                "modal-only chrome must stay hidden while closed")
            assert not page.locator("#monitor-body").is_visible()

            run_now.click()

            overlay_class = (
                page.locator("#runOverlay").get_attribute("class") or "")
            assert "open" in overlay_class.split(), "click must open the modal"
            assert page.locator(".live-head").is_visible(), (
                "modal chrome must appear once open")
            assert page.locator("#monitor-body").is_visible()
        finally:
            browser.close()


def test_startup_clears_a_pipeline_run_stranded_by_a_crash(client):
    """A process that dies mid-run leaves run_state stuck at 'running', and
    there are no pause/resume/stop endpoints for the pipeline -- the Run Now
    button stays a disabled "Running…" forever. Startup is the only place
    that can tell the difference, so it clears it."""
    conn = db.connect(web.DB_PATH)
    worker.set_run_state(conn, "pipeline", status="running", last_error=None)

    with TestClient(web.app):  # entering the context manager runs lifespan
        pass

    conn = db.connect(web.DB_PATH)
    state = worker.get_run_state(conn, "pipeline")
    assert state["status"] == "error"
    assert "restart" in state["last_error"].lower()
    # and the page offers Run Now again rather than a dead disabled button
    assert 'hx-post="/pipeline/run-now"' in client.get("/pipeline/status").text


def test_root_renders_overview_not_a_redirect(client):
    r = client.get("/", follow_redirects=False)
    assert r.status_code == 200
    assert "Dashboard" in r.text


def test_overview_shows_kpi_cards(client):
    r = client.get("/")
    assert "Discovered" in r.text
    assert "Shortlisted" in r.text


def test_overview_shows_recent_discoveries_from_fixture(client):
    r = client.get("/")
    assert "AI Engineer" in r.text  # job 1 from the client fixture
    assert "Acme" in r.text


def test_overview_ready_to_apply_cta_links_to_applications(client):
    r = client.get("/")
    assert 'href="/applications"' in r.text


def test_overview_embeds_pipeline_status_polling(client):
    r = client.get("/")
    assert 'hx-get="/pipeline/status"' in r.text
    # must be innerHTML, not outerHTML — outerHTML replaces the polling
    # element itself and htmx never re-fires the poll after that (this
    # exact regression happened once already, in the sibling plan's
    # /run/status poller)
    assert 'hx-swap="innerHTML"' in r.text


def test_the_modal_shell_is_outside_the_poller(client):
    """Open/closed is client state and the poller replaces its target every
    3s. Inside the polled region, every swap would slam the modal shut or
    reopen one the user closed -- the bug already fixed once for the
    Auto/Manual toggle."""
    r = client.get("/")
    before = r.text.split('id="pipeline-status-poller"')[0]
    assert 'id="runOverlay"' in before, "modal shell precedes the poller"


def test_overview_page_render_shows_the_pipeline_stage_track(client):
    """Review fix: overview.html includes _pipeline_status.html (which reads
    a route-supplied `stages` -- the 5-item pipeline stage list) and, further
    down the *same* template, had its own unrelated `{% set stages = [...] %}`
    for the KPI funnel. The include only resolved to the route-supplied value
    because the funnel's {% set %} sat below it in the template; reordering
    them would have silently fed the funnel's (label, count) pairs into the
    stage track instead. Every other stage-track test (e.g.
    test_pipeline_status_renders_the_stage_track above) hits /pipeline/status
    directly, whose fragment render never sees the parent template's
    variables at all -- so none of them could ever have caught this. Only a
    full render of / can."""
    r = client.get("/")
    for label in ("Discover", "Clean", "Filter", "Score", "Ready"):
        assert f'<div class="stage-label">{label}</div>' in r.text


def test_still_running_toast_clears_when_the_run_finishes_or_reopened(client):
    """Review fix: closeRunMonitor() added #runToast's `open` class when the
    modal was dismissed mid-run, but nothing ever removed it -- not when the
    run finished, not when the modal was reopened. Live-reproduced: after the
    poller reports data-status="idle", the toast kept saying "Career Agent is
    still running..." forever, right next to an Idle pill and a Run Now
    button that both correctly showed the run was over -- recreating, in the
    toast, the exact "can't tell progress from a hang" problem this whole
    plan exists to fix."""
    try:
        from playwright.sync_api import sync_playwright
    except Exception as exc:
        pytest.skip(f"playwright not importable: {exc}")

    conn = db.connect(web.DB_PATH)
    worker.set_run_state(conn, "pipeline", status="running")
    conn.execute("UPDATE run_state SET started_at = datetime('now')"
                 " WHERE kind = 'pipeline'")
    conn.commit()
    html = client.get("/").text

    with sync_playwright() as p:
        try:
            browser = p.chromium.launch(headless=True)
        except Exception as exc:
            pytest.skip(f"no browser available for playwright: {exc}")

        try:
            page = browser.new_page()
            page.route("**/*", lambda route: route.abort())  # fully offline
            page.set_content(html)

            def toast_open():
                cls = page.locator("#runToast").get_attribute("class") or ""
                return "open" in cls.split()

            # Dismiss while a run is going -> the toast appears.
            page.evaluate("closeRunMonitor()")
            assert toast_open(), "dismissing mid-run must show the toast"

            # The run finishes. In reality the poller's next swap would set
            # this; set it directly to isolate the setInterval logic itself
            # from the (already-covered-elsewhere) polling mechanics.
            page.evaluate(
                "document.getElementById('monitor-body').dataset.status = 'idle'")

            # Nobody re-opens or re-dismisses the modal here -- only the
            # once-a-second interval is running. It alone must notice and
            # clear the toast, well within a couple of ticks.
            page.wait_for_function(
                "() => !document.getElementById('runToast')"
                ".classList.contains('open')", timeout=2000)

            # Re-arm: dismiss again while running.
            page.evaluate(
                "document.getElementById('monitor-body').dataset.status = 'running'")
            page.evaluate("closeRunMonitor()")
            assert toast_open()

            # Reopening must clear it immediately, with no poll tick needed.
            page.evaluate("openRunMonitor()")
            assert not toast_open()
        finally:
            browser.close()


def test_run_now_repolls_immediately_instead_of_waiting_three_seconds(client):
    """Review fix: hx-post="/pipeline/run-now" keeps hx-swap="none" (kept
    from the previous fix round -- see
    test_run_now_discards_the_response_instead_of_blanking_the_monitor above
    -- so the raw "ok" body never blanks the monitor), which meant nothing
    refreshed #monitor-body until the poller's own next 3s tick. Live-
    reproduced: clicking Run Now from an *error* state opened the modal
    showing the *previous* run's stale error, stale stage, and stale
    percentage, under a "Career Agent is running" title, for that whole
    window.

    The reviewer's first-suggested mechanism --
    htmx.trigger('#pipeline-status-poller','load') -- turns out to be a
    no-op here: in the actual shipped htmx 2.0.4 build (unpkg.com/htmx.org@
    2.0.4, vendored below), hx-trigger's "load" keyword is handled by
    addTriggerHandler() as a one-shot initializer gated on
    `!nodeData.firstInitCompleted` that never calls addEventListener --
    unlike e.g. "revealed" or "intersect", which do. So it fires once at
    page load and, critically, registers no listener a later
    htmx.trigger(el, 'load') could ever hit. The fix instead adds a plain
    custom trigger name ("run-started") to the poller's hx-trigger list --
    which *does* go through the normal addEventListener path -- and fires
    that from the button's hx-on::after-request.

    This runs the real, unmodified htmx build against a real DOM in headless
    Chromium (mocking only the two HTTP endpoints, not htmx itself), so a
    wrong event name -- like the original "load" suggestion -- actually
    fails this test, the way a markup-only substring assertion could not."""
    try:
        from playwright.sync_api import sync_playwright
    except Exception as exc:
        pytest.skip(f"playwright not importable: {exc}")

    htmx_js = Path(__file__).parent / "vendor" / "htmx-2.0.4.min.js"
    if not htmx_js.exists():
        pytest.skip("vendored htmx build not available")

    conn = db.connect(web.DB_PATH)
    worker.set_run_state(conn, "pipeline", status="error",
                         last_error="browser crashed", stage="filter")
    conn.commit()
    html = client.get("/").text
    error_fragment = client.get("/pipeline/status").text
    assert "browser crashed" in html  # sanity: the stale error really is there
    assert "browser crashed" in error_fragment

    worker.set_run_state(conn, "pipeline", status="running", last_error=None,
                         stage=None, found=0, duplicates=0, passed=0,
                         scored=0, shortlisted=0)
    conn.execute("UPDATE run_state SET started_at = datetime('now')"
                 " WHERE kind = 'pipeline'")
    conn.commit()
    running_fragment = client.get("/pipeline/status").text
    assert "browser crashed" not in running_fragment
    assert 'data-status="running"' in running_fragment

    # #pipeline-status-poller's own hx-trigger="load, ..." fires a GET the
    # instant the page loads -- before any click -- so a single canned
    # /pipeline/status response can't tell the two apart. Answer that first,
    # automatic poll with the stale error fragment (matching what the real
    # page shows before anyone clicks anything) and only switch to the
    # "running" fragment from the second request on, i.e. the one the click
    # itself must cause.
    status_requests = {"n": 0}

    def handle_status(route):
        status_requests["n"] += 1
        body = error_fragment if status_requests["n"] == 1 else running_fragment
        route.fulfill(status=200, content_type="text/html", body=body)

    with sync_playwright() as p:
        try:
            browser = p.chromium.launch(headless=True)
        except Exception as exc:
            pytest.skip(f"no browser available for playwright: {exc}")

        try:
            page = browser.new_page()
            # Fully offline by default; specific mocks below take priority
            # over this since routes run in reverse registration order.
            page.route("**/*", lambda route: route.abort())
            page.route(
                "https://unpkg.com/htmx.org@2.0.4",
                lambda route: route.fulfill(
                    path=str(htmx_js), content_type="application/javascript"))
            page.route(
                "**/pipeline/run-now",
                lambda route: route.fulfill(
                    status=200, content_type="text/html", body="ok"))
            page.route("**/pipeline/status", handle_status)
            # A real (mocked) origin, not set_content()'s about:blank, so the
            # button's relative hx-post/hx-get URLs resolve predictably to
            # something the routes above actually match.
            page.route(
                "https://career-agent.test/",
                lambda route: route.fulfill(
                    status=200, content_type="text/html", body=html))
            t_nav = time.monotonic()  # htmx's "every 3s" interval starts
            # counting from roughly here (page/htmx init), not from the click
            # below -- so the remaining-budget math has to anchor here too.
            page.goto("https://career-agent.test/")
            page.wait_for_function("() => window.htmx !== undefined")
            # Sanity: the page's first paint (inline fragment) plus the
            # automatic first poll (mocked above) both show the stale error
            # -- so the button below really is starting from an error state,
            # not already "running" before the click ever happens.
            page.wait_for_function(
                "() => { const b = document.getElementById('monitor-body');"
                " return !!b && b.dataset.status === 'error'; }")

            # Budget the post-click wait against the real 3s poll interval,
            # not a guessed constant: click() itself carries a real (and, in
            # this sandboxed environment, fairly large -- ~1.5s, measured
            # directly, reproduces even on a bare unrelated button) input-
            # simulation overhead of its own before it even dispatches, and a
            # fixed timeout picked without accounting for that would either
            # be too tight (flaking on a correct fix) or -- worse -- so loose
            # it silently reaches the real 3s mark and starts passing for a
            # BROKEN fix too, once the ambient interval bails it out. Staying
            # a fixed margin under 3000ms regardless of how much click()
            # itself ate keeps the assertion below meaningful either way.
            # (dispatch_event() avoids that overhead but turns out not to
            # reliably reach htmx's click listener the way a real click
            # does -- confirmed empirically: it fell back to the ambient
            # interval instead, ~3s later.)
            page.click("#pipeline-status .btn.primary")
            remaining_ms = max(300, 2700 - (time.monotonic() - t_nav) * 1000)

            page.wait_for_function(
                "() => { const b = document.getElementById('monitor-body');"
                " return !!b && b.dataset.status === 'running'; }",
                timeout=remaining_ms)
            assert "browser crashed" not in page.content()
        finally:
            browser.close()


BRIEF_TOML = """\
# Keep this comment.
target_titles = ["AI Engineer"]
search_locations = ["Chennai"]
locations = ["Chennai", "Remote"]
remote_ok = true
salary_floor_inr = 1200000
daily_cap = 5
gate_threshold = 72
staleness_days = 30
"""


@pytest.fixture
def brief_path(tmp_path, monkeypatch):
    p = tmp_path / "career_brief.toml"
    p.write_text(BRIEF_TOML, encoding="utf-8")
    monkeypatch.setattr(web, "BRIEF_PATH", p)
    return p


def _form(**overrides):
    # brief_present mirrors the hidden input the rendered Career Brief
    # section always carries when the brief loaded (see brief_path fixture),
    # so this models a real submission of that section, not one with it
    # hidden.
    base = {"target_titles": "AI Engineer", "title_families": "",
            "search_locations": "Chennai", "locations": "Chennai, Remote",
            "work_authorization": "", "excluded_companies": "",
            "non_negotiables": "", "remote_ok": "on",
            "salary_floor_inr": "1200000", "daily_cap": "5",
            "gate_threshold": "72", "staleness_days": "30",
            "scoring_model": "claude-sonnet-5", "max_score_per_run": "25",
            "brief_present": "1"}
    return {**base, **overrides}


def test_settings_page_renders_current_values(client, brief_path):
    r = client.get("/settings")
    assert r.status_code == 200
    assert "Career Brief" in r.text
    assert "Agent Settings" in r.text
    assert "claude-sonnet-5" in r.text
    assert "AI Engineer" in r.text


def test_settings_nav_link_is_active_on_the_settings_page(client, brief_path):
    """base.html emits href="/settings" in the sidebar of EVERY page, so the
    bare link proves nothing about which page rendered. Only the active
    highlight, which comes from active_nav, does."""
    assert '<a href="/settings" class="active">' in client.get("/settings").text
    # ...and it is not lit on a page that isn't Settings
    assert '<a href="/settings" class="active">' not in client.get("/").text


def test_saving_persists_the_agent_settings(client, brief_path):
    r = client.post("/settings", data=_form(
        scoring_model="claude-haiku-4-5", max_score_per_run="40"))
    assert r.status_code == 200
    conn = db.connect(web.DB_PATH)
    s = store.get_settings(conn)
    assert s["scoring_model"] == "claude-haiku-4-5"
    assert s["max_score_per_run"] == 40


def test_saving_writes_the_brief_and_keeps_its_comments(client, brief_path):
    client.post("/settings", data=_form(daily_cap="9",
                                        target_titles="AI Engineer, ML Engineer"))
    text = brief_path.read_text(encoding="utf-8")
    assert "# Keep this comment." in text
    brief = load_brief(brief_path)
    assert brief.daily_cap == 9
    assert brief.target_titles == ["AI Engineer", "ML Engineer"]


def test_an_invalid_brief_leaves_both_stores_untouched(client, brief_path):
    before = brief_path.read_text(encoding="utf-8")
    r = client.post("/settings", data=_form(
        target_titles="",                      # violates min_length=1
        scoring_model="claude-haiku-4-5"))
    assert r.status_code == 200
    assert "target_titles" in r.text
    assert brief_path.read_text(encoding="utf-8") == before
    conn = db.connect(web.DB_PATH)
    assert store.get_settings(conn)["scoring_model"] == "claude-sonnet-5"


def test_an_unknown_model_is_rejected(client, brief_path):
    before = brief_path.read_text(encoding="utf-8")
    r = client.post("/settings", data=_form(scoring_model="gpt-4"))
    assert "scoring_model" in r.text
    assert brief_path.read_text(encoding="utf-8") == before


def test_a_missing_brief_file_does_not_crash_the_page(client, tmp_path,
                                                      monkeypatch):
    """Rendering brief defaults would be a trap: saving them would then
    write a brief the user never chose over the file they lost."""
    monkeypatch.setattr(web, "BRIEF_PATH", tmp_path / "gone.toml")
    r = client.get("/settings")
    assert r.status_code == 200
    assert "could not be read" in r.text
    # the Agent Settings half still works, since it does not need the file
    assert "Scoring model" in r.text


def test_a_non_numeric_daily_cap_is_rejected_without_a_422(client, brief_path):
    """daily_cap is declared as a str Form field and parsed by hand, exactly
    like salary_floor_inr already was -- an int-typed Form field with no
    `required` on its <input> would let FastAPI 422 the request before
    settings_save ever runs, skipping the friendly error page entirely. A
    non-numeric value is what used to 422 under the old int-typed
    declaration, so that's what this exercises; the blank case, which is a
    different bug, is covered by the test below."""
    before = brief_path.read_text(encoding="utf-8")
    r = client.post("/settings", data=_form(daily_cap="lots"))
    assert r.status_code == 200
    assert "daily_cap" in r.text
    assert brief_path.read_text(encoding="utf-8") == before
    conn = db.connect(web.DB_PATH)
    assert store.get_settings(conn)["scoring_model"] == "claude-sonnet-5"


def test_agent_settings_still_save_when_the_brief_is_missing(client, tmp_path,
                                                              monkeypatch):
    """The whole point of brief_present: when the TOML can't be read the
    brief section is hidden, so a real submission from that page carries no
    brief fields and no brief_present marker at all -- and Agent Settings
    must still save rather than being rejected for a blank brief."""
    monkeypatch.setattr(web, "BRIEF_PATH", tmp_path / "gone.toml")
    r = client.post("/settings", data={
        "scoring_model": "claude-haiku-4-5",
        "max_score_per_run": "40"})
    assert r.status_code == 200
    assert "Settings saved" in r.text
    assert "Nothing was saved" not in r.text
    conn = db.connect(web.DB_PATH)
    s = store.get_settings(conn)
    assert s["scoring_model"] == "claude-haiku-4-5"
    assert s["max_score_per_run"] == 40


def test_a_blank_numeric_is_an_error_not_a_silent_revert(client, brief_path):
    """The nastiest shape of the Form()-default trap. FastAPI substitutes a
    field's default for an EMPTY submitted value, so while these fields
    defaulted to "5"/"72"/"30"/"25", clearing the daily cap box and saving
    reported "Settings saved" while quietly reverting four values the user
    never chose -- rewriting the version-controlled TOML and, via
    gate_threshold, changing which jobs get submitted. The defaults are ""
    now, so a cleared box is a field error and nothing is written."""
    before = brief_path.read_text(encoding="utf-8")
    settings_before = dict(store.get_settings(db.connect(web.DB_PATH)))

    r = client.post("/settings", data=_form(
        daily_cap="", gate_threshold="80", staleness_days="45",
        max_score_per_run="40", scoring_model="claude-haiku-4-5"))

    assert r.status_code == 200
    assert "Nothing was saved" in r.text
    assert "Settings saved" not in r.text
    assert "daily_cap" in r.text
    # all-or-nothing: neither store moved, and the TOML is byte-identical
    assert brief_path.read_text(encoding="utf-8") == before
    assert dict(store.get_settings(db.connect(web.DB_PATH))) == settings_before


def test_a_blank_max_score_per_run_is_an_error_too(client, brief_path):
    """max_score_per_run is an Agent Setting, so unlike the brief numerics
    it is required on every POST, brief_present or not."""
    r = client.post("/settings", data=_form(max_score_per_run=""))
    assert "Nothing was saved" in r.text
    assert "max_score_per_run" in r.text


def test_settings_page_shows_blank_candidate_fields_on_first_run(client, tmp_path):
    missing = tmp_path / "no_candidate_profile.toml"
    web.CANDIDATE_PROFILE_PATH = missing
    r = client.get("/settings")
    assert r.status_code == 200
    assert "Candidate Profile" in r.text
    assert "could not be read" not in r.text  # first run is not an error


def test_settings_page_saves_a_new_candidate_profile(client, tmp_path):
    target = tmp_path / "fresh_candidate_profile.toml"
    web.CANDIDATE_PROFILE_PATH = target
    r = client.post("/settings", data={
        "candidate_present": "1", "candidate_name": "Jane Doe",
        "candidate_email": "jane@example.com", "candidate_phone": "+91-1",
        "linkedin_url": "", "portfolio_url": "",
        "scoring_model": "claude-sonnet-5", "max_score_per_run": "25",
    })
    assert r.status_code == 200
    assert "Settings saved" in r.text
    assert target.exists()
    saved = load_candidate_profile(target)
    assert saved.candidate_name == "Jane Doe"


def test_settings_page_rejects_a_blank_candidate_name(client, tmp_path):
    target = tmp_path / "fresh_candidate_profile.toml"
    web.CANDIDATE_PROFILE_PATH = target
    r = client.post("/settings", data={
        "candidate_present": "1", "candidate_name": "",
        "candidate_email": "jane@example.com", "candidate_phone": "+91-1",
        "linkedin_url": "", "portfolio_url": "",
        "scoring_model": "claude-sonnet-5", "max_score_per_run": "25",
    })
    assert "Nothing was saved" in r.text
    assert not target.exists()


def test_settings_saves_the_brief_even_when_candidate_present_is_blank(
        client, brief_path, tmp_path):
    """Reproduces the real page's shape: settings.html renders the
    candidate_present hidden field unconditionally (unlike brief_present,
    which is inside its own {% if brief %} block), so a fresh install with
    no candidate_profile.toml still posts candidate_present=1 with every
    candidate field blank. Before the fix, that tried to build a
    CandidateProfile from all-blank fields, raised ValidationError, and --
    since everything validates before anything saves -- blocked the brief
    and agent settings from saving too, even though the user never touched
    the candidate fields."""
    missing = tmp_path / "no_candidate_profile.toml"
    web.CANDIDATE_PROFILE_PATH = missing
    r = client.post("/settings", data=_form(
        candidate_present="1", candidate_name="", candidate_email="",
        candidate_phone="", linkedin_url="", portfolio_url=""))
    assert r.status_code == 200
    assert "Settings saved" in r.text
    assert "Nothing was saved" not in r.text
    assert not missing.exists()
    conn = db.connect(web.DB_PATH)
    assert store.get_settings(conn)["scoring_model"] == "claude-sonnet-5"


def test_settings_save_preserves_work_history(client, tmp_path):
    """The Settings form carries only the five flat candidate fields; a save
    must merge onto the saved profile, not reset work history to []."""
    from career_agent.config import (CandidateProfile, WorkEntry,
                                     save_candidate_profile)
    target = tmp_path / "full_candidate_profile.toml"
    web.CANDIDATE_PROFILE_PATH = target
    save_candidate_profile(target, CandidateProfile(
        candidate_name="Jane Doe", candidate_email="jane@example.com",
        candidate_phone="+91-1", gender="female",
        work_history=[WorkEntry(company="Acme", title="Engineer")]))
    r = client.post("/settings", data={
        "candidate_present": "1", "candidate_name": "Jane Q Doe",
        "candidate_email": "jane@example.com", "candidate_phone": "+91-1",
        "linkedin_url": "", "portfolio_url": "",
        "scoring_model": "claude-sonnet-5", "max_score_per_run": "25",
    })
    assert "Settings saved" in r.text
    saved = load_candidate_profile(target)
    assert saved.candidate_name == "Jane Q Doe"
    assert saved.gender == "female"
    assert [w.company for w in saved.work_history] == ["Acme"]


def test_the_apply_activity_log_excludes_pipeline_events(client):
    conn = db.connect(web.DB_PATH)
    for i in range(12):
        conn.execute("INSERT INTO event (job_id, type, payload)"
                     " VALUES (NULL, 'pipeline_progress', ?)", (f"step {i}",))
    conn.execute("INSERT INTO event (job_id, type) VALUES (1, 'human_applied')")
    conn.commit()

    r = client.get("/run/status")

    assert "human_applied" in r.text, (
        "a dozen pipeline rows must not push the apply events out of a "
        "LIMIT 10 log")
    assert "pipeline_progress" not in r.text


def test_marking_applied_creates_the_denominator(client):
    r = client.post("/applied/1", data={"when": "2026-08-20"})
    assert r.status_code == 200
    conn = db.connect(web.DB_PATH)
    row = conn.execute(
        "SELECT status, submitted_at FROM application WHERE job_id = 1"
    ).fetchone()
    assert row["status"] == "submitted"
    assert row["submitted_at"] == "2026-08-20"


def test_marking_applied_defaults_to_today(client):
    client.post("/applied/1", data={"when": ""})
    conn = db.connect(web.DB_PATH)
    today = conn.execute("SELECT date('now') d").fetchone()["d"]
    row = conn.execute(
        "SELECT submitted_at FROM application WHERE job_id = 1").fetchone()
    assert row["submitted_at"] == today


def test_marking_applied_blank_date_anchors_to_utc_not_host_local_clock(
        client, monkeypatch):
    """_parse_date's blank-'when' default must agree with the SQL it feeds --
    worker.guard's and this app's own date(submitted_at) = date('now'),
    which is UTC -- not the host's local calendar day. Simulates the window
    (e.g. ~00:00-05:30 IST in Chennai, UTC+5:30) where local has already
    rolled to the next day but UTC has not, by making the two clocks
    disagree on purpose -- deterministic regardless of what day it actually
    is when the suite runs. If _parse_date read the local clock (the old
    bug), submitted_at would land on 2026-01-02 and the assertion below
    would fail even though this test never touches the real system clock."""
    class _FixedLocalDate(dt.date):
        @classmethod
        def today(cls):
            return dt.date(2026, 1, 2)  # local: already the next day

    class _FixedUTCDatetime(dt.datetime):
        @classmethod
        def now(cls, tz=None):
            return dt.datetime(2026, 1, 1, 23, 0, tzinfo=tz)  # UTC: still the 1st

    monkeypatch.setattr(web.dt, "date", _FixedLocalDate)
    monkeypatch.setattr(web.dt, "datetime", _FixedUTCDatetime)

    client.post("/applied/1", data={"when": ""})

    conn = db.connect(web.DB_PATH)
    row = conn.execute(
        "SELECT submitted_at FROM application WHERE job_id = 1").fetchone()
    assert row["submitted_at"] == "2026-01-01"  # UTC's day, not local's


def test_marking_applied_twice_is_refused_readably(client):
    client.post("/applied/1", data={"when": "2026-08-20"})
    r = client.post("/applied/1", data={"when": "2026-08-21"})
    assert r.status_code == 200, "a readable message, not a 500"
    assert "already" in r.text.lower()


def test_a_malformed_date_is_an_error_not_a_422(client):
    r = client.post("/applied/1", data={"when": "last tuesday"})
    assert r.status_code == 200, "friendly error, not FastAPI's raw 422"
    assert "date" in r.text.lower()
    conn = db.connect(web.DB_PATH)
    assert conn.execute(
        "SELECT COUNT(*) n FROM application").fetchone()["n"] == 0


def test_recording_an_outcome_persists_it(client):
    client.post("/applied/1", data={"when": "2026-08-20"})
    conn = db.connect(web.DB_PATH)
    app_id = conn.execute("SELECT id FROM application").fetchone()["id"]

    r = client.post(f"/outcome/{app_id}",
                    data={"type": "screen", "occurred_at": "2026-08-21",
                          "notes": "recruiter call"})
    assert r.status_code == 200
    conn = db.connect(web.DB_PATH)
    row = conn.execute("SELECT type, derived, notes FROM outcome").fetchone()
    assert row["type"] == "screen"
    assert row["derived"] == 0
    assert row["notes"] == "recruiter call"


def test_recording_no_response_by_hand_is_refused(client):
    client.post("/applied/1", data={"when": "2026-08-20"})
    conn = db.connect(web.DB_PATH)
    app_id = conn.execute("SELECT id FROM application").fetchone()["id"]

    r = client.post(f"/outcome/{app_id}",
                    data={"type": "no_response", "occurred_at": "2026-08-21"})
    assert r.status_code == 200
    assert "derived" in r.text.lower()
    conn = db.connect(web.DB_PATH)
    assert conn.execute("SELECT COUNT(*) n FROM outcome").fetchone()["n"] == 0


def test_an_outcome_on_an_unsubmitted_application_is_refused(client):
    """The UI does not offer this, so it guards a forged request."""
    conn = db.connect(web.DB_PATH)
    conn.execute("INSERT INTO application (job_id, resume_version, status)"
                 " VALUES (1, 'base-v1', 'draft')")
    conn.commit()
    app_id = conn.execute("SELECT id FROM application").fetchone()["id"]

    r = client.post(f"/outcome/{app_id}",
                    data={"type": "screen", "occurred_at": "2026-08-21"})
    assert r.status_code == 200
    assert "submitted" in r.text.lower()
    conn = db.connect(web.DB_PATH)
    assert conn.execute("SELECT COUNT(*) n FROM outcome").fetchone()["n"] == 0


def test_all_applications_offers_mark_applied_for_an_untracked_job(client):
    r = client.get("/applications")
    assert 'hx-post="/applied/1"' in r.text
    assert "Mark applied" in r.text


def test_all_applications_offers_outcome_controls_once_submitted(client):
    client.post("/applied/1", data={"when": "2026-08-20"})
    conn = db.connect(web.DB_PATH)
    app_id = conn.execute("SELECT id FROM application").fetchone()["id"]

    r = client.get("/applications")
    assert f'hx-post="/outcome/{app_id}"' in r.text
    assert "Awaiting response" in r.text, "no outcome recorded yet"


def test_the_outcome_form_does_not_offer_no_response(client):
    client.post("/applied/1", data={"when": "2026-08-20"})
    r = client.get("/applications")
    outcome_form = r.text.split('hx-post="/outcome/')[1]
    assert 'value="screen"' in outcome_form
    assert 'value="no_response"' not in outcome_form, (
        "derived, never entered")


def test_a_recorded_outcome_is_shown(client):
    """"Interview" is on the page either way -- every MANUAL_TYPES label is an
    <option> in the outcome form -- so this has to assert on the effective
    outcome label the cell leads with, not on the page."""
    client.post("/applied/1", data={"when": "2026-08-20"})
    conn = db.connect(web.DB_PATH)
    app_id = conn.execute("SELECT id FROM application").fetchone()["id"]
    client.post(f"/outcome/{app_id}",
                data={"type": "interview", "occurred_at": "2026-08-21"})

    r = client.get("/applications")
    label = r.text.split('class="outcome-cell"')[1].split("</span>")[0]
    assert "Interview" in label
    assert "Awaiting response" not in r.text, "an outcome exists now"



def test_marking_applied_clears_the_run_state_so_the_worker_can_advance(
        client, monkeypatch):
    """Manual mode parks the run on a draft, and marking applied is now the
    sanctioned way to resolve one. If it doesn't release current_job_id,
    apply_tick returns early forever and the run is dead."""
    async def fake_submit(conn, job_id, mode, brief=None, profile=None,
                          resume_version=None, **kw):
        conn.execute("INSERT INTO application (job_id, resume_version,"
                     " status) VALUES (?, 'v1', 'draft')", (job_id,))
        conn.commit()
        return {"ok": True, "job_id": job_id, "status": "draft"}

    monkeypatch.setattr(web.ats_apply, "submit", fake_submit)
    client.post("/run/start", data={"mode": "manual"})
    conn = db.connect(web.DB_PATH)
    assert worker.get_run_state(conn, "apply")["current_job_id"] == 1

    r = client.post("/applied/1", data={"when": "2026-08-20"})
    assert r.status_code == 200
    conn = db.connect(web.DB_PATH)
    assert worker.get_run_state(conn, "apply")["current_job_id"] is None


def test_the_skipped_tab_offers_outcome_controls_for_an_applied_job(client):
    """Job 2 is the skip-verdict fixture. An override is the most valuable
    label the system produces, so it has to be able to reach the callback
    denominator from the tab it lives in."""
    client.post("/applied/2", data={"when": "2026-08-20"})
    conn = db.connect(web.DB_PATH)
    app_id = conn.execute(
        "SELECT id FROM application WHERE job_id = 2").fetchone()["id"]

    skipped = client.get("/applications").text.split('id="tab-skipped"')[1]
    assert f'hx-post="/outcome/{app_id}"' in skipped


def test_the_skipped_tab_offers_mark_applied_and_keeps_track_anyway(client):
    skipped = client.get("/applications").text.split('id="tab-skipped"')[1]
    assert 'hx-post="/applied/2"' in skipped
    assert "Track anyway" in skipped


def _all_tab(text):
    """The All Applications panel only. Slicing from 'id="tab-all"' to the
    end would run on into the Skipped panel that follows it in the markup,
    so anything found there would be credited to the wrong tab."""
    return text.split('id="tab-all"')[1].split('id="tab-skipped"')[0]


def test_all_applications_offers_apply_for_an_undrafted_job(client):
    """The Queue tab has a per-row Apply; without one here a job you are
    already looking at has to be found again in the other tab to act on."""
    assert 'hx-post="/apply/1"' in _all_tab(client.get("/applications").text)


def test_all_applications_hides_apply_once_a_draft_exists(client):
    """A drafted job is parked in the status card awaiting review. 'draft'
    is not in ats.BLOCKING, so a second Apply would not be refused -- it
    would quietly insert a second draft row."""
    conn = db.connect(web.DB_PATH)
    conn.execute("INSERT INTO application (job_id, resume_version, status)"
                 " VALUES (1, 'tailored-1-r1', 'draft')")
    conn.commit()

    all_tab = _all_tab(client.get("/applications").text)
    assert 'hx-post="/apply/1"' not in all_tab
    assert "Drafted" in all_tab


def test_all_applications_apply_uses_override_for_a_skip_verdict(client):
    """Loaded as ?show=skipped the route passes all_rows as `jobs`, so the
    All tab includes skip-verdict rows. /apply refuses those ("Use Apply
    anyway to override"), so the button has to post to /override instead."""
    all_tab = _all_tab(client.get("/applications?show=skipped").text)
    assert 'hx-post="/override/2"' in all_tab
    assert 'hx-post="/apply/2"' not in all_tab


def test_the_skipped_tab_does_not_gain_a_duplicate_apply_button(client):
    """Track anyway already posts to /override there; the new button lives
    in the All tab's caller() block, not in the shared macro."""
    skipped = client.get("/applications").text.split('id="tab-skipped"')[1]
    assert skipped.count('hx-post="/override/2"') == 1


def test_marking_applied_says_when_it_hits_the_daily_cap(client, brief_path):
    """A hand-marked row counts against daily_cap like any other and the run
    auto-pauses at it, so the click that caused the pause has to say so."""
    brief_path.write_text(BRIEF_TOML.replace("daily_cap = 5", "daily_cap = 1"),
                          encoding="utf-8")
    r = client.post("/applied/1", data={"when": ""})
    assert "Marked applied" in r.text
    assert "daily cap of 1 reached" in r.text


def test_marking_applied_stays_quiet_below_the_daily_cap(client, brief_path):
    r = client.post("/applied/1", data={"when": ""})
    assert "Marked applied" in r.text
    assert "daily cap" not in r.text


def test_resume_download_serves_the_file(client, tmp_path):
    resume_file = tmp_path / "r.docx"
    resume_file.write_bytes(b"fake docx bytes")
    conn = db.connect(web.DB_PATH)
    conn.execute("INSERT INTO resume (version, path, job_id)"
                 " VALUES ('tailored-1-r1', ?, 1)", (str(resume_file),))
    conn.commit()

    r = client.get("/resume/tailored-1-r1")
    assert r.status_code == 200
    assert r.content == b"fake docx bytes"


def test_resume_download_404s_on_an_unknown_version(client):
    r = client.get("/resume/does-not-exist")
    assert r.status_code == 404


def test_resume_download_404s_when_the_file_is_gone_from_disk(client, tmp_path):
    """A resume row can outlive its file (moved, cleaned up, disk wiped) --
    that must 404 cleanly, not blow up inside FileResponse."""
    conn = db.connect(web.DB_PATH)
    conn.execute("INSERT INTO resume (version, path, job_id)"
                 " VALUES ('tailored-1-r1', ?, 1)",
                 (str(tmp_path / "missing.docx"),))
    conn.commit()

    r = client.get("/resume/tailored-1-r1")
    assert r.status_code == 404


def test_applications_page_links_to_a_tailored_resume(client, monkeypatch):
    async def fake_submit(conn, job_id, mode, brief=None, profile=None,
                          resume_version=None, **kw):
        conn.execute(
            "INSERT INTO application (job_id, resume_version, status)"
            " VALUES (?, ?, 'draft')", (job_id, resume_version))
        conn.commit()
        return {"ok": True, "job_id": job_id, "status": "draft"}

    monkeypatch.setattr(web.ats_apply, "submit", fake_submit)
    client.post("/apply/1")

    r = client.get("/applications")
    assert 'href="/resume/tailored-1-r1"' in r.text


def test_resumes_page_shows_no_master_when_none_exists(client, monkeypatch):
    missing = Path("/no/such/file.docx")
    monkeypatch.setattr(web.tailor, "TEMPLATE_PATH", missing)
    r = client.get("/resumes")
    assert r.status_code == 200
    assert "No master template found" in r.text
    # the full path, not just the basename -- a bare "file.docx" doesn't tell
    # anyone where to put it
    assert str(missing) in r.text


def _upload(client, path, filename=None):
    return client.post("/resumes/master", files={
        "file": (filename or path.name, path.read_bytes(),
                 "application/vnd.openxmlformats-officedocument"
                 ".wordprocessingml.document")})


def test_uploading_a_marker_carrying_docx_installs_the_master(
        client, tmp_path, monkeypatch):
    """The Resumes page previously only reported whether a master existed --
    there was no way to put one there, so Apply failed with 'No master
    template' and the UI offered no way out."""
    target = tmp_path / "resume" / "master.docx"
    monkeypatch.setattr(web.tailor, "TEMPLATE_PATH", target)
    source = tmp_path / "upload.docx"
    build_tailor_template(source)

    r = _upload(client, source)
    assert r.status_code == 200
    assert "denied" not in r.text
    assert target.exists()
    assert target.read_bytes() == source.read_bytes()


def test_a_rejected_upload_does_not_clobber_the_installed_master(
        client, tmp_path, monkeypatch):
    """The property that matters most: a master already in place survives a
    bad upload untouched. Uses a docx auto-preparation cannot rescue (no
    summary section, no bullets) so it reaches the rejection path."""
    target = tmp_path / "resume" / "master.docx"
    target.parent.mkdir(parents=True)
    build_tailor_template(target)
    good_bytes = target.read_bytes()
    monkeypatch.setattr(web.tailor, "TEMPLATE_PATH", target)

    import docx
    unpreparable = tmp_path / "unpreparable.docx"
    doc = docx.Document()
    doc.add_paragraph("SATHISH R")
    doc.add_paragraph("Just a name and nothing else resembling a resume.")
    doc.save(str(unpreparable))

    r = _upload(client, unpreparable)
    assert r.status_code == 200
    assert "denied" in r.text
    assert target.read_bytes() == good_bytes, \
        "a rejected upload must not clobber the master already installed"


def test_uploading_a_non_docx_is_refused_without_a_traceback(
        client, tmp_path, monkeypatch):
    target = tmp_path / "resume" / "master.docx"
    monkeypatch.setattr(web.tailor, "TEMPLATE_PATH", target)
    pdf = tmp_path / "resume.pdf"
    pdf.write_bytes(b"%PDF-1.4 not really a docx")

    r = _upload(client, pdf)
    assert r.status_code == 200
    assert "denied" in r.text
    assert not target.exists()


def test_uploading_leaves_no_temp_file_behind_on_rejection(
        client, tmp_path, monkeypatch):
    """The temp file is written beside the target so os.replace stays a
    same-directory rename; a rejected upload must clean it up rather than
    leaving litter in the user's resume folder."""
    target = tmp_path / "resume" / "master.docx"
    monkeypatch.setattr(web.tailor, "TEMPLATE_PATH", target)
    pdf = tmp_path / "resume.pdf"
    pdf.write_bytes(b"%PDF-1.4 not really a docx")

    _upload(client, pdf)
    leftovers = list(target.parent.glob("*")) if target.parent.exists() else []
    assert leftovers == []


def test_resumes_page_offers_the_upload_form(client, monkeypatch):
    monkeypatch.setattr(web.tailor, "TEMPLATE_PATH", Path("/no/such/file.docx"))
    r = client.get("/resumes")
    assert 'hx-post="/resumes/master"' in r.text
    assert "&lt;&lt;SUMMARY&gt;&gt;" in r.text


def test_resumes_page_shows_the_master_when_present(client, monkeypatch, tmp_path):
    template = tmp_path / "master.docx"
    build_tailor_template(template)
    monkeypatch.setattr(web.tailor, "TEMPLATE_PATH", template)
    r = client.get("/resumes")
    assert "master.docx" in r.text


def test_resumes_page_lists_generated_versions_with_provenance(client, monkeypatch):
    async def fake_submit(conn, job_id, mode, brief=None, profile=None,
                          resume_version=None, **kw):
        return {"ok": True, "job_id": job_id, "status": "draft"}

    monkeypatch.setattr(web.ats_apply, "submit", fake_submit)
    client.post("/apply/1")

    r = client.get("/resumes")
    assert "AI Engineer" in r.text  # job title
    assert "Acme" in r.text          # job company
    assert "tailored-1-r1" in r.text
    assert "Relevant bullet" in r.text  # the fixture's default tailored bullet
    assert "fact 1" in r.text.lower()   # provenance: which fact backs it
    assert "claim 0" in r.text  # fact id 1's actual claim, resolved -- not a bare id
    assert "Tailored summary." in r.text  # the generated summary is shown, not dropped


def test_resumes_page_orders_versions_by_id_not_created_at(client, monkeypatch):
    """Two resumes stamped with the same created_at must still list the more
    recently inserted one first -- id is monotonic, created_at is not."""
    conn = db.connect(web.DB_PATH)
    conn.execute("INSERT INTO resume (version, path, job_id, created_at)"
                "VALUES ('tailored-1-r1', 'a.docx', 1, '2026-01-01 00:00:00')")
    conn.execute("INSERT INTO resume (version, path, job_id, created_at)"
                "VALUES ('tailored-1-r2', 'b.docx', 1, '2026-01-01 00:00:00')")
    conn.commit()

    r = client.get("/resumes")
    first = r.text.index("tailored-1-r2")
    second = r.text.index("tailored-1-r1")
    assert first < second


def test_resumes_nav_link_is_present(client):
    r = client.get("/applications")
    assert 'href="/resumes"' in r.text


def test_answer_route_saves_to_qa_bank_and_unparks(client):
    conn = db.connect(web.DB_PATH)
    job_id = conn.execute(
        "INSERT INTO job (fingerprint, source, external_id, company,"
        " company_normalized, title, title_normalized)"
        " VALUES ('fp9','ats','9','Acme','acme','AI Engineer','aiengineer')"
    ).lastrowid
    conn.execute("INSERT INTO assessment (job_id, stage, weighted_score,"
                 " verdict, rationale, model, prompt_version)"
                 " VALUES (?, 'scored', 90, 'submit', 'r', 'm', 'v1')", (job_id,))
    worker.set_run_state(conn, "apply", status="running", mode="manual",
                         current_job_id=job_id)
    conn.commit()

    r = client.post(f"/answer/{job_id}", data={
        "question": "Notice period?", "answer": "30 days", "is_volatile": "on"})
    assert r.status_code == 200

    row = store.qa_lookup(conn, "Notice period?")
    assert row["answer"] == "30 days"
    assert row["is_volatile"] == 1
    assert worker.get_run_state(conn, "apply")["current_job_id"] is None


def _open_needs_answer(conn, job_id, question):
    """What submit() opens on a needs_answer result."""
    from career_agent import chat
    chat.open_prompt(conn, job_id, "text", {
        "id": "needs_answer", "kind": "text", "question": question,
        "origin": "needs_answer", "options": [], "sensitive": False})


def test_run_status_shows_the_answer_needed_card(client):
    conn = db.connect(web.DB_PATH)
    job_id = conn.execute(
        "INSERT INTO job (fingerprint, source, external_id, company,"
        " company_normalized, title, title_normalized)"
        " VALUES ('fp9','ats','9','Acme','acme','AI Engineer','aiengineer')"
    ).lastrowid
    worker.set_run_state(conn, "apply", status="running", mode="manual",
                         current_job_id=job_id)
    _open_needs_answer(conn, job_id, "Notice period?")

    r = client.get("/run/status")
    assert "Answer needed" in r.text
    assert "Notice period?" in r.text


def test_answering_clears_the_needs_answer_card_before_a_new_draft_lands(client):
    """Regression guard for the stale-card race: the worker re-picks the job
    and sets current_job_id again while it re-tailors/re-drafts (an await
    that can take seconds), so between the answer being saved and the new
    draft landing there is a window with current_job_id set, no draft yet,
    and the OLD needs_answer event still on record. Without the
    needs_answer_resolved marker, that old event is still the latest
    matching row and the card reappears as if unanswered -- risking the
    human re-submitting the answer and orphaning the draft the worker is
    about to insert (see the final-review ledger)."""
    conn = db.connect(web.DB_PATH)
    job_id = conn.execute(
        "INSERT INTO job (fingerprint, source, external_id, company,"
        " company_normalized, title, title_normalized)"
        " VALUES ('fp9','ats','9','Acme','acme','AI Engineer','aiengineer')"
    ).lastrowid
    worker.set_run_state(conn, "apply", status="running", mode="manual",
                         current_job_id=job_id)
    store.log(conn, job_id, "needs_answer", "Notice period?")
    _open_needs_answer(conn, job_id, "Notice period?")

    r = client.post(f"/answer/{job_id}", data={
        "question": "Notice period?", "answer": "30 days"})
    assert r.status_code == 200

    # The worker re-picks the same job and re-parks it before its draft
    # lands -- no draft row exists yet, but current_job_id is set again.
    worker.set_run_state(conn, "apply", current_job_id=job_id)

    r = client.get("/run/status")
    assert "Answer needed" not in r.text
    # The activity log below the status card legitimately keeps showing the
    # old needs_answer event as history -- only the live "Answer needed"
    # card (and its re-submit form) must be gone, so this checks for the
    # form's distinguishing input rather than the question text, which the
    # log line also contains.
    assert 'name="question"' not in r.text


def test_run_status_shows_the_drafted_answers(client):
    """The README's whole pre-flip verification procedure -- draft against a
    real posting, check what the filler produced, before ever flipping
    SUBMISSION_IMPLEMENTED -- needs the drafted answers visible somewhere.
    Nothing else in the dashboard shows application.answers."""
    conn = db.connect(web.DB_PATH)
    job_id = conn.execute(
        "INSERT INTO job (fingerprint, source, external_id, company,"
        " company_normalized, title, title_normalized)"
        " VALUES ('fp9','ats','9','Acme','acme','AI Engineer','aiengineer')"
    ).lastrowid
    conn.execute("INSERT INTO application (job_id, resume_version, answers,"
                 " status) VALUES (?, 'base-v1', ?, 'draft')",
                 (job_id, json.dumps({"#first_name": "Jane"})))
    worker.set_run_state(conn, "apply", status="running", mode="manual",
                         current_job_id=job_id)
    conn.commit()

    r = client.get("/run/status")
    assert "#first_name" in r.text
    assert "Jane" in r.text


def test_do_apply_needs_answer_parks_the_run_so_the_card_shows(client, monkeypatch):
    """_do_apply (the Apply/Override button) tells the user to 'see the
    status card below' on a needs_answer result, but never set
    current_job_id itself -- so the card it points at wasn't rendered
    unless the run happened to already be parked on that exact job."""
    async def needs_answer(conn, job_id, mode, brief=None, profile=None,
                           resume_version=None, **kw):
        _open_needs_answer(conn, job_id, "Notice period?")
        return {"ok": False, "needs_answer": "Notice period?",
                "reason": "needs an answer: Notice period?"}

    monkeypatch.setattr(web.ats_apply, "submit", needs_answer)
    r = client.post("/apply/1")
    assert r.status_code == 200
    assert "Answer needed" in r.text

    conn = db.connect(web.DB_PATH)
    assert worker.get_run_state(conn, "apply")["current_job_id"] == 1

    status = client.get("/run/status")
    assert "Answer needed" in status.text
    assert "Notice period?" in status.text


def _plain_resume(path, summary_heading="PROFESSIONAL SUMMARY", bullets=3):
    """A resume shaped like a real one: a summary heading followed by prose,
    and a run of list-styled bullets. No markers anywhere."""
    import docx
    doc = docx.Document()
    doc.add_paragraph("SATHISH R")
    doc.add_paragraph("Chennai, India | someone@example.com")
    if summary_heading:
        doc.add_paragraph(summary_heading)
        doc.add_paragraph("AI engineer building production LLM systems.")
    doc.add_paragraph("EXPERIENCE")
    for i in range(bullets):
        doc.add_paragraph(f"Did notable thing number {i}.", style="List Paragraph")
    doc.add_paragraph("EDUCATION")
    doc.save(str(path))
    return path


def test_uploading_a_plain_resume_auto_inserts_the_markers(
        client, tmp_path, monkeypatch):
    """A real resume has prose under a 'SUMMARY' heading and real bullets,
    never the literal markers. Rejecting it made the user hand-edit magic
    strings into Word; preparing it here is what makes the upload usable."""
    target = tmp_path / "resume" / "master.docx"
    monkeypatch.setattr(web.tailor, "TEMPLATE_PATH", target)
    source = _plain_resume(tmp_path / "plain.docx")

    r = _upload(client, source)
    assert "denied" not in r.text
    assert target.exists()

    import docx
    texts = [p.text.strip() for p in docx.Document(str(target)).paragraphs]
    assert "<<SUMMARY>>" in texts
    assert texts.count("<<PROJECT_BULLET>>") == 1, \
        "render_docx clones the marker per bullet, so exactly one must remain"


def test_auto_prepare_reports_what_it_replaced(client, tmp_path, monkeypatch):
    """Silent rewriting is the failure mode the user rejected. Saying what
    changed is what makes an automatic edit reviewable."""
    target = tmp_path / "resume" / "master.docx"
    monkeypatch.setattr(web.tailor, "TEMPLATE_PATH", target)
    source = _plain_resume(tmp_path / "plain.docx", bullets=3)

    r = _upload(client, source)
    assert "PROFESSIONAL SUMMARY" in r.text or "summary" in r.text.lower()
    assert "3" in r.text, "must say how many bullets it collapsed"


def test_an_already_marked_docx_is_installed_unchanged(
        client, tmp_path, monkeypatch):
    """Re-uploading an already-prepared file must be idempotent, not
    double-prepared into nonsense."""
    target = tmp_path / "resume" / "master.docx"
    monkeypatch.setattr(web.tailor, "TEMPLATE_PATH", target)
    source = tmp_path / "prepared.docx"
    build_tailor_template(source)

    r = _upload(client, source)
    assert "denied" not in r.text
    assert target.read_bytes() == source.read_bytes()


def test_a_resume_with_no_summary_section_is_still_refused(
        client, tmp_path, monkeypatch):
    """Auto-preparation guesses from structure. When there is no structure
    to read, it must say so rather than pick a paragraph at random."""
    target = tmp_path / "resume" / "master.docx"
    monkeypatch.setattr(web.tailor, "TEMPLATE_PATH", target)
    source = _plain_resume(tmp_path / "nosummary.docx", summary_heading=None)

    r = _upload(client, source)
    assert "denied" in r.text
    assert not target.exists()


def test_a_resume_with_no_bullets_is_refused(client, tmp_path, monkeypatch):
    target = tmp_path / "resume" / "master.docx"
    monkeypatch.setattr(web.tailor, "TEMPLATE_PATH", target)
    source = _plain_resume(tmp_path / "nobullets.docx", bullets=0)

    r = _upload(client, source)
    assert "denied" in r.text
    assert not target.exists()


_SUBMIT_ROUTES = [("/apply/1", {}), ("/api/apply/1", {}),
                  ("/override/1", {}), ("/api/override/1", {})]
_TICK_ROUTES = [("/run/start", {"data": {"mode": "manual"}}),
                ("/api/run/start", {"json": "manual"}),
                ("/run/resume", {}), ("/api/run/resume", {}),
                ("/queue/1/skip", {}), ("/api/queue/1/skip", {})]


@pytest.mark.parametrize("path,kw", _SUBMIT_ROUTES + _TICK_ROUTES)
def test_agent_routes_hand_a_chat_conn_factory_down(client, monkeypatch, path, kw):
    """The agent's narration reaches the job chat only through this factory
    (a fresh, light connection per event, on the runner's reader thread).
    Start/Resume/Skip matter as much as Apply: their click's own tick runs
    the whole first agent session."""
    captured = {}

    async def fake_submit(conn, job_id, mode, conn_factory=None, **kw):
        captured["conn_factory"] = conn_factory
        return {"ok": True, "job_id": job_id, "status": "draft"}

    async def fake_tick(conn, brief_path, profile_path, conn_factory=None):
        captured["conn_factory"] = conn_factory

    monkeypatch.setattr(web.ats_apply, "submit", fake_submit)
    monkeypatch.setattr(web.worker, "apply_tick", fake_tick)
    monkeypatch.setattr(web.worker, "next_candidate", lambda conn: {"job_id": 2})
    conn = db.connect(web.DB_PATH)
    conn.execute("INSERT INTO application (job_id, resume_version, status)"
                 " VALUES (1, 'base-v1', 'draft')")
    conn.commit()
    worker.set_run_state(conn, "apply", status="paused", mode="manual",
                         current_job_id=1)

    client.post(path, **kw)
    assert captured.get("conn_factory") is web._chat_conn


def test_run_status_context_exposes_the_open_prompt_and_conversation(client):
    from career_agent import chat
    conn = db.connect(web.DB_PATH)
    worker.set_run_state(conn, "apply", status="running", mode="manual", current_job_id=1)
    ctx = web._run_status_context(conn)
    assert "needs_answer_question" not in ctx and "draft_answers" not in ctx
    assert ctx["open_prompt"] is None
    assert ctx["conversation_id"] == chat.conversation_for_job(conn, 1)
    pid = chat.open_prompt(conn, 1, "text", {"id": "q", "question": "Notice?"})
    assert web._run_status_context(conn)["open_prompt"] == {
        "id": pid, "kind": "text", "question": "Notice?"}
    chat.expire_open_prompts(conn, 1)
    cpid = chat.open_prompt(conn, 1, "confirm", {"fields": []})
    assert web._run_status_context(conn)["open_prompt"] == {
        "id": cpid, "kind": "confirm", "question": "Review before applying"}
