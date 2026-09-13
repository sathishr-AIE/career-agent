import docx


def build_tailor_template(path):
    """A minimal DOCX carrying tailor.py's two marker paragraphs, built with
    python-docx rather than committed as a binary fixture -- keeps the
    fixture readable and diffable in the test files that use it."""
    doc = docx.Document()
    doc.add_paragraph("Sathish R -- AI Engineer")
    doc.add_paragraph("<<SUMMARY>>")
    doc.add_paragraph("<<PROJECT_BULLET>>", style="List Bullet")
    doc.add_paragraph("Education")
    doc.save(str(path))


import pytest


@pytest.fixture(autouse=True)
def _no_live_apply_agent(monkeypatch):
    """No test may reach the real apply engine (a real Chrome plus a real
    `claude` session): tests/test_web.py's override test did, for every
    run of the suite, and spent real credits each time on the fixture's
    placeholder URL. Tests inject `run_agent` or patch `submit`; anything
    that still gets this far fails loudly instead of spawning."""
    from career_agent.apply import ats

    async def _refuse(prompt, job_id, nonce):
        raise RuntimeError("test reached ats._live_run_agent -- inject run_agent")
    monkeypatch.setattr(ats, "_live_run_agent", _refuse)
