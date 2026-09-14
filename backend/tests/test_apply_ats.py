import asyncio
import json
import logging
import sqlite3
import threading
from pathlib import Path

import docx
import pytest

from career_agent import chat, db, store
from career_agent.apply import agent as agent_mod
from career_agent.apply import ats as ats_apply
from career_agent.apply.agent import AgentResult
from career_agent.config import CandidateProfile, CareerBrief


@pytest.fixture(autouse=True)
def work_dir(tmp_path, monkeypatch):
    """submit() stages a clean copy of the résumé in the agent's work dir
    before handing the path to the prompt. Keep every test's copy inside
    its own tmp_path rather than scattering files in the real temp dir."""
    d = tmp_path / "work"
    monkeypatch.setattr(agent_mod, "WORK_DIR", d)
    return d


@pytest.fixture
def conn(tmp_path):
    # A real file on disk, but not a readable .docx: submit() refuses when
    # the résumé the agent is told to upload does not exist, while
    # _resume_text still has to degrade gracefully on one it can't parse.
    resume = tmp_path / "r.docx"
    resume.write_text("not a real docx", encoding="utf-8")
    c = db.connect(tmp_path / "t.db")
    db.init_schema(c)
    c.execute("INSERT INTO job (fingerprint, source, external_id, company,"
              " company_normalized, title, title_normalized, url)"
              " VALUES ('fp','ats','1','Acme','acme','AI Engineer',"
              " 'aiengineer','https://x/apply')")
    c.execute("INSERT INTO resume (version, path) VALUES ('base-v1', ?)",
              (str(resume),))
    c.commit()
    return c


PROFILE = CandidateProfile(candidate_name="Jane Doe",
                           candidate_email="jane@example.com",
                           candidate_phone="+91-90000-00000")

BRIEF = CareerBrief(target_titles=["AI Engineer"],
                    search_locations=["Chennai"],
                    locations=["Chennai", "Bengaluru"])


def fake_agent(result: AgentResult):
    """The one test seam. Nothing in this file may launch Chrome, spawn a
    subprocess, or need the `claude`/`npx` binaries -- every case injects
    this instead of letting submit() reach _live_run_agent."""
    async def _fake(prompt, job_id, nonce, events, session_id=None, resume=False):
        _fake.prompts.append(prompt)
        _fake.job_ids.append(job_id)
        _fake.nonces.append(nonce)
        return result
    _fake.prompts = []
    _fake.job_ids = []
    _fake.nonces = []
    return _fake


async def _submit(conn, **kw):
    kw.setdefault("brief", BRIEF)
    kw.setdefault("profile", PROFILE)
    return await ats_apply.submit(conn, 1, **kw)


def _apps(conn):
    return conn.execute("SELECT * FROM application WHERE job_id = 1"
                        " ORDER BY id").fetchall()


def _event_types(conn):
    return [r["type"] for r in conn.execute("SELECT type FROM event")]


# The four "the agent stopped reporting mid-run" reasons: unknown state when
# the run could submit (held_unknown), plain retryable failures when it couldn't.
UNKNOWN_STATE = ("agent_error", "timeout", "no_result_line",
                 "unrecognized_result:BLAH")


# -- pure helpers ---------------------------------------------------------

def test_classify_failure():
    assert ats_apply.classify_failure("sso_required", 0) == "failed_permanent"
    assert ats_apply.classify_failure("easy_apply", 0) == "failed_permanent"
    assert ats_apply.classify_failure("stuck", 0) == "failed"
    assert ats_apply.classify_failure(
        "stuck", ats_apply.MAX_ATTEMPTS - 1) == "failed_permanent"


def test_is_unknown_state():
    for reason in UNKNOWN_STATE:
        assert ats_apply.is_unknown_state(reason), reason
    for reason in ("stuck", "page_error", "sso_required", "expired",
                   "login_issue", ""):
        assert not ats_apply.is_unknown_state(reason), reason


def test_confirmed_within_days_boundary(conn):
    """Kept from the resolve_answers era: the naive-UTC window agent.py's
    prompt builder uses to mark a volatile answer stale."""
    store.qa_upsert(conn, "Current CTC?", "12 LPA", is_volatile=True)
    conn.execute("UPDATE qa_bank SET last_confirmed_at ="
                 " datetime('now', '-29 days')")
    conn.commit()
    row = store.qa_lookup(conn, "Current CTC?")
    assert ats_apply._confirmed_within_days(row["last_confirmed_at"], 30)
    conn.execute("UPDATE qa_bank SET last_confirmed_at ="
                 " datetime('now', '-31 days')")
    conn.commit()
    row = store.qa_lookup(conn, "Current CTC?")
    assert not ats_apply._confirmed_within_days(row["last_confirmed_at"], 30)


# -- draft outcomes --------------------------------------------------------

async def test_draft_ready_records_a_draft_row(conn):
    """Answers come only from an approved CONFIRM (see the relay tests below):
    an AgentResult's own answers are never recorded."""
    fake = fake_agent(AgentResult("draft_ready", answers={"Visa?": "Citizen"},
                                  transcript_path="t.txt"))
    r = await _submit(conn, mode="manual", run_agent=fake)
    assert r["ok"] and r["status"] == "draft"
    row = _apps(conn)[0]
    assert row["status"] == "draft"
    assert json.loads(row["answers"]) == {}
    assert row["transcript_path"] == "t.txt"
    assert row["submitted_at"] is None
    assert fake.job_ids == [1]


async def test_with_the_kill_switch_off_the_prompt_cannot_submit(conn, monkeypatch):
    store.qa_upsert(conn, "Notice period?", "30 days", is_volatile=False)
    fake = _use_live_fake(monkeypatch, AgentResult("draft_ready"))
    r = await _submit(conn, mode="manual")
    assert r["ok"] and r["status"] == "draft"
    prompt = fake.prompts[0]
    assert "do NOT click Submit" in prompt    # manual, can_submit=False
    assert "DRAFT_READY" in prompt and "This run is pre-approved:" not in prompt
    assert "== PREVIOUSLY ANSWERED" not in prompt
    assert "notice period -> 30 days" in prompt


async def test_draft_needs_answer_passthrough(conn):
    fake = fake_agent(AgentResult("needs_answer", "Do you have a PMP?"))
    r = await _submit(conn, mode="manual", run_agent=fake)
    assert not r["ok"] and r["needs_answer"] == "Do you have a PMP?"
    assert _apps(conn) == []


async def test_draft_captcha_holds_without_application_row(conn):
    fake = fake_agent(AgentResult("captcha"))
    r = await _submit(conn, mode="manual", run_agent=fake)
    assert not r["ok"] and r["held"]
    assert _apps(conn) == []
    assert "captcha_held" in _event_types(conn)


async def test_expired_at_draft_time_is_permanent(conn):
    fake = fake_agent(AgentResult("expired"))
    r = await _submit(conn, mode="manual", run_agent=fake)
    assert not r["ok"]
    row = _apps(conn)[0]
    assert row["status"] == "failed_permanent"
    assert row["failure_reason"] == "expired"
    assert "failed_permanent" in _event_types(conn)


async def test_draft_failure_records_a_failed_row(conn):
    fake = fake_agent(AgentResult("failed", "stuck"))
    r = await _submit(conn, mode="manual", run_agent=fake)
    assert not r["ok"]
    row = _apps(conn)[0]
    assert row["status"] == "failed"
    assert row["failure_reason"] == "stuck"


async def test_draft_that_claims_it_applied_is_held_not_retried(conn, monkeypatch):
    """With submission disabled the prompt forbids clicking Submit. An APPLIED
    there means the true state is unknown and possibly submitted -- it must
    block, never become a retryable failure that sends a second time."""
    _use_live_fake(monkeypatch, AgentResult("applied"))
    r = await _submit(conn, mode="manual")
    assert not r["ok"]
    row = _apps(conn)[0]
    assert row["status"] == "held_unknown"
    assert row["failure_reason"] == "applied_during_draft"


async def test_a_prior_draft_neither_pins_answers_nor_changes_the_mode(conn):
    """The old draft-then-Send flow is retired: the CONFIRM card is the review.
    A draft on record is not replayed into the prompt (resume answers arrive
    with checkpoints, Task 10) and the caller's mode stands."""
    await _submit(conn, mode="manual",
                  run_agent=fake_agent(AgentResult("draft_ready")))
    fake = fake_agent(AgentResult("applied"))
    r = await _submit(conn, mode="auto", run_agent=fake)
    assert r["ok"] and r["status"] == "submitted"
    assert "== PREVIOUSLY ANSWERED" not in fake.prompts[0]
    assert "This run is pre-approved:" in fake.prompts[0]
    assert [a["status"] for a in _apps(conn)] == ["draft", "submitted"]


@pytest.mark.parametrize("result,expected", [
    (AgentResult("login_issue"), "failed"),
])
async def test_send_dispatch_for_a_non_submitting_outcome(conn, result,
                                                          expected):
    """Nothing was submitted and the state is known -- retryable."""
    r = await _submit(conn, mode="auto", run_agent=fake_agent(result))
    assert not r["ok"]
    row = _apps(conn)[-1]
    assert row["status"] == expected
    assert row["failure_reason"] == (result.reason or result.code)


async def test_auto_mode_prompt_is_pre_approved_and_can_submit(conn):
    fake = fake_agent(AgentResult("applied"))
    r = await _submit(conn, mode="auto", run_agent=fake)
    assert r["ok"]
    assert "== PREVIOUSLY ANSWERED" not in fake.prompts[0]
    assert "This run is pre-approved:" in fake.prompts[0] and "click Submit" in fake.prompts[0]


async def test_send_captcha_removes_the_in_flight_row(conn):
    await _submit(conn, mode="manual",
                  run_agent=fake_agent(AgentResult("draft_ready")))
    r = await _submit(conn, mode="auto",
                      run_agent=fake_agent(AgentResult("captcha")))
    assert not r["ok"] and r["held"]
    assert [a["status"] for a in _apps(conn)] == ["draft"]
    assert "captcha_held" in _event_types(conn)


async def test_send_needs_answer_removes_the_in_flight_row(conn):
    r = await _submit(conn, mode="auto",
                      run_agent=fake_agent(AgentResult("needs_answer", "PMP?")))
    assert not r["ok"] and r["needs_answer"] == "PMP?"
    assert _apps(conn) == []


@pytest.mark.parametrize("mode", ["manual", "auto"])
async def test_a_needs_answer_with_no_question_is_still_truthy(conn, mode):
    """worker.apply_tick parks on `if result.get("needs_answer")`, so a
    bare RESULT:NEEDS_ANSWER: must not report an empty string -- that
    falls through to job_skipped and clears the park."""
    r = await _submit(conn, mode=mode,
                      run_agent=fake_agent(AgentResult("needs_answer", "")))
    assert r["needs_answer"]
    assert _apps(conn) == []


async def test_permanent_failure_writes_failed_permanent(conn):
    fake = fake_agent(AgentResult("failed", "sso_required"))
    r = await _submit(conn, mode="auto", run_agent=fake)
    assert not r["ok"]
    row = _apps(conn)[-1]
    assert row["status"] == "failed_permanent"
    assert row["failure_reason"] == "sso_required"


async def test_retryable_failure_promotes_at_max_attempts(conn):
    for _ in range(ats_apply.MAX_ATTEMPTS):
        r = await _submit(conn, mode="auto",
                          run_agent=fake_agent(AgentResult("failed", "stuck")))
        assert not r["ok"]
    statuses = [a["status"] for a in _apps(conn)]
    assert statuses == ["failed", "failed", "failed_permanent"]
    # and the permanent row now blocks any further attempt
    r = await _submit(conn, mode="auto",
                      run_agent=fake_agent(AgentResult("applied")))
    assert not r["ok"] and "already has" in r["reason"]


@pytest.mark.parametrize("reason", UNKNOWN_STATE)
async def test_an_unknown_state_send_holds_instead_of_retrying(conn, reason):
    """The agent drove a real browser and then stopped reporting -- it may
    have clicked Submit before it died. Retrying would be a double-submit
    the moment SUBMISSION_IMPLEMENTED flips, so the row must BLOCK."""
    if reason == "agent_error":
        async def runner(prompt, job_id, nonce, events, session_id=None, resume=False):
            raise RuntimeError("claude CLI not on PATH")
    else:
        runner = fake_agent(AgentResult("failed", reason))

    r = await _submit(conn, mode="auto", run_agent=runner)
    assert not r["ok"]
    row = _apps(conn)[-1]
    assert row["status"] == "held_unknown"
    assert row["failure_reason"] == reason
    # and it blocks: the queue can never re-pick this job
    again = await _submit(conn, mode="auto",
                          run_agent=fake_agent(AgentResult("applied")))
    assert not again["ok"] and "already has" in again["reason"]


@pytest.mark.parametrize("reason", ("no_result_line", "unrecognized_result:BLAH"))
async def test_the_same_reasons_stay_retryable_when_submission_is_off(conn, monkeypatch, reason):
    """With the kill switch off nothing could have been submitted, so
    re-running is free and correct. (timeout and agent_error are resumable
    stops since Task 10 -- see the S3 tests.)"""
    monkeypatch.setattr(ats_apply, "preflight", lambda: None)
    _use_live_fake(monkeypatch, AgentResult("failed", reason))

    r = await _submit(conn, mode="manual")
    assert not r["ok"]
    row = _apps(conn)[-1]
    assert row["status"] == "failed"
    assert row["failure_reason"] == reason


async def test_an_agent_crash_keeps_the_slug_and_logs_the_detail(conn):
    """failure_reason stays a queryable taxonomy slug (spec 5.3); the
    exception text lands in the event payload instead."""
    async def boom(prompt, job_id, nonce, events, session_id=None, resume=False):
        raise RuntimeError("claude CLI not on PATH")

    r = await _submit(conn, mode="auto", run_agent=boom)
    assert _apps(conn)[-1]["failure_reason"] == "agent_error"
    payload = conn.execute("SELECT payload FROM event WHERE type ="
                           " 'held_unknown'").fetchone()["payload"]
    assert "claude CLI not on PATH" in payload
    assert "claude CLI not on PATH" in r["reason"]


# -- guards ---------------------------------------------------------------

@pytest.mark.parametrize("source", ["ats", "linkedin", "naukri"])
async def test_the_kill_switch_turns_every_source_into_a_draft(conn, monkeypatch, source):
    """The switch no longer refuses the run: it decides what an approval
    means. Off, the agent fills and CONFIRMs but is told never to submit."""
    conn.execute("UPDATE job SET source = ? WHERE id = 1", (source,))
    conn.commit()
    fake = _use_live_fake(monkeypatch, AgentResult("draft_ready"))
    r = await _submit(conn, mode="auto")
    assert r["ok"] and r["status"] == "draft"
    assert "do NOT click Submit" in fake.prompts[0]


def test_submission_stays_disabled():
    assert ats_apply.SUBMISSION_IMPLEMENTED is False


async def test_blocking_status_refuses_new_attempt(conn):
    conn.execute("INSERT INTO application (job_id, resume_version, status)"
                 " VALUES (1, 'base-v1', 'submitted')")
    conn.commit()
    r = await _submit(conn, mode="manual",
                      run_agent=fake_agent(AgentResult("applied")))
    assert not r["ok"] and "already has" in r["reason"]


async def test_missing_profile_raises(conn):
    with pytest.raises(RuntimeError, match="candidate_profile"):
        await _submit(conn, mode="manual", profile=None,
                      run_agent=fake_agent(AgentResult("applied")))


async def test_unknown_job_is_reported(conn):
    r = await ats_apply.submit(conn, 99, mode="manual", brief=BRIEF,
                               profile=PROFILE,
                               run_agent=fake_agent(AgentResult("applied")))
    assert not r["ok"] and "not found" in r["reason"]


async def test_submit_refuses_when_no_resume_is_on_record(conn):
    r = await _submit(conn, mode="manual", resume_version="tailored-1-r1",
                      run_agent=fake_agent(AgentResult("draft_ready", answers={})))
    assert not r["ok"]
    assert "résumé" in r["reason"] or "resume" in r["reason"].lower()


async def test_draft_stores_the_given_resume_version(conn, tmp_path):
    (tmp_path / "r1.docx").write_text("x", encoding="utf-8")
    conn.execute("INSERT INTO resume (version, path)"
                 " VALUES ('tailored-1-r1', ?)", (str(tmp_path / "r1.docx"),))
    conn.commit()
    r = await _submit(conn, mode="manual", resume_version="tailored-1-r1",
                      run_agent=fake_agent(AgentResult("draft_ready", answers={})))
    assert r["ok"]
    assert _apps(conn)[0]["resume_version"] == "tailored-1-r1"


async def test_draft_falls_back_to_the_default_resume_version(conn):
    r = await _submit(conn, mode="manual",
                      run_agent=fake_agent(AgentResult("draft_ready", answers={})))
    assert r["ok"]
    assert _apps(conn)[0]["resume_version"] == ats_apply.RESUME_VERSION


# -- resume text ----------------------------------------------------------

async def test_prompt_resume_text_comes_from_the_docx_body(conn, tmp_path):
    """resume.content holds only the tailored summary/bullets JSON, not the
    work history an employment-history form needs -- the body must come off
    the rendered DOCX."""
    path = tmp_path / "r1.docx"
    doc = docx.Document()
    doc.add_paragraph("Jane Doe -- AI Engineer")
    doc.add_paragraph("Acme Corp, 2021-2024: shipped the thing")
    doc.save(str(path))
    conn.execute("INSERT INTO resume (version, path, content)"
                 " VALUES ('tailored-1-r1', ?, ?)",
                 (str(path), json.dumps({"summary": "Tailored summary here",
                                         "bullets": [{"text": "Bullet one",
                                                      "fact_ids": [1]}]})))
    conn.commit()
    fake = fake_agent(AgentResult("draft_ready", answers={}))
    await _submit(conn, mode="manual", resume_version="tailored-1-r1",
                  run_agent=fake)
    prompt = fake.prompts[0]
    assert "Acme Corp, 2021-2024: shipped the thing" in prompt
    assert "Tailored summary here" in prompt
    assert "Bullet one" in prompt
    # ... and the upload path is the staged clean copy, not the stored file
    assert _upload_path(prompt).name == "Jane_Doe_Resume.docx"


async def test_an_unreadable_resume_file_does_not_break_the_draft(conn):
    """'r.docx' from the fixture exists but is not a parseable .docx --
    _resume_text degrades rather than raising. (A résumé that is MISSING is
    a different case: submit() refuses, see below.)"""
    fake = fake_agent(AgentResult("draft_ready", answers={}))
    r = await _submit(conn, mode="manual", run_agent=fake)
    assert r["ok"]
    assert "== RESUME TEXT ==" in fake.prompts[0]


# -- state machine (kept) -------------------------------------------------

def test_stale_in_flight_becomes_held_unknown(conn):
    conn.execute("INSERT INTO application (job_id, resume_version, status,"
                 " started_at) VALUES (1, 'v1', 'in_flight',"
                 " datetime('now', '-30 minutes'))")
    conn.commit()
    assert ats_apply.sweep_stale_in_flight(conn, minutes=15) == 1
    assert _apps(conn)[0]["status"] == "held_unknown"


async def test_a_draft_does_not_block_a_real_submission(conn):
    await _submit(conn, mode="manual",
                  run_agent=fake_agent(AgentResult("draft_ready", answers={})))
    r = await _submit(conn, mode="auto",
                      run_agent=fake_agent(AgentResult("applied")))
    assert r["ok"]
    assert {a["status"] for a in _apps(conn)} == {"draft", "submitted"}


async def test_second_real_submission_is_refused(conn):
    await _submit(conn, mode="auto",
                  run_agent=fake_agent(AgentResult("applied")))
    r = await _submit(conn, mode="auto",
                      run_agent=fake_agent(AgentResult("applied")))
    assert not r["ok"] and "already" in r["reason"]
    assert len(_apps(conn)) == 1


# -- preconditions: a missing binary must not brick the queue --------------

@pytest.fixture
def no_live_runner(monkeypatch):
    """Belt and braces for the tests below, which are the only ones that
    reach submit() with run_agent=None on a path that could otherwise
    launch Chrome."""
    async def _never(prompt, job_id, nonce, events, session_id=None, resume=False):
        raise AssertionError("the live runner must never run in a test")
    monkeypatch.setattr(ats_apply, "_live_run_agent", _never)


@pytest.mark.parametrize("mode", ["manual", "auto"])
async def test_a_missing_binary_writes_no_application_row(conn, monkeypatch,
                                                          no_live_runner, mode):
    """`npx` off PATH means no browser launched and nothing submitted. It
    must not become held_unknown -- in auto mode that converts the whole
    queue to a permanently stuck state over a missing binary."""
    def boom():
        raise agent_mod.PreconditionError("agentic apply needs `npx` on PATH")
    monkeypatch.setattr(ats_apply, "preflight", boom)
    # the kill switch would otherwise mask the send path before preflight
    monkeypatch.setattr(ats_apply, "SUBMISSION_IMPLEMENTED", True)

    r = await _submit(conn, mode=mode)
    assert not r["ok"] and "npx" in r["reason"]
    assert conn.execute("SELECT COUNT(*) n FROM application").fetchone()["n"] == 0
    assert _event_types(conn) == []
    # ... and the refusal must carry the flag worker.apply_tick pauses on:
    # no row means QUEUE_WHERE re-admits this job, so clearing the park
    # instead would re-pick it on the very next tick, forever.
    assert r["unsupported"]


async def test_preflight_is_skipped_when_a_runner_is_injected(conn, monkeypatch):
    """Preconditions describe what _live_run_agent needs; no test on this
    machine may require the claude/npx/Chrome binaries."""
    def boom():
        raise AssertionError("preflight must not run for an injected runner")
    monkeypatch.setattr(ats_apply, "preflight", boom)
    r = await _submit(conn, mode="manual",
                      run_agent=fake_agent(AgentResult("draft_ready", answers={})))
    assert r["ok"]


async def test_a_precondition_escaping_mid_run_is_not_unknown_state(conn):
    """The backstop check inside run_session: nothing launched, so
    the send path must not hold it as 'possibly submitted'."""
    async def boom(prompt, job_id, nonce, events, session_id=None, resume=False):
        raise agent_mod.PreconditionError("Chrome not found -- set CHROME_PATH")

    r = await _submit(conn, mode="auto", run_agent=boom)
    assert not r["ok"]
    row = _apps(conn)[-1]
    assert row["status"] == "failed"            # retryable, not held_unknown
    assert row["failure_reason"] == "precondition"


# -- the sweep/late-return double-submit race ------------------------------

def _sweeping_agent(conn, result):
    """A run that outlives sweep_stale_in_flight's window:
    if the deadline's kill fails to land, the sweep flips the row to
    held_unknown while the agent is still driving the browser."""
    async def _fake(prompt, job_id, nonce, events, session_id=None, resume=False):
        conn.execute("UPDATE application SET started_at ="
                     " datetime('now', '-45 minutes') WHERE status = 'in_flight'")
        conn.commit()
        assert ats_apply.sweep_stale_in_flight(conn) == 1
        return result
    return _fake


async def test_a_late_failure_does_not_reopen_a_swept_row(conn):
    """Turning held_unknown back into 'failed' re-admits the job to the
    queue and applies a second time to a form the first run may already
    have submitted."""
    r = await _submit(conn, mode="auto",
                      run_agent=_sweeping_agent(conn, AgentResult("failed", "stuck")))
    assert not r["ok"]
    row = _apps(conn)[-1]
    assert row["status"] == "held_unknown"
    assert row["status"] in ats_apply.BLOCKING     # ... so QUEUE_WHERE skips it
    again = await _submit(conn, mode="auto",
                          run_agent=fake_agent(AgentResult("applied")))
    assert not again["ok"] and "already has" in again["reason"]


async def test_a_late_applied_still_upgrades_a_swept_row(conn):
    """The truthful upgrade must never be suppressed: the form really was
    submitted."""
    r = await _submit(conn, mode="auto",
                      run_agent=_sweeping_agent(conn, AgentResult("applied")))
    assert r["ok"] and r["status"] == "submitted"
    assert _apps(conn)[-1]["status"] == "submitted"


# -- the résumé the agent is told to upload --------------------------------

async def test_the_agent_gets_an_absolute_resume_path(conn, tmp_path, monkeypatch):
    """tailor.py stores resume.path relative to backend/, but the agent's
    cwd is a temp dir outside the repo -- a relative path there resolves to
    nothing, and uploading the tailored résumé is the whole point of the
    run. Same bug class as the relative --mcp-config path."""
    rel = Path("resume/generated/tailored-1-r1.docx")
    (tmp_path / rel).parent.mkdir(parents=True)
    (tmp_path / rel).write_text("x", encoding="utf-8")
    monkeypatch.chdir(tmp_path)                       # ... so `rel` is live
    conn.execute("INSERT INTO resume (version, path) VALUES ('tailored-1-r1', ?)",
                 (str(rel),))
    conn.commit()

    fake = fake_agent(AgentResult("draft_ready", answers={}))
    r = await _submit(conn, mode="manual", resume_version="tailored-1-r1",
                      run_agent=fake)
    assert r["ok"]
    given = _upload_path(fake.prompts[0])
    assert given.is_absolute(), given
    assert given.exists()
    assert given.parent == (tmp_path / "work" / "job1").resolve()


async def test_a_missing_resume_file_refuses_and_writes_nothing(conn):
    """Better to refuse than to start a browser session that can only fail
    at the upload step -- same spirit as the binary preflight."""
    conn.execute("INSERT INTO resume (version, path)"
                 " VALUES ('tailored-1-r1', 'no/such/resume.docx')")
    conn.commit()
    r = await _submit(conn, mode="manual", resume_version="tailored-1-r1",
                      run_agent=fake_agent(AgentResult("draft_ready", answers={})))
    assert not r["ok"]
    assert "résumé" in r["reason"] or "resume" in r["reason"].lower()
    assert _apps(conn) == []
    assert _event_types(conn) == []
    # pauses the worker rather than being re-picked every 0.1s: no row means
    # QUEUE_WHERE still admits this job (same reason as the preflight above)
    assert r["unsupported"]


# -- F1: the filename a recruiter sees --------------------------------------

def _upload_path(prompt: str) -> Path:
    line = next(l for l in prompt.splitlines() if "upload this exact file" in l)
    return Path(line.rsplit(": ", 1)[1])


async def test_the_resume_is_staged_under_the_candidates_own_name(
        conn, tmp_path, monkeypatch):
    """A recruiter opening the attachment must not read
    'tailored-1-r1.docx' and learn the résumé was machine-generated per
    job (spec 3.2 FILES). The agent is handed a clean copy instead."""
    src = tmp_path / "tailored-1-r1.docx"
    src.write_text("the rendered docx", encoding="utf-8")
    conn.execute("INSERT INTO resume (version, path)"
                 " VALUES ('tailored-1-r1', ?)", (str(src),))
    conn.commit()

    fake = fake_agent(AgentResult("draft_ready", answers={}))
    r = await _submit(conn, mode="manual", resume_version="tailored-1-r1",
                      run_agent=fake)
    assert r["ok"]
    given = _upload_path(fake.prompts[0])
    assert given.name == "Jane_Doe_Resume.docx"     # PROFILE is "Jane Doe"
    assert given.is_absolute() and given.exists()
    assert given.read_text(encoding="utf-8") == "the rendered docx"
    assert given.parent == (tmp_path / "work" / "job1").resolve()
    assert src.exists()                             # the original is untouched


async def test_a_name_with_path_characters_still_makes_a_legal_filename(
        conn, tmp_path):
    profile = CandidateProfile(candidate_name="A/B  C:D",
                               candidate_email="a@example.com",
                               candidate_phone="+91-90000-00000")
    fake = fake_agent(AgentResult("draft_ready", answers={}))
    r = await _submit(conn, mode="manual", profile=profile, run_agent=fake)
    assert r["ok"]
    assert _upload_path(fake.prompts[0]).name == "A_B_C_D_Resume.docx"


# -- F3: what a run cost survives the run ----------------------------------

async def test_the_run_cost_and_duration_are_logged(conn, caplog):
    """AgentResult carries cost_usd/duration_ms and nothing read them: after
    a live trial there would be no record of what any run cost. No DB column
    for P0 -- the log line and the transcript footer are the record."""
    fake = fake_agent(AgentResult("draft_ready", answers={},
                                  cost_usd=0.0421, duration_ms=12345))
    with caplog.at_level(logging.INFO, logger="career_agent.apply.ats"):
        await _submit(conn, mode="manual", run_agent=fake)
    assert "0.0421" in caplog.text
    assert "12345" in caplog.text
    assert "job 1" in caplog.text


# -- F4: a real submission with an empty audit trail -----------------------

async def test_an_auto_submit_with_no_answers_is_recorded_and_flagged(conn):
    """A RESULT:APPLIED with no approved CONFIRM records a REAL submission with
    answers = '{}'. The send happened and must still be recorded; the gap has
    to be visible."""
    r = await _submit(conn, mode="auto",
                      run_agent=fake_agent(AgentResult("applied")))
    assert r["ok"] and r["status"] == "submitted"
    row = _apps(conn)[-1]
    assert row["status"] == "submitted"          # never lose the send itself
    assert row["submitted_at"] is not None
    assert row["failure_reason"] == "confirm_missing"
    assert "confirm_missing" in _event_types(conn)


async def test_a_submit_with_an_approved_confirm_is_not_flagged(conn, tmp_path):
    factory = lambda: db.connect(tmp_path / "t.db")

    async def fake(prompt, job_id, nonce, events, session_id=None, resume=False):
        events.on_confirm({"fields": [{"label": "q", "value": "a"}], "files": [],
                           "account_actions": [], "memory_used": [], "notes": ""})
        return AgentResult("applied")

    await _submit(conn, mode="auto", run_agent=fake, conn_factory=factory)
    row = _apps(conn)[-1]
    assert row["status"] == "submitted"
    assert row["failure_reason"] is None
    assert json.loads(row["answers"]) == {"q": "a"}
    assert "confirm_missing" not in _event_types(conn)


def _second_job(conn):
    conn.execute("INSERT INTO job (fingerprint, source, external_id, company,"
                 " company_normalized, title, title_normalized, url)"
                 " VALUES ('fp2','ats','2','Beta','beta','AI Engineer',"
                 " 'aiengineer','https://y/apply')")
    conn.commit()
    return conn.execute("SELECT id FROM job WHERE fingerprint = 'fp2'"
                        ).fetchone()["id"]


async def test_two_agent_runs_never_overlap(conn):
    """chrome.launch_chrome _kill_port(9222)s before every launch, and
    run_session wipes the shared session dir and rewrites the shared
    .mcp-apply.json -- one port, one profile, one session dir. A dashboard
    click during an auto-worker run would taskkill the in-flight Chrome
    mid-submission."""
    job2 = _second_job(conn)
    live = 0
    peak = 0

    async def runner(prompt, job_id, nonce, events, session_id=None, resume=False):
        nonlocal live, peak
        live += 1
        peak = max(peak, live)
        await asyncio.sleep(0.02)          # the browser session
        live -= 1
        return AgentResult("draft_ready", answers={})

    await asyncio.gather(
        ats_apply.submit(conn, 1, mode="manual", brief=BRIEF, profile=PROFILE,
                         run_agent=runner),
        ats_apply.submit(conn, job2, mode="manual", brief=BRIEF, profile=PROFILE,
                         run_agent=runner))
    assert peak == 1, f"{peak} agent runs were in flight at once"
    assert {a["status"] for a in conn.execute(
        "SELECT status FROM application").fetchall()} == {"draft"}


async def test_a_queued_send_does_not_age_its_own_in_flight_row(conn):
    """The in_flight row is written inside the lock, not before it: a run
    waiting its turn would otherwise have started_at ticking toward
    sweep_stale_in_flight's window while it had not begun."""
    job2 = _second_job(conn)
    started = []

    async def runner(prompt, job_id, nonce, events, session_id=None, resume=False):
        row = conn.execute("SELECT COUNT(*) n FROM application"
                           " WHERE status = 'in_flight'").fetchone()
        started.append(row["n"])
        await asyncio.sleep(0.02)
        return AgentResult("applied", answers={"q": "a"})

    await asyncio.gather(
        ats_apply.submit(conn, 1, mode="auto", brief=BRIEF, profile=PROFILE,
                         run_agent=runner),
        ats_apply.submit(conn, job2, mode="auto", brief=BRIEF, profile=PROFILE,
                         run_agent=runner))
    assert started == [1, 1], f"in_flight rows seen per run: {started}"


async def test_two_same_job_sends_produce_one_row_and_one_refusal(conn):
    """The BLOCKING check and the in_flight INSERT are separated by the
    lock's await, so while a third run holds the lock the auto worker and a
    dashboard Send can both pass the check seeing no live row. Without a
    re-check inside the lock the loser's INSERT hits
    one_live_application_per_job and the IntegrityError escapes submit() --
    worker.apply_tick turns that into run_state 'error' and the whole apply
    queue stops."""
    job2 = _second_job(conn)
    holding = asyncio.Event()

    async def holder(prompt, job_id, nonce, events, session_id=None, resume=False):
        holding.set()
        await asyncio.sleep(0.05)
        return AgentResult("draft_ready", answers={})

    async def runner(prompt, job_id, nonce, events, session_id=None, resume=False):
        await asyncio.sleep(0.01)
        return AgentResult("applied", answers={"q": "a"})

    held = asyncio.create_task(
        ats_apply.submit(conn, job2, mode="manual", brief=BRIEF,
                         profile=PROFILE, run_agent=holder))
    await holding.wait()          # another run owns the lock: both sends queue

    results = await asyncio.gather(
        ats_apply.submit(conn, 1, mode="auto", brief=BRIEF, profile=PROFILE,
                         run_agent=runner),
        ats_apply.submit(conn, 1, mode="auto", brief=BRIEF, profile=PROFILE,
                         run_agent=runner))
    await held

    assert [r["ok"] for r in results].count(True) == 1
    loser = [r for r in results if not r["ok"]][0]
    assert "already has" in loser["reason"]
    assert [a["status"] for a in _apps(conn)] == ["submitted"]


def test_the_agent_deadline_fires_before_the_sweep_window():
    """Ordering invariant: a timing-out run must always resolve its own
    in_flight row before sweep_stale_in_flight can touch it, or the
    sweep-vs-returning-run race opens. Raise one of these and you raise
    both."""
    import inspect

    deadline_s = inspect.signature(
        agent_mod.run_session).parameters["timeout_s"].default
    sweep_s = inspect.signature(
        ats_apply.sweep_stale_in_flight).parameters["minutes"].default * 60
    assert deadline_s >= 1200, (
        "a multi-page ATS form (Workday/iCIMS/SuccessFactors -- a live"
        " SuccessFactors run needed over 10 min: snapshot -> upload -> parse ->"
        f" several screens) needs more than {deadline_s}s of healthy run")
    assert sweep_s == 30 * 60
    assert deadline_s < sweep_s, (
        f"agent deadline {deadline_s}s must fire before the {sweep_s}s sweep")


# -- F7: clearing a held_unknown (web/actions.queue_retry) -----------------
# Imported inside each test: this file must stay collectable even if the web
# package's import chain breaks on a given machine.

def _held(conn, reason="agent_error"):
    conn.execute("INSERT INTO application (job_id, resume_version, status,"
                 " failure_reason) VALUES (1, 'base-v1', 'held_unknown', ?)",
                 (reason,))
    conn.commit()


def test_a_held_application_is_not_cleared_without_confirmation(conn):
    """Held means 'the agent may already have submitted'. Releasing it on a
    plain retry click would re-admit the job and apply a second time."""
    from career_agent.web import actions

    _held(conn)
    r = actions.queue_retry(conn, 1)
    assert not r["ok"]
    assert "held" in r["message"].lower() or "confirm" in r["message"].lower()
    assert _apps(conn)[-1]["status"] == "held_unknown"


def test_confirming_it_was_not_submitted_clears_the_hold(conn):
    """The applications page says 'Held -- confirm manually, then clear it'
    and no control did: queue_retry took only 'failed', and mark_applied hit
    the one_live_application_per_job index. Raw SQL was the only exit."""
    from career_agent.web import actions

    _held(conn)
    r = actions.queue_retry(conn, 1, confirm_not_submitted=True)
    assert r["ok"]
    row = _apps(conn)[-1]
    assert row["status"] == "failed"               # retryable, not BLOCKING
    assert row["status"] not in ats_apply.BLOCKING
    assert row["failure_reason"] == "agent_error"  # why it was held is kept
    assert "hold_cleared" in _event_types(conn)


async def test_clearing_a_hold_lets_the_job_be_attempted_again(conn):
    from career_agent.web import actions

    _held(conn)
    actions.queue_retry(conn, 1, confirm_not_submitted=True)
    r = await _submit(conn, mode="manual",
                      run_agent=fake_agent(AgentResult("draft_ready", answers={})))
    assert r["ok"], r


def test_confirmation_does_nothing_for_any_other_status(conn):
    """The flag is an escape hatch for held_unknown only -- it must not turn
    a 'submitted' row back into a retry."""
    from career_agent.web import actions

    conn.execute("INSERT INTO application (job_id, resume_version, status)"
                 " VALUES (1, 'base-v1', 'submitted')")
    conn.commit()
    r = actions.queue_retry(conn, 1, confirm_not_submitted=True)
    assert not r["ok"]
    assert _apps(conn)[-1]["status"] == "submitted"


def test_a_failed_application_still_retries_without_the_flag(conn):
    from career_agent.web import actions

    conn.execute("INSERT INTO application (job_id, resume_version, status)"
                 " VALUES (1, 'base-v1', 'failed')")
    conn.commit()
    assert actions.queue_retry(conn, 1)["ok"]


def test_the_human_can_say_it_WAS_submitted(conn):
    """The other exit the page promises: confirm on the employer's site that
    it went through. mark_applied promotes the held row rather than hitting
    the one_live_application_per_job index."""
    from career_agent import store as store_mod

    _held(conn)
    conn.execute("UPDATE application SET answers = '{\"q\": \"a\"}'")
    conn.commit()
    app_id = store_mod.mark_applied(conn, 1, "2026-09-01")
    rows = _apps(conn)
    assert len(rows) == 1 and rows[0]["id"] == app_id
    assert rows[0]["status"] == "submitted"
    assert rows[0]["submitted_at"] == "2026-09-01"
    assert json.loads(rows[0]["answers"]) == {"q": "a"}, (
        "unlike a draft, a held row's answers are what the agent actually"
        " typed into the live form -- the only audit trail there is")


# -- F8: submit() is what keeps the two sides of the contract together -----

async def test_the_runner_is_handed_the_prompts_own_nonce(conn):
    """build_prompt stamps it and parse_result demands it -- submit() is the
    only place both are chosen, so they cannot drift apart at runtime."""
    fake = fake_agent(AgentResult("draft_ready", answers={}))
    await _submit(conn, mode="manual", run_agent=fake)
    nonce = fake.nonces[0]
    assert nonce
    assert f"RESULT:{nonce}:APPLIED" in fake.prompts[0]
    assert "RESULT:" not in fake.prompts[0].replace(f"RESULT:{nonce}:", "")


async def test_every_run_gets_a_fresh_nonce(conn):
    fake = fake_agent(AgentResult("failed", "stuck"))
    await _submit(conn, mode="manual", run_agent=fake)
    await _submit(conn, mode="manual", run_agent=fake)
    assert fake.nonces[0] != fake.nonces[1]


# -- the sweep race on the captcha / needs_answer delete paths -------------

@pytest.mark.parametrize("result,key", [
    (AgentResult("captcha"), "held"),
    (AgentResult("needs_answer", "PMP?"), "needs_answer"),
])
async def test_a_late_hold_does_not_delete_a_swept_row(conn, result, key):
    """Same race the failure path already closed: a run can outlive
    sweep_stale_in_flight and come back to find its own row held_unknown.
    Deleting it would re-admit a job the first run may already have
    submitted."""
    r = await _submit(conn, mode="auto",
                      run_agent=_sweeping_agent(conn, result))
    assert not r["ok"] and r[key]
    row = _apps(conn)[-1]
    assert row["status"] == "held_unknown"
    again = await _submit(conn, mode="auto",
                          run_agent=fake_agent(AgentResult("applied")))
    assert not again["ok"] and "already has" in again["reason"]


async def test_two_jobs_do_not_share_one_staged_resume(conn, tmp_path):
    """Staging happens before _agent_lock() is taken, and the basename is
    the same for every job -- one shared path would let a second submit()
    overwrite the copy the first run is about to upload, and send the wrong
    résumé to a real employer."""
    job2 = _second_job(conn)
    for version, body, job in (("tailored-1-r1", "resume for job 1", 1),
                               ("tailored-2-r1", "resume for job 2", job2)):
        src = tmp_path / f"{version}.docx"
        src.write_text(body, encoding="utf-8")
        conn.execute("INSERT INTO resume (version, path) VALUES (?, ?)",
                     (version, str(src)))
    conn.commit()

    staged = {}

    async def runner(prompt, job_id, nonce, events, session_id=None, resume=False):
        staged[job_id] = _upload_path(prompt)
        await asyncio.sleep(0.01)
        return AgentResult("draft_ready", answers={})

    await asyncio.gather(
        ats_apply.submit(conn, 1, mode="manual", brief=BRIEF, profile=PROFILE,
                         resume_version="tailored-1-r1", run_agent=runner),
        ats_apply.submit(conn, job2, mode="manual", brief=BRIEF, profile=PROFILE,
                         resume_version="tailored-2-r1", run_agent=runner))

    assert staged[1] != staged[job2]
    assert staged[1].name == staged[job2].name == "Jane_Doe_Resume.docx"
    assert staged[1].read_text(encoding="utf-8") == "resume for job 1"
    assert staged[job2].read_text(encoding="utf-8") == "resume for job 2"


# -- live-safety FIX 1: account_required is a one-time human fix -----------

def test_account_required_is_retryable_not_permanent_nor_unknown():
    assert "account_required" not in ats_apply.PERMANENT_REASONS
    assert not ats_apply.is_unknown_state("account_required")
    assert ats_apply.classify_failure("account_required", 0) == "failed"


@pytest.mark.parametrize("mode", ["manual", "auto"])
async def test_account_required_records_a_retryable_failure(conn, mode):
    """The human creates the account (or gives the consent) in the agent's
    Chrome profile once, then the same job is re-attempted."""
    r = await _submit(conn, mode=mode,
                      run_agent=fake_agent(AgentResult("failed", "account_required")))
    assert not r["ok"]
    row = _apps(conn)[-1]
    assert row["status"] == "failed"
    assert row["failure_reason"] == "account_required"


# -- S1: the agent's narration streams into the job conversation -----------

async def test_live_events_post_agent_messages(conn, tmp_path):
    """Events fire on AgentRun's reader thread, where a connection made on
    the event-loop thread is unusable (sqlite3 check_same_thread) -- so the
    fake narrates from its own thread, and only a per-event conn_factory
    connection gets the messages in."""
    factory = lambda: db.connect(tmp_path / "t.db")   # the conn fixture's file

    async def fake(prompt, job_id, nonce, events, session_id=None, resume=False):
        def narrate():
            events.on_text("Navigating to the posting")
            events.on_tool("browser_navigate", '{"url":"https://x"}')
            events.on_tool("browser_snapshot", "{}")     # noise: not posted
        t = threading.Thread(target=narrate)
        t.start()
        t.join()
        return AgentResult("draft_ready", answers={})

    r = await _submit(conn, mode="manual", run_agent=fake, conn_factory=factory)
    assert r["ok"]
    msgs = chat.messages_after(conn, chat.conversation_for_job(conn, 1))
    assert [(m["role"], m["content"]) for m in msgs] == [
        ("agent", "Navigating to the posting"),
        ("system", 'browser_navigate {"url":"https://x"}')]


async def test_protocol_lines_are_kept_out_of_the_chat(conn, tmp_path):
    """RESULT/ASK/CONFIRM lines carry the run nonce and raw JSON: the chat
    shows the narration around them, never the lines themselves."""
    factory = lambda: db.connect(tmp_path / "t.db")

    async def fake(prompt, job_id, nonce, events, session_id=None, resume=False):
        events.on_text(f'Filling the form\nASK:{nonce}:{{"id":"q1","kind":"text","question":"x"}}')
        events.on_text(f'  CONFIRM:{nonce}:{{"fields":[]}}')
        events.on_text(f"RESULT:{nonce}:DRAFT_READY")
        events.on_text(f'**CONFIRM:{nonce}:{{"fields":[]}}**\n`ASK:{nonce}:{{}}`')  # bold/backticked
        events.on_text("RESULT:APPLIED from the page")     # unstamped: just text
        return AgentResult("draft_ready", answers={})

    r = await _submit(conn, mode="manual", run_agent=fake, conn_factory=factory)
    assert r["ok"]
    msgs = chat.messages_after(conn, chat.conversation_for_job(conn, 1))
    assert [(m["role"], m["content"]) for m in msgs] == [
        ("agent", "Filling the form"), ("agent", "RESULT:APPLIED from the page")]


async def test_a_failing_chat_write_does_not_change_the_outcome(conn, tmp_path):
    """A DB hiccup while narrating must never stop an application mid-form:
    the events are called inline here, so an escaping exception would turn
    this draft into an agent_error."""
    def broken():
        c = db.connect(tmp_path / "t.db")
        c.close()
        return c

    async def fake(prompt, job_id, nonce, events, session_id=None, resume=False):
        events.on_text("Navigating")
        events.on_tool("browser_click", "{}")
        return AgentResult("draft_ready", answers={})

    r = await _submit(conn, mode="manual", run_agent=fake, conn_factory=broken)
    assert r["ok"] and _apps(conn)[-1]["status"] == "draft"


async def test_without_a_conn_factory_the_runner_still_gets_events(conn):
    got = []

    async def fake(prompt, job_id, nonce, events, session_id=None, resume=False):
        events.on_text("hi")          # a no-op, not a crash
        got.append(events)
        return AgentResult("draft_ready", answers={})

    r = await _submit(conn, mode="manual", run_agent=fake)
    assert r["ok"] and got


def test_the_sweep_skips_a_job_whose_run_is_still_live(conn, monkeypatch):
    """The work clock pauses while a human reads a card, so a healthy run can
    outlive the 30-min window -- and web/app.py sweeps on every request."""
    conn.execute("INSERT INTO application (job_id, resume_version, status,"
                 " started_at) VALUES (1, 'v1', 'in_flight',"
                 " datetime('now', '-45 minutes'))")
    conn.commit()
    monkeypatch.setattr(agent_mod, "RUNS", {1: object()})
    assert ats_apply.sweep_stale_in_flight(conn) == 0
    assert _apps(conn)[0]["status"] == "in_flight"
    monkeypatch.setattr(agent_mod, "RUNS", {})
    assert ats_apply.sweep_stale_in_flight(conn) == 1

# -- S2: the two-way relay (Task 7) ----------------------------------------

class FakeRun:
    """Stands in for agent.RUNS[job_id]: same send() contract as AgentRun, and
    the same `events` (whose prompt_baseline binds cards to this run)."""
    def __init__(self, nonce, events=None):
        from career_agent.apply.runner import RunEvents
        self.events = events or RunEvents()
        self.nonce, self.sent, self.secrets = nonce, [], set()
        self.waiting, self.done = threading.Event(), threading.Event()

    def send(self, text):
        if self.done.is_set() or not self.waiting.is_set():
            return False
        self.waiting.clear()
        self.sent.append(text)
        return True

    def kill(self):
        self.killed = True
        self.done.set()


@pytest.fixture
def runs(monkeypatch):
    d = {}
    monkeypatch.setattr(agent_mod, "RUNS", d)
    return d


@pytest.fixture
def factory(tmp_path):
    return lambda: db.connect(tmp_path / "t.db")


def _confirm(**fields):
    return {"fields": [{"label": k, "value": v} for k, v in fields.items()],
            "files": [], "account_actions": [], "memory_used": [], "notes": ""}


def _waiting_on(run, events, kind, payload):
    """What the reader thread does at an ASK/CONFIRM turn end."""
    run.waiting.set()
    (events.on_confirm if kind == "confirm" else events.on_ask)(payload)


async def test_ask_opens_a_prompt_and_answer_reaches_the_run(conn, runs, factory):
    from career_agent.web import actions

    async def fake(prompt, jid, nonce, events, session_id=None, resume=False):
        run = runs[jid] = FakeRun(nonce, events)
        _waiting_on(run, events, "ask", {
            "id": "q1", "kind": "choice", "question": "Notice?", "options": ["30", "60"],
            "why": "", "memory_key": "notice_period", "default": None, "sensitive": False})
        pid = chat.open_prompt_for_job(factory(), jid)["id"]
        r = actions.answer_prompt(factory(), pid, {"answer": "30", "remember": True}, factory)
        assert r["ok"], r
        assert run.sent[0].startswith(f"ANSWER:{nonce}:")
        assert json.loads(run.sent[0].split(":", 2)[2]) == {"id": "q1", "answer": "30", "remember": True}
        return AgentResult("draft_ready")

    r = await _submit(conn, mode="manual", run_agent=fake, conn_factory=factory)
    assert r["ok"] and r["status"] == "draft"
    assert chat.open_prompt_for_job(conn, 1) is None
    msgs = chat.messages_after(conn, chat.conversation_for_job(conn, 1))
    assert ("user", "Notice? → 30") in [(m["role"], m["content"]) for m in msgs]


async def test_confirm_in_auto_mode_is_auto_approved_and_recorded(conn, runs, factory):
    async def fake(prompt, jid, nonce, events, session_id=None, resume=False):
        run = runs[jid] = FakeRun(nonce, events)
        _waiting_on(run, events, "confirm", _confirm(Name="Asha"))
        assert run.sent == [f'DECISION:{nonce}:{{"decision": "approve"}}']
        return AgentResult("applied")

    r = await _submit(conn, mode="auto", run_agent=fake, conn_factory=factory)
    assert r["ok"] and r["status"] == "submitted"
    row = _apps(conn)[-1]
    assert row["status"] == "submitted" and row["failure_reason"] is None
    assert json.loads(row["answers"]) == {"Name": "Asha"}
    p = conn.execute("SELECT * FROM agent_prompt WHERE kind = 'confirm'").fetchone()
    assert p["status"] == "answered" and json.loads(p["answer"]) == {"decision": "approve"}


async def test_auto_approval_of_a_confirm_in_the_result_turn_still_records(conn, runs, factory):
    """The pre-approved agent CONFIRMs and submits without ending its turn:
    the run is not waiting, nothing is sent, and the CONFIRM still counts."""
    async def fake(prompt, jid, nonce, events, session_id=None, resume=False):
        run = runs[jid] = FakeRun(nonce, events)
        events.on_confirm(_confirm(Name="Asha"))          # not waiting
        assert run.sent == []
        return AgentResult("applied")

    await _submit(conn, mode="auto", run_agent=fake, conn_factory=factory)
    assert json.loads(_apps(conn)[-1]["answers"]) == {"Name": "Asha"}


async def test_confirm_in_manual_mode_waits_for_a_decision(conn, runs, factory):
    from career_agent.web import actions

    async def fake(prompt, jid, nonce, events, session_id=None, resume=False):
        run = runs[jid] = FakeRun(nonce, events)
        _waiting_on(run, events, "confirm", _confirm(Name="Asha", Phone="+1"))
        await asyncio.sleep(0.05)
        first = chat.open_prompt_for_job(factory(), jid)
        assert first is not None and run.sent == [], "manual never auto-approves"
        # From another thread, as the HTTP route does -- its connection too.
        r = await asyncio.to_thread(lambda: actions.answer_prompt(
            factory(), first["id"], {"decision": "change", "changes": {"Phone": "+91"}}, factory))
        assert r["ok"], r
        assert run.sent[-1].startswith(f"DECISION:{nonce}:")
        assert json.loads(run.sent[-1].split(":", 2)[2]) == {
            "decision": "change", "changes": {"Phone": "+91"}}
        _waiting_on(run, events, "confirm", _confirm(Name="Asha", Phone="+91"))
        second = chat.open_prompt_for_job(factory(), jid)
        assert actions.answer_prompt(factory(), second["id"], {"decision": "approve"}, factory)["ok"]
        return AgentResult("applied")

    r = await _submit(conn, mode="manual", run_agent=fake, conn_factory=factory)
    assert r["ok"]
    assert json.loads(_apps(conn)[-1]["answers"]) == {"Name": "Asha", "Phone": "+91"}


async def test_a_manual_run_with_no_approved_confirm_records_no_answers(conn, runs, factory):
    async def fake(prompt, jid, nonce, events, session_id=None, resume=False):
        run = runs[jid] = FakeRun(nonce, events)
        _waiting_on(run, events, "confirm", _confirm(Name="Asha"))   # never approved
        return AgentResult("draft_ready")

    r = await _submit(conn, mode="manual", run_agent=fake, conn_factory=factory)
    assert r["ok"] and r["status"] == "draft"
    assert json.loads(_apps(conn)[-1]["answers"]) == {}
    # a CONFIRM was shown but never approved; confirm_missing means none at all
    assert "draft_without_decision" in _event_types(conn)
    assert chat.open_prompt_for_job(conn, 1) is None, "a finished run leaves no open card"


def _open(conn, kind, payload):
    return chat.open_prompt(conn, 1, kind, payload)


def test_answer_prompt_validates_by_kind_and_refuses_closed(conn, runs):
    from career_agent.web import actions
    run = runs[1] = FakeRun("n0nce")
    run.waiting.set()
    pid = _open(conn, "choice", {"id": "q", "question": "Notice?", "options": ["30", "60"]})
    r = actions.answer_prompt(conn, pid, {"answer": "90"})
    assert not r["ok"] and r["code"] == 422 and run.sent == []
    assert actions.answer_prompt(conn, pid, {"answer": "60"})["ok"]
    run.waiting.set()
    again = actions.answer_prompt(conn, pid, {"answer": "60"})
    assert again == {"ok": False, "code": 409, "message": "That question is no longer open"}

    for kind, payload, bad, good in [
            ("text", {"id": "t", "question": "Why?"}, {"answer": "  "}, {"answer": "because"}),
            ("approve", {"id": "a", "question": "OK?"}, {"answer": "maybe"}, {"answer": "reject"}),
            ("confirm", _confirm(Name="A"), {"decision": "change", "changes": {}}, {"decision": "cancel"})]:
        run.waiting.set()
        pid = _open(conn, kind, payload)
        assert actions.answer_prompt(conn, pid, bad)["code"] == 422, kind
        assert actions.answer_prompt(conn, pid, good)["ok"], kind
    assert actions.answer_prompt(conn, 999, {})["code"] == 404


# -- Task 15: approve_account / need_password (backend fills and submits) ----

@pytest.fixture
def key(monkeypatch):
    from cryptography.fernet import Fernet
    monkeypatch.setenv("CREDENTIAL_KEY", Fernet.generate_key().decode())


@pytest.fixture
def browser_on(monkeypatch):
    """Point secret_fill at fake CDP pages -- never a real Chrome -- with no
    post-submit wait (an SPA fake never navigates)."""
    from test_secret_fill import connect_to
    from career_agent.apply import secret_fill
    monkeypatch.setattr(secret_fill, "SUBMIT_WAIT_S", 0)

    def on(*pages):
        monkeypatch.setattr(secret_fill, "_live_connect", connect_to(*pages))
        return pages
    return on


def _chat_texts(conn):
    return [m["content"] for m in chat.messages_after(conn, chat.conversation_for_job(conn, 1))]


def _db_dump(conn):
    return "\n".join(conn.iterdump())


def _sent_body(line):
    return json.loads(line.split(":", 2)[2])


def _in_flight(conn):
    """The live run's row: a refused answer reopens its card only while it exists."""
    conn.execute("INSERT INTO application (job_id, resume_version, status, started_at)"
                 " VALUES (1, 'v1', 'in_flight', datetime('now'))")
    conn.commit()


_ACCT = {"id": "acct", "kind": "approve_account", "question": "Create an account?",
         "domain": "careers.ses.com", "email": "asha@example.com",
         "login_url": "https://careers.ses.com/join"}


def _waiting_run(runs):
    run = runs[1] = FakeRun("n")
    run.waiting.set()
    return run


def test_approve_account_fills_submits_then_stores_and_never_sends_the_password(conn, runs, key, browser_on):
    from test_secret_fill import Field, Page
    from career_agent import credentials
    from career_agent.web import actions
    [page] = browser_on(Page("https://careers.ses.com/join", [Field(), Field()]))
    run = _waiting_run(runs)
    pid = _open(conn, "approve_account", _ACCT)
    assert actions.answer_prompt(conn, pid, {"answer": "approve"})["ok"]

    pw = page.main.fields[0].fills[0]
    assert len(pw) == 20 and page.main.fields[1].fills == [pw]  # confirm field too
    assert page.main.submits == ["requestSubmit"]               # the backend's submit is Create
    assert [_sent_body(s) for s in run.sent] == [{"id": "acct", "answer": "approve", "submitted": True}]
    assert pw not in "".join(run.sent)
    cred = credentials.get(conn, "careers.ses.com")
    assert cred["password"] == pw and cred["created_by"] == "agent"
    assert cred["login_url"] == "https://careers.ses.com/join"
    assert pw in run.secrets                                     # defence in depth
    assert any(t.startswith("Saved login for careers.ses.com") for t in _chat_texts(conn))
    assert pw not in _db_dump(conn)
    answer = json.loads(conn.execute("SELECT answer FROM agent_prompt WHERE id = ?",
                                     (pid,)).fetchone()[0])
    assert answer == {"id": "acct", "answer": "approve"}


def test_approve_account_on_a_foreign_real_page_submits_and_stores_nothing(conn, runs, key, browser_on):
    from test_secret_fill import Field, Page
    from career_agent import credentials
    from career_agent.web import actions
    [page] = browser_on(Page("https://evil.com/join?token=s3cr3t#frag", [Field()]))
    _in_flight(conn)
    run = _waiting_run(runs)
    pid = _open(conn, "approve_account", _ACCT)
    r = actions.answer_prompt(conn, pid, {"answer": "approve"})
    assert r["code"] == 409 and "https://evil.com/join" in r["message"]
    assert "s3cr3t" not in r["message"]                          # scheme+host+path only
    assert run.sent == [] and credentials.list_(conn) == []
    assert page.main.fields[0].fills == [] and page.main.submits == []
    assert chat.open_prompt_for_job(conn, 1)["id"] == pid


def test_approve_account_refuses_a_login_url_off_the_domain(conn, runs, key, browser_on):
    from test_secret_fill import Field, Page
    from career_agent import credentials
    from career_agent.web import actions
    [page] = browser_on(Page("https://careers.ses.com/join", [Field()]))
    run = _waiting_run(runs)
    pid = _open(conn, "approve_account", dict(_ACCT, login_url="https://evil.com/join"))
    r = actions.answer_prompt(conn, pid, {"answer": "approve"})
    assert r["code"] == 409 and run.sent == [] and credentials.list_(conn) == []
    assert page.main.fields[0].fills == []
    assert chat.open_prompt_for_job(conn, 1)["id"] == pid


@pytest.mark.parametrize("domain", ["co.in", "myworkdayjobs.com", "wd3.myworkdayjobs.com",
                                    "vercel.app", "localhost", "com"])
def test_approve_account_refuses_shared_suffix_and_bare_domains(conn, runs, key, domain):
    from career_agent import credentials
    from career_agent.web import actions
    run = _waiting_run(runs)
    pid = _open(conn, "approve_account",
                dict(_ACCT, domain=domain, login_url=f"https://{domain}/join"))
    r = actions.answer_prompt(conn, pid, {"answer": "approve"})
    assert r["code"] == 422 and run.sent == [] and credentials.list_(conn) == []


def test_approve_account_with_an_existing_login_tells_the_agent_it_exists(conn, runs, key, browser_on):
    from test_secret_fill import Field, Page
    from career_agent import credentials
    from career_agent.web import actions
    [page] = browser_on(Page("https://careers.ses.com/join", [Field()]))
    credentials.put(conn, "careers.ses.com", "", "a@x.com", "agent-made-pw", "agent")
    run = _waiting_run(runs)
    pid = _open(conn, "approve_account", _ACCT)
    r = actions.answer_prompt(conn, pid, {"answer": "approve"})
    assert r["code"] == 409
    assert r["message"] == ("a login already exists for careers.ses.com; use it, or delete "
                            "it in Logins first")
    assert [_sent_body(s) for s in run.sent] == [{"id": "acct", "answer": "exists"}]
    assert page.main.fields[0].fills == []
    assert credentials.get(conn, "careers.ses.com")["password"] == "agent-made-pw"


def test_reject_account_sends_reject_and_stores_nothing(conn, runs, key):
    from career_agent import credentials
    from career_agent.web import actions
    run = _waiting_run(runs)
    pid = _open(conn, "approve_account", _ACCT)
    assert actions.answer_prompt(conn, pid, {"answer": "reject"})["ok"]
    assert _sent_body(run.sent[0]) == {"id": "acct", "answer": "reject"}
    assert credentials.list_(conn) == []
    assert actions.answer_prompt(conn, _open(conn, "approve_account", _ACCT),
                                 {"answer": "maybe"})["code"] == 422


def test_an_exception_after_the_approve_claim_reopens_the_card_with_503(conn, runs, key, monkeypatch):
    from career_agent import credentials
    from career_agent.apply import secret_fill
    from career_agent.web import actions

    def boom(*a, **kw):
        raise RuntimeError("cdp went away")
    monkeypatch.setattr(secret_fill, "fill_and_submit", boom)
    _in_flight(conn)
    run = _waiting_run(runs)
    pid = _open(conn, "approve_account", _ACCT)
    r = actions.answer_prompt(conn, pid, {"answer": "approve"})
    assert r["code"] == 503 and run.sent == [] and credentials.list_(conn) == []
    assert chat.open_prompt_for_job(conn, 1)["id"] == pid


def test_a_store_failure_after_submit_clears_the_field_and_reopens_the_card(conn, runs, key, browser_on, monkeypatch):
    from test_secret_fill import Field, Page
    from career_agent import credentials
    from career_agent.web import actions
    [page] = browser_on(Page("https://careers.ses.com/join", [Field()], navigates=False))

    def boom(*a, **kw):
        raise RuntimeError("database is locked")
    monkeypatch.setattr(credentials, "put", boom)
    _in_flight(conn)
    run = _waiting_run(runs)
    pid = _open(conn, "approve_account", _ACCT)
    r = actions.answer_prompt(conn, pid, {"answer": "approve"})
    assert r["code"] == 503 and run.sent == []
    assert page.main.submits == ["requestSubmit"] and page.main.fields[0].value == ""
    assert chat.open_prompt_for_job(conn, 1)["id"] == pid


def _need(domain="careers.ses.com", url="https://careers.ses.com/login"):
    return {"id": "np", "kind": "need_password", "question": f"Password for {domain}",
            "domain": domain, "url": url}


def test_need_password_is_refused_on_the_human_answer_path(conn, runs, key):
    from career_agent.web import actions
    run = _waiting_run(runs)
    pid = _open(conn, "need_password", _need())
    r = actions.answer_prompt(conn, pid, {"answer": "x"})
    assert r["code"] == 422 and run.sent == []
    assert chat.open_prompt_for_job(conn, 1)["id"] == pid


def test_need_password_fills_and_submits_the_real_matching_page(conn, runs, key, browser_on):
    from test_secret_fill import Field, Page
    from career_agent import credentials
    from career_agent.web import actions
    pw = "Stored!Pass_4321abcd"
    credentials.put(conn, "careers.ses.com", "", "a@x.com", pw, "agent")
    [page] = browser_on(Page("https://jobs.careers.ses.com/login?next=%2Fhome", [Field()]))
    run = _waiting_run(runs)
    pid = _open(conn, "need_password", _need())
    assert actions.answer_prompt(conn, pid, {}, auto=True)["ok"]
    assert page.main.fields[0].fills == [pw] and page.main.submits == ["requestSubmit"]
    assert [_sent_body(s) for s in run.sent] == [{"id": "np", "submitted": True}]
    assert pw not in "".join(run.sent) and pw not in _db_dump(conn)
    assert ("Submitted the saved login for careers.ses.com on "
            "https://jobs.careers.ses.com/login") in _chat_texts(conn)


def test_c1_a_claimed_url_never_releases_the_login_on_a_foreign_real_page(conn, runs, key, browser_on):
    """C1 regression: an injected agent claims it is on linkedin while the
    real browser page is evil.com -- nothing is filled, the answer is none."""
    from test_secret_fill import Field, Page
    from career_agent import credentials
    from career_agent.web import actions
    pw = "Users!Own_Pass_9876"
    credentials.put(conn, "linkedin.com", "", "me@x.com", pw, "user")
    [page] = browser_on(Page("https://evil.com/login", [Field()]))
    run = _waiting_run(runs)
    pid = _open(conn, "need_password", _need("linkedin.com", "https://www.linkedin.com/login"))
    assert actions.answer_prompt(conn, pid, {}, auto=True)["ok"]
    assert page.main.fields[0].fills == [] and page.main.submits == []
    assert [_sent_body(s) for s in run.sent] == [{"id": "np", "answer": "none"}]
    assert pw not in "".join(run.sent) and pw not in _db_dump(conn)
    assert any("https://evil.com/login" in t for t in _chat_texts(conn))


def test_need_password_never_uses_a_stored_login_for_a_bare_suffix(conn, runs, key, browser_on):
    from test_secret_fill import Field, Page
    from career_agent import credentials
    from career_agent.web import actions
    credentials.put(conn, "co.in", "", "me@x.com", "Users!Own_Pass_9876", "user")
    [page] = browser_on(Page("https://anything.co.in/login", [Field()]))
    run = _waiting_run(runs)
    pid = _open(conn, "need_password", _need("co.in", "https://anything.co.in/login"))
    assert actions.answer_prompt(conn, pid, {}, auto=True)["ok"]
    assert page.main.fields[0].fills == []
    assert [_sent_body(s) for s in run.sent] == [{"id": "np", "answer": "none"}]


def test_need_password_with_no_saved_login_answers_none(conn, runs, key):
    from career_agent.web import actions
    run = _waiting_run(runs)
    pid = _open(conn, "need_password", _need())
    assert actions.answer_prompt(conn, pid, {}, auto=True)["ok"]
    assert _sent_body(run.sent[0]) == {"id": "np", "answer": "none"}


async def test_need_password_is_submitted_on_ask_and_logins_reach_the_prompt(conn, runs, factory, key, browser_on):
    from test_secret_fill import Field, Page
    from career_agent import credentials
    pw = "Auto!Pass_1234wxyz"
    credentials.put(conn, "careers.ses.com", "", "a@x.com", pw, "agent")
    [page] = browser_on(Page("https://careers.ses.com/login", [Field()]))
    seen = {}

    async def fake(prompt, jid, nonce, events, session_id=None, resume=False):
        run = runs[jid] = FakeRun(nonce, events)
        seen["prompt"] = prompt
        # on_ask runs on AgentRun's reader thread in production, never the loop
        await asyncio.to_thread(_waiting_on, run, events, "ask", _need())
        seen["sent"] = list(run.sent)
        return AgentResult("draft_ready")

    await _submit(conn, mode="manual", run_agent=fake, conn_factory=factory)
    assert [_sent_body(s) for s in seen["sent"]] == [{"id": "np", "submitted": True}]
    assert page.main.fields[0].fills == [pw]
    assert "careers.ses.com (sign in as a@x.com)" in seen["prompt"]
    assert pw not in seen["prompt"] and pw not in _db_dump(conn)


async def test_a_failed_auto_need_password_answers_none_at_once(conn, runs, factory, key, monkeypatch):
    from career_agent import credentials
    from career_agent.web import actions
    credentials.put(conn, "careers.ses.com", "", "a@x.com", "pw-Fail-123!", "agent")
    seen = {}

    def boom(*a, **kw):
        raise RuntimeError("answer_prompt blew up")

    async def fake(prompt, jid, nonce, events, session_id=None, resume=False):
        run = runs[jid] = FakeRun(nonce, events)
        _waiting_on(run, events, "ask", _need())
        seen["sent"] = list(run.sent)
        return AgentResult("draft_ready")

    monkeypatch.setattr(actions, "answer_prompt", boom)
    await _submit(conn, mode="manual", run_agent=fake, conn_factory=factory)
    assert [_sent_body(s) for s in seen["sent"]] == [{"id": "np", "answer": "none"}]
    assert any("careers.ses.com" in t and "none" in t for t in _chat_texts(conn))


def test_an_approve_account_card_carries_the_real_page_urls(conn, factory, browser_on):
    from test_secret_fill import Page
    browser_on(Page("https://careers.ses.com/join?invite=abc"), Page("https://evil.com/x"))
    ats_apply._chat_events(factory, 1, "nn").on_ask(dict(_ACCT, page_urls=["https://lie.example/"]))
    payload = json.loads(chat.open_prompt_for_job(conn, 1)["payload"])
    assert payload["page_urls"] == ["https://careers.ses.com/join", "https://evil.com/x"]


def test_a_tool_call_echoing_a_sent_password_is_redacted_in_chat_and_transcript(conn, factory):
    from test_runner import FakePopen, _assistant, _result
    from career_agent.apply.runner import AgentRun
    pw = "Echo!Pass_9876wxyz"
    fake = FakePopen()
    run = AgentRun(["x"], ".", {}, "nn", ats_apply._chat_events(factory, 1, "nn"),
                   popen=lambda *a, **kw: fake)
    run.secrets.add(pw)
    run.start("go")
    fake.emit(_assistant(f"Signing in with {pw}", ("mcp__playwright__browser_fill_form",
              {"fields": [{"name": "Account key", "value": pw}]})))
    fake.emit(_assistant("RESULT:nn:APPLIED")); fake.emit(_result()); fake.close()
    assert run.wait(5)
    texts = _chat_texts(conn)
    assert any("browser_fill_form" in t for t in texts)
    assert not any(pw in t for t in texts) and pw not in run.transcript


def test_a_sensitive_answer_is_hidden_in_the_chat(conn, runs):
    from career_agent.web import actions
    run = runs[1] = FakeRun("n")
    run.waiting.set()
    pid = _open(conn, "text", {"id": "s", "question": "Expected CTC?", "sensitive": True})
    assert actions.answer_prompt(conn, pid, {"answer": "42 LPA"})["ok"]
    texts = [m["content"] for m in chat.messages_after(conn, chat.conversation_for_job(conn, 1))]
    assert "Expected CTC? → (hidden)" in texts and not any("42 LPA" in t for t in texts)
    assert "42 LPA" in run.sent[0]                     # the agent still gets it


@pytest.mark.parametrize("state", ["no_run", "not_waiting", "done"])
def test_a_refused_send_leaves_the_prompt_open(conn, runs, state):
    from career_agent.web import actions
    if state != "no_run":
        run = runs[1] = FakeRun("n")
        if state == "done":
            run.waiting.set(); run.done.set()
    pid = _open(conn, "confirm", _confirm(Name="A"))
    r = actions.answer_prompt(conn, pid, {"decision": "approve"})
    assert r == {"ok": False, "code": 409, "message": "No live agent run for this job"}
    assert chat.open_prompt_for_job(conn, 1)["id"] == pid


# -- Task 13: answer_prompt remembers choice/text answers -------------------

def test_answer_prompt_remembers_a_choice_answer_by_default(conn, runs):
    from career_agent.web import actions
    run = runs[1] = FakeRun("n")
    run.waiting.set()
    pid = _open(conn, "choice", {"id": "q", "kind": "choice", "question": "Notice period?",
                                 "options": ["30", "60"], "memory_key": "notice_period",
                                 "sensitive": False})
    assert actions.answer_prompt(conn, pid, {"answer": "30"})["ok"]
    assert store.qa_lookup(conn, "Notice period?")["answer"] == "30"
    keyed = store.qa_by_key(conn, "notice_period")
    assert keyed is not None and keyed["answer"] == "30"
    assert keyed["source_job_id"] == 1


def test_answer_prompt_remember_false_writes_nothing(conn, runs):
    from career_agent.web import actions
    run = runs[1] = FakeRun("n")
    run.waiting.set()
    pid = _open(conn, "text", {"id": "q", "kind": "text", "question": "Why?",
                               "memory_key": None, "sensitive": False})
    assert actions.answer_prompt(conn, pid, {"answer": "because", "remember": False})["ok"]
    assert conn.execute("SELECT COUNT(*) n FROM qa_bank").fetchone()["n"] == 0


def test_answer_prompt_sensitive_card_is_never_remembered(conn, runs):
    from career_agent.web import actions
    run = runs[1] = FakeRun("n")
    run.waiting.set()
    pid = _open(conn, "text", {"id": "q", "kind": "text", "question": "SSN?",
                               "memory_key": "ssn", "sensitive": True})
    assert actions.answer_prompt(conn, pid, {"answer": "123-45-6789"})["ok"]
    assert conn.execute("SELECT COUNT(*) n FROM qa_bank").fetchone()["n"] == 0


class RefusingRun(FakeRun):
    """A run that always refuses send() -- same as a run that ended between
    the waiting-check and the send call."""
    def send(self, text):
        return False


def test_answer_prompt_a_refused_send_remembers_nothing(conn, runs):
    from career_agent.web import actions
    run = runs[1] = RefusingRun("n")
    run.waiting.set()
    pid = _open(conn, "choice", {"id": "q", "kind": "choice", "question": "Notice period?",
                                 "options": ["30", "60"], "memory_key": "notice_period"})
    r = actions.answer_prompt(conn, pid, {"answer": "30"})
    assert not r["ok"] and r["code"] == 409
    assert conn.execute("SELECT COUNT(*) n FROM qa_bank").fetchone()["n"] == 0


def test_answer_prompt_confirm_approve_bumps_memory_use_count(conn, runs):
    from career_agent.web import actions
    store.qa_remember(conn, "Notice period?", "30 days", memory_key="notice_period")
    run = runs[1] = FakeRun("n")
    run.waiting.set()
    pid = _open(conn, "confirm", {**_confirm(Name="Asha"), "memory_used": ["notice_period"]})
    assert actions.answer_prompt(conn, pid, {"decision": "approve"})["ok"]
    assert store.qa_by_key(conn, "notice_period")["use_count"] == 1


# -- Fix round 1 -------------------------------------------------------------

def test_answer_relayed_to_the_agent_defaults_remember_true(conn, runs):
    """The relay's own default must match the storage default (both True) --
    a review found the relay still defaulting to False while storage
    defaulted to True."""
    from career_agent.web import actions
    run = runs[1] = FakeRun("n")
    run.waiting.set()
    pid = _open(conn, "text", {"id": "q", "kind": "text", "question": "Why?"})
    assert actions.answer_prompt(conn, pid, {"answer": "because"})["ok"]
    assert json.loads(run.sent[0].split(":", 2)[2])["remember"] is True


def test_answer_prompt_survives_a_qa_remember_failure(conn, runs, monkeypatch):
    """A DB hiccup in qa_remember (e.g. 'database is locked') must not turn
    an already-delivered answer into a 500/refusal -- the send already
    happened and the summary must still post."""
    from career_agent.web import actions

    def boom(*a, **kw):
        raise sqlite3.OperationalError("database is locked")
    monkeypatch.setattr(store, "qa_remember", boom)

    run = runs[1] = FakeRun("n")
    run.waiting.set()
    pid = _open(conn, "choice", {"id": "q", "kind": "choice", "question": "Notice period?",
                                 "options": ["30", "60"], "memory_key": "notice_period"})
    r = actions.answer_prompt(conn, pid, {"answer": "30"})
    assert r["ok"], r
    texts = [m["content"] for m in chat.messages_after(conn, chat.conversation_for_job(conn, 1))]
    assert "Notice period? → 30" in texts


def test_answer_prompt_survives_a_qa_touch_failure(conn, runs, monkeypatch):
    from career_agent.web import actions

    def boom(*a, **kw):
        raise sqlite3.OperationalError("database is locked")
    monkeypatch.setattr(store, "qa_touch", boom)

    run = runs[1] = FakeRun("n")
    run.waiting.set()
    pid = _open(conn, "confirm", {**_confirm(Name="Asha"), "memory_used": ["notice_period"]})
    r = actions.answer_prompt(conn, pid, {"decision": "approve"})
    assert r["ok"], r
    texts = [m["content"] for m in chat.messages_after(conn, chat.conversation_for_job(conn, 1))]
    assert "Approved the application" in texts


def _in_flight(conn):
    conn.execute("INSERT INTO application (job_id, resume_version, status, started_at)"
                 " VALUES (1, 'base-v1', 'in_flight', datetime('now'))")
    conn.commit()


def _prompt_status(conn, pid):
    return conn.execute("SELECT status, answer FROM agent_prompt WHERE id = ?", (pid,)).fetchone()


def test_a_send_refused_after_recording_reopens_a_live_runs_card(conn, runs):
    """The agent can leave the wait between the live check and the send."""
    from career_agent.web import actions

    class Racing(FakeRun):
        def send(self, text):
            self.waiting.clear()
            return super().send(text)
    run = runs[1] = Racing("n")
    run.waiting.set()
    _in_flight(conn)
    pid = _open(conn, "confirm", _confirm(Name="A"))
    assert not actions.answer_prompt(conn, pid, {"decision": "approve"})["ok"]
    row = _prompt_status(conn, pid)
    assert row["status"] == "open" and row["answer"] is None


@pytest.mark.parametrize("ended", ["run_done", "no_in_flight_row"])
def test_a_refused_card_of_an_ended_run_is_never_reopened(conn, runs, ended):
    """I5: reopening after the run's expire_open_prompts leaves a zombie card
    a later run could receive."""
    from career_agent.web import actions

    class Racing(FakeRun):
        def send(self, text):
            if ended == "run_done":
                self.done.set()
            else:
                self.waiting.clear()
            return super().send(text)
    run = runs[1] = Racing("n")
    run.waiting.set()
    if ended == "run_done":
        _in_flight(conn)
    pid = _open(conn, "confirm", _confirm(Name="A"))
    assert not actions.answer_prompt(conn, pid, {"decision": "approve"})["ok"]
    assert _prompt_status(conn, pid)["status"] == "expired"
    assert chat.open_prompt_for_job(conn, 1) is None


def test_a_card_older_than_the_live_run_is_refused(conn, runs):
    """I5(b): a crash leftover must never answer a later run's wait."""
    from career_agent.apply.runner import RunEvents
    from career_agent.web import actions
    stale = _open(conn, "confirm", _confirm(Name="old"))
    run = runs[1] = FakeRun("n", RunEvents(prompt_baseline=stale))
    run.waiting.set()
    r = actions.answer_prompt(conn, stale, {"decision": "approve"})
    assert r == {"ok": False, "code": 409, "message": "That question is no longer open"}
    assert run.sent == []


def test_only_the_newest_open_card_is_answerable(conn, runs):
    from career_agent.web import actions
    run = runs[1] = FakeRun("n")
    run.waiting.set()
    older = _open(conn, "text", {"id": "a", "question": "Old?"})
    newer = _open(conn, "text", {"id": "b", "question": "New?"})
    assert actions.answer_prompt(conn, older, {"answer": "x"})["code"] == 409
    assert run.sent == []
    assert actions.answer_prompt(conn, newer, {"answer": "y"})["ok"]


async def test_a_stale_card_after_a_requeue_never_reaches_the_new_run(conn, runs, factory):
    """I5(a): a crashed run's CONFIRM stays open; the requeued run waiting on
    an ASK must not receive DECISION approve from it."""
    from career_agent.web import actions
    stale = _open(conn, "confirm", _confirm(Name="old"))

    async def fake(prompt, jid, nonce, events, session_id=None, resume=False):
        run = runs[jid] = FakeRun(nonce, events)
        assert _prompt_status(factory(), stale)["status"] == "expired"
        _waiting_on(run, events, "ask", {"id": "q", "kind": "text", "question": "Notice?",
                                          "options": [], "why": "", "memory_key": None,
                                          "default": None, "sensitive": False})
        r = actions.answer_prompt(factory(), stale, {"decision": "approve"})
        assert not r["ok"] and run.sent == []
        return AgentResult("draft_ready")

    r = await _submit(conn, mode="manual", run_agent=fake, conn_factory=factory)
    assert r["ok"], r


@pytest.mark.parametrize("code,status,event", [
    ("applied", "submitted", "submitted_without_decision"),
    ("draft_ready", "draft", "draft_without_decision")])
async def test_a_later_unapproved_confirm_is_not_recorded_as_approved(
        conn, runs, factory, code, status, event):
    """I1: the human approved CONFIRM #1; the agent then changed a field and
    emitted CONFIRM #2 with its RESULT in the same turn. #2's values went out
    and nobody approved them -- #1 must not be recorded as their review."""
    from career_agent.web import actions

    async def fake(prompt, jid, nonce, events, session_id=None, resume=False):
        run = runs[jid] = FakeRun(nonce, events)
        _waiting_on(run, events, "confirm", _confirm(Name="Asha", Phone="+1"))
        first = chat.open_prompt_for_job(factory(), jid)["id"]
        assert actions.answer_prompt(factory(), first, {"decision": "approve"}, factory)["ok"]
        events.on_confirm(_confirm(Name="Asha", Phone="+91"))   # RESULT turn: not waiting
        return AgentResult(code)

    await _submit(conn, mode="manual", run_agent=fake, conn_factory=factory)
    row = _apps(conn)[-1]
    assert row["status"] == status and json.loads(row["answers"]) == {}
    assert event in _event_types(conn) and "confirm_missing" not in _event_types(conn)
    if code == "applied":
        assert row["failure_reason"] == "submitted_without_decision"


async def test_confirm_missing_means_no_confirm_at_all(conn, runs, factory):
    await _submit(conn, mode="manual", run_agent=fake_agent(AgentResult("applied")),
                  conn_factory=factory)
    assert _apps(conn)[-1]["failure_reason"] == "confirm_missing"


async def test_auto_approval_never_overrides_a_human_decision(conn, runs, factory, monkeypatch):
    """I3: if the human's answer lands first, the auto path must not send
    approve (nor claim it approved)."""
    real_open = chat.open_prompt

    def human_cancels_first(c, job_id, kind, payload):
        pid = real_open(c, job_id, kind, payload)
        chat.answer_prompt_row(c, pid, {"decision": "cancel"})
        return pid
    monkeypatch.setattr(ats_apply.chat, "open_prompt", human_cancels_first)

    async def fake(prompt, jid, nonce, events, session_id=None, resume=False):
        run = runs[jid] = FakeRun(nonce, events)
        _waiting_on(run, events, "confirm", _confirm(Name="Asha"))
        assert run.sent == [], "auto-approve overrode the human's cancel"
        return AgentResult("draft_ready")

    await _submit(conn, mode="auto", run_agent=fake, conn_factory=factory)
    texts = [m["content"] for m in chat.messages_after(conn, chat.conversation_for_job(conn, 1))]
    assert not any("approved without review" in t for t in texts)


@pytest.mark.parametrize("can_submit,expected", [(True, "held_unknown"), (False, None)])
async def test_answer_timeout_after_an_approve_holds_when_it_could_have_submitted(
        conn, runs, factory, monkeypatch, can_submit, expected):
    """I4: approve -> Submit -> a post-submit questionnaire ASKs -> silence.
    That run may have submitted; requeueing it would apply twice."""
    from career_agent.web import actions

    async def fake(prompt, jid, nonce, events, session_id=None, resume=False):
        run = runs[jid] = FakeRun(nonce, events)
        _waiting_on(run, events, "confirm", _confirm(Name="Asha"))
        pid = chat.open_prompt_for_job(factory(), jid)["id"]
        assert actions.answer_prompt(factory(), pid, {"decision": "approve"}, factory)["ok"]
        return AgentResult("failed", "answer_timeout")

    if can_submit:
        await _submit(conn, mode="manual", run_agent=fake, conn_factory=factory)
    else:
        monkeypatch.setattr(ats_apply, "_live_run_agent", fake)
        monkeypatch.setattr(ats_apply, "preflight", lambda: None)
        await _submit(conn, mode="manual", conn_factory=factory)
    if expected is None:        # nothing could have been sent: a resumable stop
        assert _apps(conn) == [] and checkpoint.get(conn, 1)["status"] == "resumable"
        return
    row = _apps(conn)[-1]
    assert (row["status"], row["failure_reason"]) == (expected, "answer_timeout")


async def test_a_cancel_before_the_run_registers_marks_it_cancelled(conn, runs):
    """Fold: cancellation can land before run_session registers in RUNS; the
    run's own cancel flag is what stops it spawning claude."""
    got = {}
    started = asyncio.Event()

    async def hangs(prompt, jid, nonce, events, session_id=None, resume=False):
        got["events"] = events
        started.set()
        await asyncio.sleep(30)

    task = asyncio.create_task(_submit(conn, mode="manual", run_agent=hangs))
    await started.wait()
    task.cancel()
    with pytest.raises(asyncio.CancelledError):
        await task
    assert got["events"].cancelled.is_set()


# -- outcomes by can_submit ------------------------------------------------

async def test_unknown_state_holds_only_when_submission_was_possible(conn, monkeypatch, no_live_runner):
    r = await _submit(conn, mode="manual", run_agent=fake_agent(AgentResult("failed", "no_result_line")))
    assert _apps(conn)[-1]["status"] == "held_unknown"


async def test_cancelled_is_a_human_decision_not_a_retry(conn):
    r = await _submit(conn, mode="manual", run_agent=fake_agent(AgentResult("failed", "cancelled")))
    row = _apps(conn)[-1]
    assert (row["status"], row["failure_reason"]) == ("failed_permanent", "cancelled")


@pytest.mark.parametrize("can_submit", [True, False])
async def test_answer_timeout_without_an_approve_is_resumable(conn, monkeypatch, can_submit):
    result = AgentResult("failed", "answer_timeout")
    if can_submit:
        runner = fake_agent(result)
    else:
        runner = None
        _use_live_fake(monkeypatch, result)
    r = await _submit(conn, mode="manual", run_agent=runner)
    assert r["resumable"] and _apps(conn) == []
    assert checkpoint.get(conn, 1)["status"] == "resumable"


def _use_live_fake(monkeypatch, result):
    """can_submit=False needs run_agent=None: stand a fake in for the live
    runner (and its preflight) instead."""
    fake = fake_agent(result)
    monkeypatch.setattr(ats_apply, "_live_run_agent", fake)
    monkeypatch.setattr(ats_apply, "preflight", lambda: None)
    return fake


async def test_cancelling_the_awaiting_task_kills_the_live_run(conn, runs):
    started = asyncio.Event()

    async def hangs(prompt, jid, nonce, events, session_id=None, resume=False):
        runs[jid] = FakeRun(nonce, events)
        started.set()
        await asyncio.sleep(30)

    task = asyncio.create_task(_submit(conn, mode="manual", run_agent=hangs))
    await started.wait()
    task.cancel()
    with pytest.raises(asyncio.CancelledError):
        await task
    assert getattr(runs[1], "killed", False)


# -- Task 9: a needs_answer park surfaces as a chat text card ---------------

async def test_needs_answer_opens_a_text_card_in_the_job_chat(conn):
    r = await _submit(conn, mode="manual",
                      run_agent=fake_agent(AgentResult("needs_answer", "Do you have a PMP?")))
    assert r["needs_answer"] == "Do you have a PMP?"
    row = chat.open_prompt_for_job(conn, 1)
    assert row is not None and row["kind"] == "text"
    assert json.loads(row["payload"]) == {
        "id": "needs_answer", "kind": "text", "question": "Do you have a PMP?",
        "why": "The agent stopped to ask this before continuing.",
        "origin": "needs_answer", "memory_key": None, "default": None,
        "options": [], "sensitive": False}


# -- S3: checkpoints and resume (Task 10) -----------------------------------

from career_agent.apply import checkpoint  # noqa: E402


def _recording(result, calls, on_run=None):
    """A runner that records the resume kwargs submit() hands it."""
    async def _fake(prompt, jid, nonce, events, session_id=None, resume=False):
        calls.append({"prompt": prompt, "nonce": nonce, "session_id": session_id,
                      "resume": resume, "cp": checkpoint.get(_fake.conn, jid)})
        if on_run:
            return on_run(prompt, jid, nonce, events, session_id, resume)
        return result
    return _fake


def _queue_ready(conn):
    conn.execute("INSERT INTO assessment (job_id, stage, weighted_score, verdict, rationale,"
                 " prompt_version, model) VALUES (1, 'scored', 80, 'submit', 'r', 'v', 'm')")
    conn.commit()


async def test_a_fresh_run_starts_a_checkpoint_with_its_session_id(conn):
    calls = []
    fake = _recording(AgentResult("draft_ready"), calls)
    fake.conn = conn
    await _submit(conn, mode="manual", run_agent=fake)
    c = calls[0]
    assert c["resume"] is False and c["session_id"]
    assert c["cp"]["status"] == "running" and c["cp"]["session_id"] == c["session_id"]
    assert c["cp"]["nonce"] == c["nonce"]
    assert checkpoint.get(conn, 1)["status"] == "done"


async def test_ask_and_confirm_mark_the_checkpoint_waiting(conn, runs, factory):
    seen = []

    async def fake(prompt, jid, nonce, events, session_id=None, resume=False):
        run = runs[jid] = FakeRun(nonce, events)

        def on_thread():      # reader-thread callbacks: the caller's conn is unusable there
            _waiting_on(run, events, "ask", {"id": "q1", "kind": "text", "question": "Why?",
                                             "options": [], "sensitive": False})
        t = threading.Thread(target=on_thread)
        t.start()
        t.join()
        seen.append(checkpoint.get(factory(), jid))
        return AgentResult("draft_ready")

    await _submit(conn, mode="manual", run_agent=fake, conn_factory=factory)
    pid = conn.execute("SELECT id FROM agent_prompt").fetchone()["id"]
    assert (seen[0]["status"], seen[0]["open_prompt_id"]) == ("waiting", pid)


@pytest.mark.parametrize("reason", ["timeout", "answer_timeout", "agent_error"])
async def test_a_resumable_stop_consumes_no_attempt(conn, reason):
    _queue_ready(conn)
    if reason == "agent_error":
        async def runner(prompt, jid, nonce, events, session_id=None, resume=False):
            raise RuntimeError("claude died")
    else:
        runner = fake_agent(AgentResult("failed", reason))
    for _ in range(ats_apply.MAX_ATTEMPTS + 1):
        r = await _submit(conn, mode="manual", run_agent=runner)
        assert not r["ok"] and r["resumable"]
        assert checkpoint.get(conn, 1)["status"] == "resumable"
    assert _apps(conn) == []                       # no failed / failed_permanent rows
    assert _event_types(conn).count("resumable") == ats_apply.MAX_ATTEMPTS + 1
    from career_agent.web import worker
    assert worker.next_candidate(conn) is None     # not re-picked as a fresh job
    checkpoint.finish(conn, 1)
    assert worker.next_candidate(conn)["job_id"] == 1


async def test_a_resumable_stop_holds_after_an_approve_when_it_could_submit(conn, runs, factory):
    """Narrowed I4: the approve may have been followed by a Submit click."""
    from career_agent.web import actions

    async def fake(prompt, jid, nonce, events, session_id=None, resume=False):
        run = runs[jid] = FakeRun(nonce, events)
        _waiting_on(run, events, "confirm", _confirm(Name="Asha"))
        pid = chat.open_prompt_for_job(factory(), jid)["id"]
        assert actions.answer_prompt(factory(), pid, {"decision": "approve"}, factory)["ok"]
        return AgentResult("failed", "timeout")

    r = await _submit(conn, mode="manual", run_agent=fake, conn_factory=factory)
    assert not r.get("resumable")
    assert _apps(conn)[-1]["status"] == "held_unknown"
    assert checkpoint.get(conn, 1)["status"] == "done"


async def test_a_resumable_stop_in_auto_mode_still_holds_when_it_could_submit(conn):
    """Auto mode is pre-approved from the start and CONFIRM+Submit can share
    one turn, so no recorded approve does not mean nothing was sent."""
    r = await _submit(conn, mode="auto", run_agent=fake_agent(AgentResult("failed", "timeout")))
    assert _apps(conn)[-1]["status"] == "held_unknown" and not r.get("resumable")
    assert checkpoint.get(conn, 1)["status"] == "done"


async def test_a_cancelled_run_is_left_resumable(conn):
    started = asyncio.Event()

    async def hangs(prompt, jid, nonce, events, session_id=None, resume=False):
        started.set()
        await asyncio.sleep(30)

    task = asyncio.create_task(_submit(conn, mode="manual", run_agent=hangs))
    await started.wait()
    task.cancel()
    with pytest.raises(asyncio.CancelledError):
        await task
    assert _apps(conn) == []
    assert checkpoint.get(conn, 1)["status"] == "resumable"


def _resumable(conn, answers):
    # an injected runner means can_submit; these tests resume in manual mode
    checkpoint.start(conn, 1, "sess-old", "oldnonce", mode="manual", can_submit=True)
    checkpoint.mark_running(conn, 1, "answered 4", answers)
    checkpoint.mark_resumable(conn, 1)


async def test_resume_reuses_the_session_and_nonce_and_sends_continue(conn):
    _resumable(conn, {"Notice?": "30 days"})
    calls = []
    fake = _recording(AgentResult("draft_ready"), calls)
    fake.conn = conn
    r = await _submit(conn, mode="manual", run_agent=fake, resume=True)
    assert r["ok"]
    c = calls[0]
    assert (c["resume"], c["session_id"], c["nonce"]) == (True, "sess-old", "oldnonce")
    assert c["prompt"].startswith("CONTINUE:oldnonce:")
    assert "== PREVIOUSLY ANSWERED" in c["prompt"] and "- Notice? -> 30 days" in c["prompt"]
    assert checkpoint.get(conn, 1)["status"] == "done"


async def test_resume_refuses_without_a_resumable_checkpoint(conn):
    fake = fake_agent(AgentResult("draft_ready"))
    r = await _submit(conn, mode="manual", run_agent=fake, resume=True)
    assert not r["ok"] and "resumable" in r["reason"] and fake.prompts == []
    checkpoint.start(conn, 1, "s", "n")                     # running, not resumable
    r = await _submit(conn, mode="manual", run_agent=fake, resume=True)
    assert not r["ok"] and fake.prompts == [] and _apps(conn) == []


async def test_a_failed_resume_falls_back_to_a_fresh_pinned_run(conn, caplog):
    _resumable(conn, {"Notice?": "30 days"})
    calls = []

    def on_run(prompt, jid, nonce, events, session_id, resume):
        if resume:
            return AgentResult("failed", "agent_error")     # no output at all
        return AgentResult("draft_ready")

    fake = _recording(None, calls, on_run)
    fake.conn = conn
    with caplog.at_level(logging.INFO):
        r = await _submit(conn, mode="manual", run_agent=fake, resume=True)
    assert r["ok"] and r["status"] == "draft"
    assert [c["resume"] for c in calls] == [True, False]
    fresh = calls[1]
    assert fresh["session_id"] not in (None, "sess-old")
    assert not fresh["prompt"].startswith("CONTINUE:")
    assert "- Notice? -> 30 days" in fresh["prompt"] and "== JOB" in fresh["prompt"]
    assert fresh["cp"]["session_id"] == fresh["session_id"]
    assert fresh["cp"]["answers"] == {"Notice?": "30 days"}
    assert "resume_fallback" in caplog.text
    assert [a["status"] for a in _apps(conn)] == ["draft"]


async def test_a_resume_that_produced_output_does_not_fall_back(conn):
    _resumable(conn, {})
    calls = []

    def on_run(prompt, jid, nonce, events, session_id, resume):
        events.on_text("Back on the form")
        return AgentResult("failed", "stuck")

    fake = _recording(None, calls, on_run)
    fake.conn = conn
    await _submit(conn, mode="manual", run_agent=fake, resume=True)
    assert len(calls) == 1 and _apps(conn)[-1]["status"] == "failed"


async def test_answer_prompt_records_the_answer_into_the_checkpoint(conn, runs, factory):
    from career_agent.web import actions

    async def fake(prompt, jid, nonce, events, session_id=None, resume=False):
        run = runs[jid] = FakeRun(nonce, events)
        _waiting_on(run, events, "ask", {"id": "q1", "kind": "text", "question": "Notice?",
                                         "options": [], "sensitive": False})
        pid = chat.open_prompt_for_job(factory(), jid)["id"]
        assert actions.answer_prompt(factory(), pid, {"answer": "30"}, factory)["ok"]
        cp = checkpoint.get(factory(), jid)
        assert (cp["status"], cp["step"], cp["answers"]) == ("running", f"answered {pid}",
                                                              {"Notice?": "30"})
        # a refused send (the run is no longer waiting) records nothing
        _waiting_on(run, events, "ask", {"id": "q2", "kind": "text", "question": "Visa?",
                                         "options": [], "sensitive": False})
        run.waiting.clear()
        pid2 = chat.open_prompt_for_job(factory(), jid)["id"]
        assert not actions.answer_prompt(factory(), pid2, {"answer": "no"}, factory)["ok"]
        assert checkpoint.get(factory(), jid)["answers"] == {"Notice?": "30"}
        return AgentResult("draft_ready")

    assert (await _submit(conn, mode="manual", run_agent=fake, conn_factory=factory))["ok"]


# -- Task 10 fix round 1 ------------------------------------------------------

@pytest.mark.parametrize("broken", ["mark_waiting", "mark_running"])
async def test_a_failing_checkpoint_write_never_blocks_the_auto_approve(conn, runs, factory,
                                                                        monkeypatch, broken):
    """I1: the approve is on record, so it must be sent -- else answer_timeout
    with an approve on record holds a job that sent nothing."""
    def boom(*a, **kw):
        raise sqlite3.OperationalError("database is locked")
    monkeypatch.setattr(checkpoint, broken, boom)

    async def fake(prompt, jid, nonce, events, session_id=None, resume=False):
        run = runs[jid] = FakeRun(nonce, events)
        _waiting_on(run, events, "confirm", _confirm(Name="Asha"))
        assert run.sent == [f'DECISION:{nonce}:{{"decision": "approve"}}']
        return AgentResult("applied")

    r = await _submit(conn, mode="auto", run_agent=fake, conn_factory=factory)
    assert r["ok"] and _apps(conn)[-1]["status"] == "submitted"


async def test_an_error_result_alone_falls_back_to_a_fresh_run(conn):
    """I4: a resumed session that only returns an is_error result gave no
    assistant content -- a turn end is not a sign of life."""
    _resumable(conn, {"Notice?": "30 days"})
    calls = []

    def on_run(prompt, jid, nonce, events, session_id, resume):
        if resume:
            events.on_turn_end(0.0, {})
            return AgentResult("failed", "no_result_line")
        return AgentResult("draft_ready")

    fake = _recording(None, calls, on_run)
    fake.conn = conn
    r = await _submit(conn, mode="manual", run_agent=fake, resume=True)
    assert r["ok"] and [c["resume"] for c in calls] == [True, False]


def _brief_file(tmp_path):
    p = tmp_path / "career_brief.toml"
    p.write_text('target_titles = ["AI Engineer"]\nsearch_locations = ["Chennai"]\ndaily_cap = 5\n')
    return p


def test_retry_supersedes_a_resumable_checkpoint(conn):
    """I2: crash -> resumable -> held -> human confirms not submitted -> Retry
    must actually requeue."""
    from career_agent.web import actions, worker

    _queue_ready(conn)
    _resumable(conn, {})
    _held(conn)
    assert actions.queue_retry(conn, 1, confirm_not_submitted=True)["ok"]
    assert checkpoint.get(conn, 1)["status"] == "done"
    assert worker.next_candidate(conn)["job_id"] == 1


def test_mark_applied_supersedes_a_resumable_checkpoint(conn, tmp_path):
    from career_agent.web import actions

    _resumable(conn, {})
    assert actions.mark_applied(conn, 1, "2026-09-01", _brief_file(tmp_path))["ok"]
    assert checkpoint.get(conn, 1)["status"] == "done"


async def test_what_the_checkpoint_pins_from_answers(conn, runs, factory):
    """M3: a sensitive ASK answer is never pinned; a CONFIRM pins only its
    changes, never the unedited field list."""
    from career_agent.web import actions

    async def fake(prompt, jid, nonce, events, session_id=None, resume=False):
        run = runs[jid] = FakeRun(nonce, events)
        _waiting_on(run, events, "ask", {"id": "q1", "kind": "text", "question": "SSN?",
                                         "options": [], "sensitive": True})
        pid = chat.open_prompt_for_job(factory(), jid)["id"]
        assert actions.answer_prompt(factory(), pid, {"answer": "123"}, factory)["ok"]
        _waiting_on(run, events, "confirm", _confirm(Name="Asha", Phone="+1"))
        pid = chat.open_prompt_for_job(factory(), jid)["id"]
        assert actions.answer_prompt(factory(), pid, {"decision": "change",
                                                      "changes": {"Phone": "+91"}}, factory)["ok"]
        assert checkpoint.get(factory(), jid)["answers"] == {"Phone": "+91"}
        return AgentResult("draft_ready")

    assert (await _submit(conn, mode="manual", run_agent=fake, conn_factory=factory))["ok"]


# -- Task 11: resume cap, mode mismatch, approve_sent, fallback nonce ------------

def _resumable_as(conn, mode="manual", can_submit=True, answers=None):
    checkpoint.start(conn, 1, "sess-old", "oldnonce", mode=mode, can_submit=can_submit)
    checkpoint.mark_running(conn, 1, "answered 4", answers or {})
    checkpoint.mark_resumable(conn, 1)


async def test_resume_increments_the_resume_count(conn):
    _resumable_as(conn)
    await _submit(conn, mode="manual", run_agent=fake_agent(AgentResult("failed", "timeout")),
                  resume=True)
    cp = checkpoint.get(conn, 1)
    assert (cp["status"], cp["resume_count"]) == ("resumable", 1)


async def test_resume_past_the_cap_is_a_retryable_failure(conn):
    _resumable_as(conn)
    conn.execute("UPDATE apply_checkpoint SET resume_count = ?", (ats_apply.MAX_RESUMES,))
    conn.commit()
    fake = fake_agent(AgentResult("draft_ready"))
    r = await _submit(conn, mode="manual", run_agent=fake, resume=True)
    assert not r["ok"] and "resume_limit" in r["reason"] and fake.prompts == []
    assert [(a["status"], a["failure_reason"]) for a in _apps(conn)] == [("failed", "resume_limit")]
    assert checkpoint.get(conn, 1)["status"] == "done"
    cid = chat.conversation_for_job(conn, 1)
    assert any("resume limit" in m["content"] for m in chat.messages_after(conn, cid, 0))


@pytest.mark.parametrize("stored", [("auto", True), ("manual", False), (None, None)])
async def test_a_mode_or_can_submit_mismatch_starts_fresh(conn, stored):
    """Never resume an auto session as manual (it could submit without a DECISION)."""
    _resumable_as(conn, *stored)
    calls = []
    fake = _recording(AgentResult("draft_ready"), calls)
    fake.conn = conn
    r = await _submit(conn, mode="manual", run_agent=fake, resume=True)
    assert r["ok"] and [c["resume"] for c in calls] == [False]
    assert calls[0]["session_id"] != "sess-old" and calls[0]["nonce"] != "oldnonce"
    assert calls[0]["cp"]["mode"] == "manual" and calls[0]["cp"]["can_submit"] == 1
    cid = chat.conversation_for_job(conn, 1)
    assert any("fresh session" in m["content"] for m in chat.messages_after(conn, cid, 0))


async def test_a_matching_resume_says_so_in_chat(conn):
    _resumable_as(conn)
    await _submit(conn, mode="manual", run_agent=fake_agent(AgentResult("draft_ready")), resume=True)
    cid = chat.conversation_for_job(conn, 1)
    assert any("Resuming" in m["content"] for m in chat.messages_after(conn, cid, 0))


async def test_the_fallback_gets_a_new_nonce_and_ignores_the_old_one(conn):
    _resumable_as(conn, answers={"Notice?": "30 days"})
    calls = []

    def on_run(prompt, jid, nonce, events, session_id, resume):
        if resume:
            return AgentResult("failed", "agent_error")
        # a stale line stamped with the old nonce is no protocol line for this run
        return agent_mod.parse_result("RESULT:oldnonce:DRAFT_READY", nonce)

    fake = _recording(None, calls, on_run)
    fake.conn = conn
    r = await _submit(conn, mode="manual", run_agent=fake, resume=True)
    fresh = calls[1]
    assert fresh["nonce"] != "oldnonce" and "oldnonce" not in fresh["prompt"]
    assert f"RESULT:{fresh['nonce']}:" in fresh["prompt"]
    assert fresh["cp"]["nonce"] == fresh["nonce"] and fresh["cp"]["session_id"] == fresh["session_id"]
    assert fresh["cp"]["answers"] == {"Notice?": "30 days"}
    assert not r["ok"] and r.get("status") != "draft"          # the stale RESULT did not count


async def test_an_approve_is_recorded_on_the_checkpoint(conn, runs, factory):
    from career_agent.web import actions

    async def fake(prompt, jid, nonce, events, session_id=None, resume=False):
        run = runs[jid] = FakeRun(nonce, events)
        _waiting_on(run, events, "confirm", _confirm(Name="Asha"))
        pid = chat.open_prompt_for_job(factory(), jid)["id"]
        assert checkpoint.get(factory(), jid)["approve_sent"] == 0
        assert actions.answer_prompt(factory(), pid, {"decision": "approve"}, factory)["ok"]
        assert checkpoint.get(factory(), jid)["approve_sent"] == 1
        return AgentResult("draft_ready")

    await _submit(conn, mode="manual", run_agent=fake, conn_factory=factory)


async def test_an_auto_approve_is_recorded_on_the_checkpoint(conn, runs, factory):
    async def fake(prompt, jid, nonce, events, session_id=None, resume=False):
        run = runs[jid] = FakeRun(nonce, events)
        _waiting_on(run, events, "confirm", _confirm(Name="Asha"))
        assert checkpoint.get(factory(), jid)["approve_sent"] == 1
        return AgentResult("applied")

    await _submit(conn, mode="auto", run_agent=fake, conn_factory=factory)


# -- Task 11 fix round 1 ------------------------------------------------------

async def test_a_queued_resume_is_refused_when_the_session_changed_under_it(conn):
    """I5: B read the checkpoint before the lock; A fell back meanwhile (new
    session and nonce) and stopped resumable again. B must not continue A's
    stale session."""
    _resumable_as(conn)
    fake = fake_agent(AgentResult("draft_ready"))
    lock = ats_apply._agent_lock()
    await lock.acquire()
    try:
        task = asyncio.create_task(_submit(conn, mode="manual", run_agent=fake, resume=True))
        await asyncio.sleep(0.1)                            # B is now waiting on the lock
        checkpoint.restart(conn, 1, "sess-new", "newnonce")
        checkpoint.mark_resumable(conn, 1)
    finally:
        lock.release()
    r = await task
    assert not r["ok"] and "resumable" in r["reason"] and fake.prompts == []
    assert _apps(conn) == [] and checkpoint.get(conn, 1)["session_id"] == "sess-new"


async def test_a_mismatch_fresh_start_carries_the_resume_count(conn):
    """M5."""
    _resumable_as(conn, "auto", True)
    conn.execute("UPDATE apply_checkpoint SET resume_count = 2")
    conn.commit()
    calls = []
    fake = _recording(AgentResult("draft_ready"), calls)
    fake.conn = conn
    await _submit(conn, mode="manual", run_agent=fake, resume=True)
    assert calls[0]["resume"] is False and calls[0]["cp"]["resume_count"] == 2


async def test_after_a_fallback_an_ask_with_the_old_nonce_opens_no_card(conn, factory):
    """M3: through a real AgentRun reader on the fallback's events."""
    import test_runner as tr
    from career_agent.apply.runner import AgentRun

    _resumable_as(conn)
    seen = []

    async def fake(prompt, jid, nonce, events, session_id=None, resume=False):
        if resume:
            return AgentResult("failed", "agent_error")      # silent: falls back
        popen = tr.FakePopen()
        run = AgentRun(cmd=["x"], cwd=".", env={}, nonce=nonce, events=events,
                       popen=lambda *a, **kw: popen)
        run.start("fresh")
        popen.emit(tr._assistant('ASK:oldnonce:{"id":"q1","kind":"text","question":"Stale?"}'))
        popen.emit(tr._result())
        popen.close()
        assert run.wait(5)
        seen.append(nonce)
        return agent_mod.parse_result(run.transcript, nonce)

    r = await _submit(conn, mode="manual", run_agent=fake, conn_factory=factory, resume=True)
    assert seen and seen[0] != "oldnonce"
    assert conn.execute("SELECT COUNT(*) n FROM agent_prompt").fetchone()["n"] == 0
    assert not r["ok"]
